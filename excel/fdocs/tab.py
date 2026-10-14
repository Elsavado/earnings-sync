"""Transactions and payroll data (CSV, XLSX, JSON): checks the file is the right kind of table
and anonymises it column by column.

People columns (employee, payee, cardholder...) become 'Redacted Name', organisation columns
(vendor, merchant, supplier...) 'Redacted Company Name Inc'. Account, card, routing and tax
numbers, e-mails, phones and street addresses are removed. Other text cells go through the
document anonymiser (names, links, tickers).
"""
import csv
import itertools
import io
import json
import re

from models.xl import REDACTED

from .anon import PERSON, TICKER  # noqa: F401  (TICKER: same markers everywhere)

# A column of people: 'name', 'first_name', 'employee', 'payee_name'... but not 'object_name' or 'fund_name'.
PERSON_COL = re.compile(r'^(?:(?:full|first|last|middle|given|family|sur|employee|emp|payee|cardholder|card_holder|'
                        r'customer|client|recipient|beneficiary|sender|receiver|owner|approver|requester|worker|'
                        r'staff|person|official|elected_official)_?)?(?:name|init|initial)?$|^nameorig$|^namedest$', re.I)
# Public officials named by title in any text cell: 'Commissioner Arrington' → 'Commissioner Redacted Name'.
OFFICIAL = re.compile(r"\b(Commissioner|Council ?member|Councilwoman|Councilman|Mayor|Judge|Justice|Senator|Supervisor|"
                      r"Alderman|Alderwoman|Alder|Representative|Rep\.|Sen\.|Sheriff|Clerk|Treasurer|Assessor|Comptroller|"
                      r"Director|Chief|Chairman|Chairwoman|Chair|Trustee|Dr\.?|Mr\.?|Mrs\.?|Ms\.?)\s+"
                      r"(?:[A-Z]\.\s+)?[A-Z][a-z'’\-]+(?:\s+[A-Z][a-z'’\-]+)?(?![a-z])")
ORG_COL = re.compile(r'vendor|supplier|merchant|company|organi[sz]ation|business|counterparty|employer|bank_?name|'
                     r'contractor|payee_?business|issuer|institution', re.I)
SECRET_COL = re.compile(r'ssn|social|tax_?id|\btin\b|ein|account_?(no|num)|acct|iban|swift|routing|card_?(no|num)|'
                        r'\bpan\b|cvv|email|e-mail|phone|mobile|address|street|zip|postal|ip_?addr|device|passport|'
                        r'licen[cs]e_?(no|num)|dob|birth|identifier|employee_?id|emp_?id|badge|person_?id', re.I)
MONEY_COL = re.compile(r'amount|amt|debit|credit|payment|paid|price|value|total|balance|fee|charge|cost|expense|'
                       r'salary|pay|wage|gross|net|earnings|overtime|bonus|compensation', re.I)
PAY_COL = re.compile(r'salary|pay\b|pay_|gross|wage|earnings|overtime|bonus|compensation|regular', re.I)
TXN_COL = re.compile(r'amount|debit|credit|payment|transaction|txn|vendor|merchant|invoice|check|cheque|balance|'
                     r'charge|payout|refund|ledger|journal|account', re.I)
ORG_WORDS = re.compile(r'\b(inc|llc|ltd|corp|co|company|corporation|services|group|associates|bank|city of|county|'
                       r'department|university|hospital|systems|solutions|partners|holdings|trust|foundation|'
                       r'authority|district|school|church|lp|llp|plc|gmbh)\b\.?', re.I)


def _official(m):
    """'Commissioner Arrington' → 'Commissioner Redacted Name'; 'Chief Financial Officer' stays."""
    from wordfreq import zipf_frequency
    from .anon import first_names
    words = m.group(0)[len(m.group(1)):].split()
    if words[0].strip('.').lower() in first_names() and len(words) > 1:
        return f'{m.group(1)} {PERSON}'  # 'Mayor Jane Smith'
    if any(zipf_frequency(w.strip('.').lower(), 'en') >= 4.5 for w in words if len(w.strip('.')) > 1):
        return m.group(0)
    return f'{m.group(1)} {PERSON}'


def gate(kind, header, rows):
    """None when the table belongs to kind, otherwise the reason it does not."""
    if len(rows) < 100:
        return f'too few rows ({len(rows)})'
    if len(header) < 4:
        return 'too few columns'
    h = ' '.join(header)
    if not MONEY_COL.search(h):
        return 'no money column'
    if kind == 'payroll' and not PAY_COL.search(h):
        return 'not payroll data'
    if kind == 'transactions' and not TXN_COL.search(h):
        return 'not transaction data'
    return None


def _role(col):
    if SECRET_COL.search(col):
        return 'secret'
    if ORG_COL.search(col):
        return 'org'
    norm = re.sub(r'[\s\-.]+', '_', col.strip()).strip('_')
    if norm and PERSON_COL.match(norm):
        return 'person'
    return 'text'


class Table:
    def __init__(self, anon):
        self.anon = anon
        self.cache = {}

    def cell(self, role, v):
        if not isinstance(v, str) or not v.strip() or not any(ch.isalpha() or ch.isdigit() for ch in v):
            return v
        if role == 'secret':
            return 'Redacted'
        if role == 'org':
            return REDACTED
        if role == 'person':
            return REDACTED if ORG_WORDS.search(v) else PERSON
        if re.fullmatch(r'[\d.,\-+$€£% ()/:]+', v):
            return v
        if ' ' not in v.strip() and any(ch.isdigit() for ch in v) and '@' not in v:
            return v  # a code or reference number (check, invoice, document id): not text to anonymise
        if v not in self.cache:
            if len(self.cache) > 200000:
                self.cache.clear()
            out = OFFICIAL.sub(_official, v)
            self.cache[v] = self.anon.scrub(out)
        return self.cache[v]

    # --- CSV ---------------------------------------------------------------
    def csv(self, src, dst, kind):
        """Streams the file row by row, so datasets of a gigabyte or more never sit in memory."""
        head = open(src, 'rb').read(2 * 2**20)
        enc = 'latin-1'
        for e in ('utf-8-sig', 'cp1252'):
            try:
                head.decode(e)
                enc = e
                break
            except UnicodeDecodeError as err:
                if err.start > len(head) - 4:  # a character cut at the sample's end
                    enc = e
                    break
        try:
            dialect = csv.Sniffer().sniff(head[:20000].decode(enc, 'replace'), delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(16 * 2**20)
        with open(src, encoding=enc, errors='replace', newline='') as f:
            reader = csv.reader(f, dialect)
            header = next(reader, None)
            if not header:
                return 'empty'
            sample = []
            for r in reader:
                if any(c.strip() for c in r):
                    sample.append(r)
                if len(sample) >= 100:
                    break
            why = gate(kind, header, sample)
            if why:
                return why
            roles = [_role(c) for c in header]
            with open(dst, 'w', encoding='utf-8', newline='') as out:
                w = csv.writer(out, lineterminator='\n')
                w.writerow([self.anon.scrub(c) for c in header])
                for r in itertools.chain(sample, reader):
                    if any(c.strip() for c in r):
                        w.writerow([self.cell(roles[i] if i < len(roles) else 'text', v) for i, v in enumerate(r)])
        return None

    # --- XLSX --------------------------------------------------------------
    def xlsx(self, src, dst, kind):
        from openpyxl import load_workbook
        wb = load_workbook(src)
        best = None
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(c or '') for c in rows[0]]
            if gate(kind, header, rows[1:]) is None:
                best = ws
            roles = [_role(c) for c in header]
            for row in ws.iter_rows(min_row=2):
                for i, c in enumerate(row):
                    if isinstance(c.value, str) and not c.value.startswith('='):
                        c.value = self.cell(roles[i] if i < len(roles) else 'text', c.value)
            for c in ws[1]:
                if isinstance(c.value, str):
                    c.value = self.anon.scrub(c.value)
            ws.title = (self.anon.scrub(ws.title).replace(REDACTED, 'Redacted Co') or 'Sheet')[:31]
        if best is None:
            return 'no transaction or payroll table'
        for ws in wb.worksheets:
            for img in list(getattr(ws, '_images', [])):
                ws._images.remove(img)
            for row in ws.iter_rows():
                for c in row:
                    c.hyperlink = None
        p = wb.properties
        p.creator = p.lastModifiedBy = PERSON
        p.title = p.subject = p.description = p.keywords = None
        wb.save(dst)
        return None

    # --- JSON (e.g. payment API fixtures) --------------------------------------
    def json(self, src, dst, kind):
        data = json.load(open(src, encoding='utf-8'))
        records = _records(data)
        if len(records) < 100:
            return f'too few records ({len(records)})'
        keys = {k for r in records[:500] for k in r}
        why = gate(kind, sorted(keys), records)
        if why:
            return why

        def walk(v, key=''):
            if isinstance(v, dict):
                return {k: walk(x, k) for k, x in v.items()}
            if isinstance(v, list):
                return [walk(x, key) for x in v]
            if isinstance(v, str):
                return self.cell(_role(key), v)
            return v
        json.dump(walk(data), open(dst, 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
        return None


def _records(data):
    if isinstance(data, list):
        return [r for r in data if isinstance(r, dict)]
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list) and len(v) >= 10 and all(isinstance(r, dict) for r in v[:10]):
                return v
    return []
