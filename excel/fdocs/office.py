"""Workspace office kinds: spreadsheets (uploaded as Google Sheets), documents (Google Docs), slide decks (Google
Slides), PDF presentations and bookkeeping records.

Sources are openly licensed only: Zenodo records under a Creative Commons, CC0 or open-source licence (searched by
file type), the Government of Canada open-data portal (English resources), and US government open-data portals
(Socrata) for bookkeeping ledgers. People's names become 'Salvador Sekwa' (user, 2026-10-22), companies
'Redacted Company Name Inc'; e-mails, links, pictures and embedded objects are removed as for the other types.

English content from the US and Canada only, and unique (user, 2026-10-24; usca.py): one file per record or
dataset, records from US / Canadian authors or governments, English metadata, title and text, and no two files with
the same normalised title.
"""
import re
import time
from pathlib import Path

import requests

from . import anon as _anon
from . import sources
from . import tab as _tab
from . import usca
from .sources import _add as _queue, late

RULES = 2  # queued candidates from before the English / US-Canada / one-per-record rules are dropped


def _records(state):
    """Zenodo records and Canadian datasets already used (any file of them), so each is collected once."""
    rec = state.setdefault('office_records', {})
    if not rec:
        for key in list(state['seen']) + list(state['queued']):
            if key.startswith(('zenodo:', 'ckan:')):
                rec[':'.join(key.split(':', 2)[:2]) if key.startswith('zenodo:') else key] = 1
        rec[''] = 1  # marks the set as seeded
    return rec


def _add(state, kind, cand, record):
    """Queues one candidate per record and per normalised title."""
    rec, titles = _records(state), state.setdefault('office_titles', {}).setdefault(kind, {})
    t = usca.norm_title(cand.get('title') or '')
    if record in rec or (t and t in titles):
        return False
    cand['rules'] = RULES
    if not _queue(state, kind, cand):
        return False
    rec[record] = 1
    if t:
        titles[t] = 1
    return True


def prune(state, kind):
    """Drops candidates queued under the old rules (any language, any country, every file of a record)."""
    q = state['queue'].get(kind, [])
    keep = [c for c in q if c.get('rules') == RULES or c.get('source') == 'socrata']
    for c in q:
        if c not in keep:
            state['queued'].pop(c['key'], None)
    state['queue'][kind] = keep
    return len(q) - len(keep)

KINDS = ('gsheets', 'gdocs', 'slides', 'presentations', 'bookkeeping')
PERSON = 'Salvador Sekwa'
GOOGLE = {'gsheets': 'application/vnd.google-apps.spreadsheet', 'gdocs': 'application/vnd.google-apps.document',
          'slides': 'application/vnd.google-apps.presentation'}
ZENODO_LIC = re.compile(r'^(?:cc-by|cc0|cc-zero|mit|apache|bsd|gpl|lgpl|agpl|mpl|isc|unlicense|odc|pddl)', re.I)
# Topics rotate so the collection covers business, finance, public sector, education and science.
TOPICS = ['budget', 'financial', 'accounting', 'business plan', 'project management', 'marketing', 'sales', 'economics',
          'annual report', 'policy', 'education', 'training', 'workshop', 'health', 'engineering', 'energy', 'climate',
          'agriculture', 'transport', 'survey', 'statistics', 'data analysis', 'software', 'research methods', 'strategy',
          'governance', 'procurement', 'human resources', 'operations', 'supply chain', 'risk', 'audit', 'tourism',
          'water', 'urban planning', 'public health', 'innovation', 'entrepreneurship', 'sustainability', '']
# A spreadsheet expands many times in memory when opened (a 40 MB workbook ran a runner out of memory).
MAX_SIZE = {'xlsx': 8 * 2**20, 'docx': 40 * 2**20, 'pptx': 40 * 2**20, 'pdf': 40 * 2**20}
ZFILE = {'gsheets': ('xlsx', None), 'gdocs': ('docx', None), 'slides': ('pptx', None), 'presentations': ('pdf', 'presentation'),
         'bookkeeping': ('xlsx', None)}
BOOK_Q = ['"general ledger"', 'bookkeeping', '"trial balance"', '"journal entries"', '"cash book"', '"accounts payable"',
          '"accounts receivable"', '"chart of accounts"', '"balance sheet"', '"income statement"', 'ledger', 'expenditure']
SOCRATA_BOOK = ['general ledger', 'journal entries', 'trial balance', 'check register', 'cash receipts', 'accounts receivable',
                'accounts payable', 'chart of accounts', 'fund balance', 'revenue ledger', 'expenditure ledger',
                'bank reconciliation', 'fixed assets', 'checkbook', 'vendor payments', 'disbursements', 'deposits']
LABEL = {'gsheets': 'Spreadsheet', 'gdocs': 'Document', 'slides': 'Slide deck', 'presentations': 'Presentation',
         'bookkeeping': 'Bookkeeping records'}
BOOK_COL = re.compile(r'debit|credit|ledger|journal|account|balance|fund|object|gl_?|posting|fiscal|period|voucher|'
                      r'check|cheque|receipt|deposit|invoice|payment|vendor|amount', re.I)


def zenodo_files(state, kind, s, log, pages=2):
    ext, rtype = ZFILE[kind]
    added = 0
    for topic in (BOOK_Q if kind == 'bookkeeping' else TOPICS):
        cur = state['cursors'].setdefault(f'office:zenodo:{kind}:{topic}', {'page': 1})
        if cur.get('done'):
            continue
        for _ in range(pages):
            if late():
                return added
            params = {'q': topic, 'file_type': ext, 'size': 25, 'page': cur['page'], 'sort': 'mostrecent'}
            if rtype:
                params['type'] = rtype
            try:
                r = s.get('https://zenodo.org/api/records', timeout=60, params=params)
                if r.status_code == 429:
                    time.sleep(60)
                    return added
                hits = r.json()['hits']['hits'] if r.ok else None
            except (requests.RequestException, ValueError, KeyError):
                hits = None
            if hits is None:
                break
            time.sleep(2.5)  # Zenodo allows about 30 searches a minute
            for h in hits:
                md = h.get('metadata') or {}
                if not ZENODO_LIC.match((md.get('license') or {}).get('id', '') or ''):
                    continue
                title = md.get('title') or ''
                if not (usca.zenodo_language(md) and usca.english_title(title) and usca.zenodo_usca(md)):
                    continue
                for f in h.get('files', [])[:20]:
                    key = f.get('key', '')
                    url = (f.get('links') or {}).get('self')
                    if key.lower().endswith('.' + ext) and url and 0 < f.get('size', 0) <= MAX_SIZE[ext]:
                        added += _add(state, kind, {'key': f'zenodo:{h["id"]}:{key}'[:300], 'url': url, 'ext': ext,
                                                    'orgs': [], 'source': 'zenodo', 'topic': topic.strip('"'),
                                                    'title': title[:200]}, f'zenodo:{h["id"]}')
                        break  # one file per record
            cur['page'] += 1
            if len(hits) < 25 or cur['page'] > 400:  # Zenodo serves at most 10,000 results per search
                cur['done'] = True
                break
    return added


CKAN_CA = 'https://open.canada.ca/data/api/action/package_search'
CKAN_FMT = {'gsheets': 'XLSX', 'gdocs': 'DOCX', 'slides': 'PPTX', 'presentations': 'PDF', 'bookkeeping': 'XLSX'}


def canada(state, kind, s, log, pages=2):
    """Government of Canada open data (Open Government Licence - Canada): the English resource of each dataset."""
    fmt, ext = CKAN_FMT[kind], ZFILE[kind][0]
    added = 0
    for topic in (BOOK_Q if kind == 'bookkeeping' else TOPICS):
        cur = state['cursors'].setdefault(f'office:canada:{kind}:{topic}', {'start': 0})
        if cur.get('done'):
            continue
        for _ in range(pages):
            if late():
                return added
            try:
                r = s.get(CKAN_CA, timeout=60, params={'q': topic.strip('"') or '*:*', 'fq': f'res_format:{fmt}',
                                                       'rows': 50, 'start': cur['start'], 'sort': 'metadata_modified desc'})
                res = r.json()['result'] if r.ok else None
            except (requests.RequestException, ValueError, KeyError):
                res = None
            if res is None:
                log(f'  canada {kind}: failed')
                break
            for p in res.get('results', []):
                if not str(p.get('license_id', '')).startswith('ca-ogl'):
                    continue
                title = (p.get('title_translated') or {}).get('en') or p.get('title') or ''
                if not usca.english_title(title):
                    continue
                for f in p.get('resources', []):
                    url = (f.get('url') or '').strip()
                    if url.startswith('/'):  # files hosted on the portal itself are listed by path
                        url = 'https://open.canada.ca' + url
                    if (str(f.get('format', '')).upper() == fmt and 'en' in (f.get('language') or [])
                            and url.lower().split('?')[0].endswith('.' + ext) and (f.get('size') or 0) <= MAX_SIZE[ext]):
                        added += _add(state, kind, {'key': f'ckan:open.canada.ca:{p["id"]}', 'url': url, 'ext': ext,
                                                    'orgs': [], 'source': 'canada', 'topic': topic.strip('"'),
                                                    'title': title[:200]}, f'ckan:open.canada.ca:{p["id"]}')
                        break  # one file per dataset
            cur['start'] += 50
            if cur['start'] >= res.get('count', 0) or cur['start'] >= 5000:
                cur['done'] = True
                break
    return added


def socrata_book(state, kind, s, log, pages=3):
    saved = sources.SOCRATA.get(kind)
    sources.SOCRATA[kind] = SOCRATA_BOOK
    try:
        return sources.socrata(state, kind, s, log, pages=pages)
    finally:
        if saved is None:
            sources.SOCRATA.pop(kind, None)


def _vet_socrata(state, kind):
    """Socrata datasets (US portals) queued by sources.socrata: English titles, one per normalised title."""
    titles = state.setdefault('office_titles', {}).setdefault(kind, {})
    keep = []
    for c in state['queue'].get(kind, []):
        if c.get('source') == 'socrata' and c.get('rules') != RULES:
            t = usca.norm_title(c.get('title') or '')
            if not usca.english_title(c.get('title') or 'x') or (t and t in titles):
                state['queued'].pop(c['key'], None)
                state['seen'][c['key']] = 'rej:duplicate or not English'
                continue
            c['rules'] = RULES
            if t:
                titles[t] = 1
        keep.append(c)
    state['queue'][kind] = keep


def start(state, kind, log):
    """At the start of a run, before any queued candidate is processed (main.Run.go)."""
    dropped = prune(state, kind)
    _vet_socrata(state, kind)
    if dropped:
        log(f'  {kind}: {dropped} queued candidate(s) from before the English / US-Canada rules dropped')


def office_sources(state, kind, s, log):
    turn = state['cursors'].setdefault(f'office:turn:{kind}', {'n': 0})
    turn['n'] += 1
    fns = [zenodo_files, canada] + ([socrata_book] if kind == 'bookkeeping' else [])
    fns = fns[turn['n'] % len(fns):] + fns[:turn['n'] % len(fns)]  # each source leads in turn
    added = 0
    for fn in fns:
        if not late():
            added += fn(state, kind, s, log)
    _vet_socrata(state, kind)
    return added


def _gate(kind, header, rows):
    """Table gate for the office kinds (replaces tab.gate's transaction/payroll checks)."""
    if kind == 'gsheets':
        if len(rows) < 10:
            return f'too few rows ({len(rows)})'
        if len([h for h in header if str(h).strip()]) < 3:
            return 'too few columns'
        return None if _english_table(header, rows) else 'not English'
    if kind == 'bookkeeping':
        if len(rows) < 30:
            return f'too few rows ({len(rows)})'
        h = ' '.join(str(c) for c in header)
        if len(set(m.group(0).lower() for m in BOOK_COL.finditer(h))) < 2 or not _tab.MONEY_COL.search(h):
            return 'not bookkeeping records'
        return None if _english_table(header, rows) else 'not English'
    return _ORIG_GATE(kind, header, rows)


def _english_table(header, rows):
    """Headers and text cells in English: Latin letters without accents, and few foreign function words."""
    cells = [str(c) for c in header] + [v for r in rows[:300] for v in r if isinstance(v, str) and not v.startswith('=')]
    text = ' '.join(cells)[:30000]
    letters = [ch for ch in text if ch.isalpha()]
    if letters and sum(not ch.isascii() for ch in letters) / len(letters) > 0.03:
        return False
    words = re.findall(r'[a-z]+', ' '.join(str(c) for c in header).lower())
    return not words or sum(w in usca.FOREIGN for w in words) / len(words) < 0.2


_ORIG_GATE = _tab.gate
_tab.gate = _gate


_NAME = r"[A-Z][\w'’.-]*"
# Author mentions in titles: 'for Smith et al.', 'by J. Smith and K. Lee', 'Smith & Lee (2021)'.
AUTHORS = re.compile(rf"(?:\b(?:for|from|by|of)\s+)?(?:{_NAME}\s+){{0,3}}{_NAME}\s+et\.?\s+al\.?,?|"
                     rf"\bby\s+{_NAME}(?:\s+{_NAME}){{0,2}}(?:\s*(?:,|and|&)\s*{_NAME}(?:\s+{_NAME}){{0,2}})*|"
                     rf"\b{_NAME}\s*(?:and|&)\s*{_NAME}\s*\(\d{{4}}\)")


def label(kind, cand, A=None):
    """'Document - <the record's own title>' (anonymised, short, safe as a file name); the search topic if no title."""
    t = re.sub(r'<[^>]+>', ' ', cand.get('title') or '')
    t = AUTHORS.sub(' ', t)  # 'Source Data for Josephson et al. ...' -> 'Source Data ...'
    if t and A is not None:
        t = A.scrub(t)
    t = re.sub(r'[\\/:*?"<>|\s]+', ' ', t).strip(' .-')
    if len(t) > 70:
        t = t[:70].rsplit(' ', 1)[0].rstrip(' ,;:-')
    if not t:
        t = (cand.get('topic') or '').strip().title()
    return f'{LABEL[kind]} - {t}' if t else LABEL[kind]


def table_text(path, ext, limit=800000):
    """Every text cell of a spreadsheet or CSV (up to limit characters): learnt before anonymising, checked after."""
    out, n = [], 0
    if ext == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        for ws in wb.worksheets:
            out.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and not v.startswith('='):
                        out.append(v)
                        n += len(v) + 1
                if n > limit:
                    break
        wb.close()
    else:
        with open(path, encoding='utf-8', errors='replace') as f:
            out.append(f.read(limit))
    return '\n'.join(out)[:limit]


def setup(kind):
    """People become Salvador Sekwa in every office kind (documents, tables and file properties)."""
    _anon.PERSON = PERSON
    _tab.PERSON = PERSON


def slide_count(path):
    import zipfile
    with zipfile.ZipFile(path) as z:
        return sum(1 for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n))


def landscape_share(path):
    import fitz
    with fitz.open(path) as d:
        pages = [p.rect for p in d]
    return len(pages), (sum(r.width > r.height for r in pages) / len(pages) if pages else 0)


DISCOVER = {k: [office_sources] for k in KINDS}
