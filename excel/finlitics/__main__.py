"""Financelitics collection (financial analytics firm: earnings call analysis, 10-K/10-Q review, risk modelling,
portfolio performance reporting, sector teams for banking, technology, healthcare, energy & utilities, industrials).

python -m finlitics --sources ... [--live]   (run from excel/). Files go to the Workspace shared drive folder
"Financelitics" with rclone; SEC web documents are converted to PDF. Logs print counts only."""
import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import time
import traceback

import requests

from avala.sink import Sink, State, safe

UA = os.environ.get("SEC_USER_AGENT") or "Financelitics research collector joseph@irtranscripts.com"
EARN = "01 Earnings Call Analysis"
FILE = "02 Filing Review (10-K & 10-Q)"
COMM = "03 SEC Comment Letters"
RISK = "04 Risk Modeling & Stress Testing"
PORT = "05 Portfolio Performance & Benchmarks"
BOOKS = "06 Books & Research"

SECTORS = {
    "Banking": ["JPM", "BAC", "WFC", "C", "GS", "MS", "USB", "PNC", "TFC", "COF", "SCHW", "BK", "STT", "FITB", "MTB",
                "HBAN", "RF", "KEY", "CFG", "ALLY", "ZION", "CMA", "FHN", "WAL", "EWBC"],
    "Technology": ["AAPL", "MSFT", "NVDA", "GOOGL", "META", "AMZN", "ORCL", "CRM", "ADBE", "INTC", "AMD", "CSCO", "IBM",
                   "NOW", "INTU", "QCOM", "TXN", "AVGO", "SNOW", "PANW", "CRWD", "WDAY", "ADSK", "MU", "PLTR"],
    "Healthcare": ["UNH", "JNJ", "LLY", "PFE", "MRK", "ABBV", "TMO", "ABT", "DHR", "BMY", "AMGN", "GILD", "CVS", "CI",
                   "ELV", "HCA", "MDT", "ISRG", "SYK", "BSX", "REGN", "VRTX", "ZTS", "HUM", "CNC"],
    "Energy & Utilities": ["XOM", "CVX", "COP", "EOG", "SLB", "OXY", "PSX", "MPC", "VLO", "KMI", "WMB", "NEE", "DUK",
                           "SO", "D", "AEP", "EXC", "SRE", "XEL", "PEG", "ED", "EIX", "WEC", "ETR", "PCG"],
    "Industrials": ["GE", "HON", "CAT", "DE", "UPS", "RTX", "LMT", "BA", "UNP", "MMM", "GD", "NOC", "ETN", "EMR", "ITW",
                    "CSX", "NSC", "FDX", "WM", "PH", "ROK", "CMI", "PCAR", "JCI", "TT"],
}
COMMENT_TOPICS = ["revenue recognition", "non-GAAP", "segment reporting", "goodwill impairment",
                  "allowance for credit losses", "critical audit matters", "income taxes valuation allowance",
                  "loss contingencies", "fair value measurements", "liquidity and capital resources",
                  "cybersecurity disclosure", "climate-related disclosure", "derivatives hedge accounting",
                  "business combinations purchase price allocation", "leases ASC 842", "risk factors"]
FRED = {  # workbook: series
    "Interest rates & yield curve": ["DGS1MO", "DGS3MO", "DGS6MO", "DGS1", "DGS2", "DGS5", "DGS7", "DGS10", "DGS20",
                                     "DGS30", "T10Y2Y", "T10Y3M", "DFF", "SOFR", "MORTGAGE30US"],
    "Credit spreads": ["BAMLC0A0CM", "BAMLH0A0HYM2", "BAMLC0A4CBBB", "BAA10Y", "AAA10Y", "TEDRATE"],
    "Macro": ["GDPC1", "UNRATE", "CPIAUCSL", "PCEPILFE", "INDPRO", "PAYEMS", "HOUST", "RSAFS", "UMCSENT"],
    "Markets & volatility": ["SP500", "NASDAQCOM", "VIXCLS", "DTWEXBGS", "DCOILWTICO", "GOLDAMGBD228NLBM"],
    "Banking system": ["TOTLL", "DPSACBW027SBOG", "DRCCLACBS", "DRALACBN", "TOTRESNS", "WALCL"],
}
FAMA_FRENCH = ["F-F_Research_Data_Factors", "F-F_Research_Data_5_Factors_2x3", "F-F_Momentum_Factor",
               "F-F_Research_Data_Factors_daily", "49_Industry_Portfolios", "12_Industry_Portfolios",
               "25_Portfolios_5x5", "Portfolios_Formed_on_ME", "F-F_ST_Reversal_Factor", "F-F_LT_Reversal_Factor"]
HF_QUERIES = ["earnings call", "earnings calls transcripts", "financial qa", "10-k", "sec filings", "financial news",
              "financial sentiment", "financial statements"]
HF_OPEN = {"apache-2.0", "mit", "cc-by-4.0", "cc-by-sa-4.0", "cc0-1.0", "odc-by", "odbl", "cc-by-nc-4.0",
           "cc-by-nc-sa-4.0", "bsd-3-clause", "openrail", "cc-by-3.0"}
BOOK_TOPICS = ["financial statement analysis", "financial reporting", "earnings management", "risk management",
               "credit risk", "market risk", "stress testing", "portfolio management", "asset pricing",
               "investment analysis", "corporate finance", "bank regulation", "financial econometrics", "valuation"]


def session(sec=False):
    s = requests.Session()
    s.headers["User-Agent"] = UA if sec else "Financelitics research collector (github.com/Elsavado/earnings-sync)"
    return s


def get(s, url, tries=4, **kw):
    for i in range(tries):
        try:
            r = s.get(url, timeout=kw.pop("timeout", 120), **kw)
            if r.status_code == 200:
                return r
            if r.status_code in (403, 404):
                return None
        except requests.RequestException:
            pass
        time.sleep(5 * 2 ** i)
    return None


def to_pdf(p):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not p.lower().endswith((".htm", ".html", ".txt")):
        return p
    try:
        filt = "pdf:writer_web_pdf_Export" if p.lower().endswith(("htm", "html")) else "pdf"
        subprocess.run([soffice, "--headless", "--norestore", "--convert-to", filt, "--outdir", os.path.dirname(p), p],
                       capture_output=True, timeout=300)
    except (subprocess.TimeoutExpired, OSError):
        return p
    pdf = os.path.splitext(p)[0] + ".pdf"
    if os.path.exists(pdf) and os.path.getsize(pdf):
        os.remove(p)
        return pdf
    return p


def fetch(s, sink, url, rel, cap=500 * 2 ** 20, pdf=True):
    p = sink.path(rel)
    n = 0
    try:
        with s.get(url, stream=True, timeout=180) as r:
            if r.status_code != 200:
                return False
            with open(p, "wb") as f:
                for chunk in r.iter_content(1 << 20):
                    n += len(chunk)
                    if n > cap:
                        break
                    f.write(chunk)
    except requests.RequestException:
        n = 0
    if not n or n > cap:
        if os.path.exists(p):
            os.remove(p)
        return False
    sink.wrote(to_pdf(p) if pdf else p)
    return True


def ciks(s, st):
    if not st.get("tickers"):
        r = get(s, "https://www.sec.gov/files/company_tickers.json")
        st["tickers"] = {v["ticker"]: v["cik_str"] for v in r.json().values()} if r else {}
    return st["tickers"]


# ---------------------------------------------------------------- sources

def filings(sink, st, years=3):
    """Latest annual and quarterly reports (10-K, 10-Q) and earnings releases (8-K item 2.02, EX-99.1) by sector."""
    s = session(True)
    tick = ciks(s, st)
    done = set(st.setdefault("acc_done", []))
    n = 0
    since = (dt.date.today() - dt.timedelta(days=365 * years)).isoformat()
    for sector, tickers in SECTORS.items():
        for t in tickers:
            cik = tick.get(t)
            if not cik or sink.full():
                continue
            time.sleep(0.15)
            r = get(s, f"https://data.sec.gov/submissions/CIK{int(cik):010d}.json")
            if not r:
                continue
            d = r.json()
            name = safe(d.get("name") or t, 60)
            rec = d["filings"]["recent"]
            for form, date, acc, doc, items in zip(rec["form"], rec["filingDate"], rec["accessionNumber"],
                                                   rec["primaryDocument"], rec.get("items", [""] * len(rec["form"]))):
                if date < since or acc in done or sink.full():
                    continue
                folder = f"{acc.replace('-', '')}"
                base = f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/{folder}"
                if form in ("10-K", "10-Q") and doc:
                    time.sleep(0.15)
                    if fetch(s, sink, f"{base}/{doc}", f"{FILE}/{sector}/{t} {name}/{date} {form}.{doc.rsplit('.', 1)[-1]}"):
                        n += 1
                elif form == "8-K" and "2.02" in (items or ""):
                    time.sleep(0.15)
                    idx = get(s, f"{base}/index.json")
                    for item in (idx.json().get("directory", {}).get("item", []) if idx else []):
                        nm = item.get("name", "")
                        if re.search(r"ex[-_]?99", nm, re.I) and nm.lower().endswith((".htm", ".html", ".txt", ".pdf")):
                            time.sleep(0.15)
                            if fetch(s, sink, f"{base}/{nm}",
                                     f"{EARN}/Earnings releases/{sector}/{t} {name}/{date} earnings release {nm}"):
                                n += 1
                            break
                done.add(acc)
            st["acc_done"] = sorted(done)
    return n


def comment_letters(sink, st, per_topic=40):
    """SEC staff comment letters (UPLOAD) and company responses (CORRESP): how filings are reviewed and challenged."""
    s = session(True)
    offs = st.setdefault("cl_from", {})
    seen = set(st.setdefault("cl_seen", []))
    n = 0
    for topic in COMMENT_TOPICS:
        took, start = 0, offs.get(topic, 0)
        while took < per_topic and not sink.full():
            time.sleep(0.3)
            r = get(s, "https://efts.sec.gov/LATEST/search-index", params={"q": f'"{topic}"', "forms": "UPLOAD,CORRESP",
                                                                          "from": start})
            hits = r.json().get("hits", {}).get("hits", []) if r else []
            if not hits:
                start = 0
                break
            start += len(hits)
            for h in hits:
                if h["_id"] in seen:
                    continue
                seen.add(h["_id"])
                src = h["_source"]
                acc, fname = h["_id"].split(":", 1)
                cik = (src.get("ciks") or ["0"])[0].lstrip("0")
                company = re.sub(r"\s*\(.*", "", (src.get("display_names") or ["Company"])[0])
                kind = "SEC staff letters" if src.get("file_type") == "UPLOAD" or "UPLOAD" in (src.get("root_forms") or []) \
                    else "Company responses"
                if fetch(s, sink, f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc.replace('-', '')}/{fname}",
                         f"{COMM}/{safe(topic.title(), 50)}/{kind}/{src.get('file_date', '')} {safe(company, 50)} {safe(fname, 30)}"):
                    took += 1
                    n += 1
                if took >= per_topic:
                    break
        offs[topic] = start
    st["cl_seen"] = sorted(seen)[-200000:]
    return n


def fred(sink, st):
    """FRED series for risk models (rates, spreads, macro, volatility, bank aggregates); refreshed weekly."""
    from openpyxl import Workbook
    stamp = st.get("fred")
    if stamp and stamp > (dt.date.today() - dt.timedelta(days=7)).isoformat():
        return 0
    s = session()
    n = 0
    for book, series in FRED.items():
        wb = Workbook()
        wb.remove(wb.active)
        for sid in series:
            r = get(s, "https://fred.stlouisfed.org/graph/fredgraph.csv", params={"id": sid})
            if not r:
                continue
            ws = wb.create_sheet(sid[:31])
            for line in r.text.strip().splitlines():
                parts = line.split(",")
                if len(parts) == 2:
                    try:
                        ws.append([parts[0], float(parts[1])])
                    except ValueError:
                        ws.append(parts)
            n += 1
        if wb.sheetnames:
            p = sink.path(f"{RISK}/Market & Macro Data (FRED)/{book} {dt.date.today()}.xlsx")
            wb.save(p)
            sink.wrote(p)
    st["fred"] = dt.date.today().isoformat()
    return n


def fdic(sink, st):
    """Quarterly financials of every FDIC-insured bank (assets, deposits, loans, capital, income, credit quality)."""
    s = session()
    done = set(st.setdefault("fdic_q", []))
    fields = ("CERT,NAMEFULL,STNAME,REPDTE,ASSET,DEP,LNLSNET,EQ,NETINC,ROA,ROE,NIMY,ELNATRY,NCLNLSR,RBC1AAJ,RBCRWAJ,"
              "LNRE,LNCI,LNCON,SC,INTINC,EINTEXP,NONII,NONIX")
    n = 0
    today = dt.date.today()
    for y in range(2015, today.year + 1):
        for md in ("0331", "0630", "0930", "1231"):
            rep = f"{y}{md}"
            if rep in done or dt.date(y, int(md[:2]), int(md[2:])) > today - dt.timedelta(days=75) or sink.full():
                continue
            r = get(s, "https://api.fdic.gov/banks/financials", params={"filters": f"REPDTE:{rep}", "fields": fields,
                                                                       "limit": 10000, "format": "csv"}, timeout=300)
            if r and len(r.content) > 1000:
                p = sink.path(f"{RISK}/Bank Financials (FDIC)/{y}/FDIC bank financials {rep}.csv")
                open(p, "wb").write(r.content)
                sink.wrote(p)
                done.add(rep)
                n += 1
    st["fdic_q"] = sorted(done)
    return n


def famafrench(sink, st):
    """Fama-French factor and portfolio returns: the standard benchmarks for performance attribution."""
    stamp = st.get("ff")
    if stamp and stamp > (dt.date.today() - dt.timedelta(days=30)).isoformat():
        return 0
    s = session()
    n = 0
    for name in FAMA_FRENCH:
        if sink.full():
            break
        if fetch(s, sink, f"https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/{name}_CSV.zip",
                 f"{PORT}/Factor & Benchmark Returns (Fama-French)/{name}.zip", pdf=False):
            n += 1
    st["ff"] = dt.date.today().isoformat()
    return n


def treasury(sink, st):
    """US Treasury daily par yield curve rates, one file per year."""
    s = session()
    done = set(st.setdefault("ust", []))
    n = 0
    for y in range(2000, dt.date.today().year + 1):
        key = str(y)
        if (key in done and y != dt.date.today().year) or sink.full():
            continue
        url = ("https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/"
               f"{y}/all?type=daily_treasury_yield_curve&field_tdr_date_value={y}&page&_format=csv")
        if fetch(s, sink, url, f"{RISK}/Treasury Yield Curve/Daily par yield curve {y}.csv", pdf=False):
            done.add(key)
            n += 1
    st["ust"] = sorted(done)
    return n


def transcripts(sink, st, per_query=4, cap=5 * 2 ** 30):
    """Openly licensed earnings-call transcript and financial text datasets on Hugging Face."""
    s = session()
    seen = set(st.setdefault("hf_seen", []))
    n = 0
    for q in HF_QUERIES:
        if sink.full():
            break
        r = get(s, "https://huggingface.co/api/datasets", params={"search": q, "sort": "downloads", "limit": 40, "full": "true"})
        took = 0
        for d in (r.json() if r else []):
            if took >= per_query or sink.full():
                break
            rid = d["id"]
            lic = (d.get("cardData") or {}).get("license")
            lic = (lic[0] if isinstance(lic, list) and lic else lic) or ""
            if rid in seen or d.get("gated") or str(lic).lower() not in HF_OPEN:
                continue
            seen.add(rid)
            t = get(s, f"https://huggingface.co/api/datasets/{rid}/tree/main", params={"recursive": "true"})
            files = [f for f in (t.json() if t else []) if f.get("type") == "file" and not f["path"].startswith(".")]
            if not files or sum(int(f.get("size") or 0) for f in files) > cap:
                continue
            folder = f"{EARN}/Earnings Call Transcripts" if "earn" in q else f"{BOOKS}/Financial Text Datasets"
            for f in files:
                fetch(s, sink, f"https://huggingface.co/datasets/{rid}/resolve/main/{f['path']}",
                      f"{folder}/{safe(rid.replace('/', ' - '), 80)}/{f['path']}", cap=3 * 2 ** 30, pdf=False)
            took += 1
            n += 1
    st["hf_seen"] = sorted(seen)
    return n


def books(sink, st, per_topic=8):
    """Open-access books on financial analysis, reporting, risk and portfolio management (OpenAlex, CC/public domain)."""
    s = session()
    seen = set(st.setdefault("book_seen", []))
    pages = st.setdefault("book_page", {})
    n = 0
    for topic in BOOK_TOPICS:
        if sink.full():
            break
        page = pages.get(topic, 1)
        r = get(s, "https://api.openalex.org/works", params={
            "filter": f"title.search:{topic},type:book|book-chapter,open_access.is_oa:true", "per-page": 50,
            "page": page, "sort": "cited_by_count:desc", "select": "id,title,publication_year,best_oa_location,type"})
        res = r.json().get("results", []) if r else []
        pages[topic] = page + 1 if res else 1
        took = 0
        for w in res:
            loc = w.get("best_oa_location") or {}
            if took >= per_topic or w["id"] in seen or not loc.get("pdf_url") or not str(loc.get("license", "")).startswith(("cc", "public")):
                continue
            seen.add(w["id"])
            if fetch(s, sink, loc["pdf_url"], f"{BOOKS}/{safe(topic.title(), 40)}/{safe(w.get('title') or w['id'], 100)} "
                                              f"({w.get('publication_year')}).pdf", cap=300 * 2 ** 20, pdf=False):
                took += 1
                n += 1
    st["book_seen"] = sorted(seen)
    return n


SOURCES = {"filings": filings, "comments": comment_letters, "fred": fred, "fdic": fdic, "famafrench": famafrench,
           "treasury": treasury, "transcripts": transcripts, "books": books}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sources", default=",".join(SOURCES))
    ap.add_argument("--out", default="out")
    ap.add_argument("--state", default="state/finlitics.json")
    ap.add_argument("--remote", default="ws:Financelitics")
    ap.add_argument("--live", action="store_true")
    ap.add_argument("--budget-gb", type=float, default=40)
    ap.add_argument("--minutes", type=float, default=320)
    a = ap.parse_args()
    st = State(a.state)
    sink = Sink(a.out, a.remote, a.live, st, int(a.budget_gb * 2 ** 30), time.time() + a.minutes * 60)
    names = [x.strip() for x in a.sources.split(",") if x.strip() in SOURCES]
    failed = []
    for i, name in enumerate(names):
        left = sink.deadline - time.time()
        slice_end = time.time() + max(15 * 60, left / (len(names) - i))
        real_deadline, sink.deadline = sink.deadline, min(sink.deadline, slice_end)
        t0 = time.time()
        try:
            n = SOURCES[name](sink, st.section(name))
            sink.flush()
            print(f"{name}: {n} items in {time.time() - t0:.0f}s; total {sink.files} files, {sink.sent / 2 ** 20:.0f} MB")
        except Exception:
            failed.append(name)
            print(f"{name}: FAILED")
            traceback.print_exc()
            try:
                sink.flush()
            except Exception:
                traceback.print_exc()
        sink.deadline = real_deadline
    print(f"done: {sink.files} files, {sink.sent / 2 ** 30:.2f} GB, live={a.live}")
    if failed:
        raise SystemExit(f"failed: {', '.join(failed)}")


if __name__ == "__main__":
    main()
