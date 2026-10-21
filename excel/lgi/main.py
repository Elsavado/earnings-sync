"""Builds every record of Little Games Incorporated and uploads them to the Shared Drive "Gaming Operations Data"
under "Little Games Incorporated/<system>". Environment: LGI_SCALE (1.0 = full size), DRY_RUN (write ./samples only).

python -m lgi.main"""
import io
import os
import shutil
import time
from pathlib import Path

import numpy as np
import pandas as pd

from gaming.out import Out, log
from . import aggregate, crm, dw, finance, mail, players as P, pm, support
from .config import COMPANY, END, GAMES, SEED, SYNTHETIC_NOTE

ROOT = COMPANY
CRM, ACC, SUP, PMF, DWF, MAIL = (f'{ROOT}/1 CRM', f'{ROOT}/2 Accounting', f'{ROOT}/3 Customer Support',
                                 f'{ROOT}/4 Project Management', f'{ROOT}/5 Data Warehouse', f'{ROOT}/6 Email Exports')
PART = 1_000_000
XLSX_MAX = 200_000


class Sink:
    def __init__(self, out):
        self.out = out
        self.parts = {}
        self.catalog = []

    def df(self, folder, base, frame, xlsx=None, part=False, about=''):
        if frame is None or not len(frame):
            return
        if part:
            k = (folder, base)
            n = self.parts.get(k, 0)
            for s in range(0, len(frame), PART):
                n += 1
                self.out.put(f'{folder}/{base}_part{n:03d}.csv', frame.iloc[s:s + PART].to_csv(index=False).encode('utf-8'), 'text/csv')
            self.parts[k] = n
        else:
            self.out.put(f'{folder}/{base}.csv', frame.to_csv(index=False).encode('utf-8'), 'text/csv')
            if xlsx is not False and len(frame) <= XLSX_MAX:
                self.out.put(f'{folder}/{base}.xlsx', xlsx_bytes({base: frame}),
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        if not any(c['folder'] == folder and c['table'] == base for c in self.catalog):
            self.catalog.append({'folder': folder.split('/', 1)[1], 'table': base, 'columns': ', '.join(map(str, frame.columns)), 'about': about})
        log(f'  {folder}/{base}: {len(frame):,} rows')


def xlsx_bytes(sheets):
    from openpyxl.styles import Font, PatternFill
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xw:
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor='1F3864')
            ws.freeze_panes = 'B2' if name.lower().startswith(('income', 'balance', 'cash')) else 'A2'
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(48, max(10, len(str(col[0].value)) + 4))
        xw.book.properties.creator = 'Salvador Sekwa'
    return buf.getvalue()


def crm_players(p):
    return pd.DataFrame({
        'player_id': p.player_id, 'account_number': p.account_number, 'username': p.username, 'first_name': p.first_name,
        'last_name': p.last_name, 'email': p.email, 'country': p.country, 'language': p.language, 'age_band': p.age_band,
        'game': p.game, 'platform': p.platform, 'signup_date': P.daystr(p.signup_day.values), 'last_active_date': P.daystr(p.last_active_day.values),
        'status': p.status, 'segment': p.segment, 'lifetime_spend_usd': p.lifetime_spend_usd, 'purchases': p.purchases,
        'first_purchase_date': P.daystr(p.first_purchase_day.values), 'acquisition': p.acquisition,
        'campaign_id': np.where(p.acquisition == 'Paid', [f'UA-{g}-{pl[:3].upper()}-T{aggregate.TIER[c]}-{d[:7]}' for g, pl, c, d in
                                                           zip(p.game, p.platform, p.country, P.daystr(p.signup_day.values))], ''),
        'marketing_opt_in': p.marketing_opt_in})


def main():
    t0 = time.time()
    dry = os.environ.get('DRY_RUN', 'false') == 'true'
    scale = float(os.environ.get('LGI_SCALE', '1.0'))
    out = Out(dry)
    sink = Sink(out)
    rng = np.random.default_rng(SEED + 4242)
    emp = finance.staff(np.random.default_rng(SEED + 99))
    agents = support.agents(emp, rng)
    work = Path('lgi_mail')
    shutil.rmtree(work, ignore_errors=True)
    box = mail.Box(work)
    aggs, tickets_monthly, start_no = [], {}, 0
    for gi, g in enumerate(GAMES):
        r = P.simulate_game(gi, g, scale)
        if r is None:
            continue
        players, txn = r
        log(f'{g["name"]}: {len(players):,} accounts, {len(txn):,} purchases')
        kpis = P.daily_kpis(gi, g, players, txn)
        aggs.append(aggregate.game_aggregates(players, txn, kpis, scale))
        sink.df(CRM, 'player_accounts', crm_players(players), part=True, about='Every player account (one row per account and game).')
        t = txn.assign(purchased_at=P.ts(txn.day.values, txn.seconds.values))[
            ['txn_id', 'player_id', 'purchased_at', 'game', 'platform', 'country', 'sku', 'item', 'kind', 'price_usd', 'status']]
        for y, grp in t.groupby(t.purchased_at.dt.year):
            sink.df(f'{DWF}/fact_purchase/{y}', 'fact_purchase', grp, part=True, about='Every purchase, by year.')
        sink.df(f'{DWF}/fact_daily_game_kpi', 'fact_daily_game_kpi', kpis, part=True,
                about='Daily KPIs per game, platform and country (aggregated from player activity and purchases).')
        sink.df(f'{DWF}/fact_retention_cohort', 'fact_retention_cohort', dw.retention(players, g), part=True,
                about='Monthly sign-up cohorts: share still playing after N days.')
        sink.df(f'{DWF}/dim_player', 'dim_player', pd.DataFrame({
            'player_id': players.player_id, 'game_code': players.game, 'platform': players.platform, 'country_code': players.country,
            'signup_date': P.daystr(players.signup_day.values), 'last_active_date': P.daystr(players.last_active_day.values),
            'acquisition': players.acquisition, 'segment': players.segment, 'status': players.status, 'age_band': players.age_band}), part=True)
        tk = support.tickets(gi, g, players, txn, kpis, agents, start_no)
        start_no += len(tk)
        msgs = support.messages(tk, g, np.random.default_rng(SEED + 1000 * gi + 12))
        for mo, n in pd.Series(tk.created_at).dt.to_period('M').astype(str).value_counts().items():
            tickets_monthly[mo] = tickets_monthly.get(mo, 0) + n / scale
        cols = ['ticket_id', 'player_id', 'created_at', 'solved_at', 'game', 'platform', 'country', 'category', 'group', 'channel',
                'priority', 'status', 'assignee_id', 'escalated', 'first_response_minutes', 'resolution_hours', 'sla_breached', 'csat_score', 'txn_id']
        sink.df(SUP, 'tickets', tk[cols], part=True, about='Every support ticket; refund and chargeback tickets link to their purchase (txn_id).')
        sink.df(SUP, 'ticket_messages', msgs, part=True, about='Conversation messages of each ticket.')
        sink.df(f'{DWF}/fact_support_ticket', 'fact_support_ticket', tk[cols].drop(columns=['group', 'assignee_id', 'escalated']), part=True)
        mail.support_threads(box, tk, msgs, agents)
        camps, recips = crm.email_campaigns(gi, g, players, np.random.default_rng(SEED + 1000 * gi + 13))
        sink.df(CRM, 'email_campaigns', camps, part=True, about='Lifecycle e-mails to opted-in active players: results per send.')
        sink.df(CRM, 'email_campaign_recipients_sample', recips, part=True, about='Recipient-level log for a 1% sample of each send.')
        del players, txn, kpis, tk, msgs
        log(f'  done {g["code"]} at {time.time() - t0:.0f}s')

    km, tm, pi = [pd.concat([a[i] for a in aggs], ignore_index=True) for i in range(3)]
    emp, pay, gl, ap, po, fa, ua, ar = finance.build(km, tm, pi, tickets_monthly, emp)
    tb, is_m, is_y, bs, cf, check = finance.statements(gl)
    log(f'accounting: {len(gl):,} ledger lines; balance check {check:.6f}')
    if check > 1:
        raise RuntimeError(f'balance sheet does not balance: {check}')
    sink.df(ACC, 'chart_of_accounts', pd.DataFrame(finance.COA, columns=['account', 'name', 'type', 'statement_section']))
    sink.df(ACC, 'general_ledger', gl, part=True, about='Every journal line (double entry).')
    sink.df(ACC, 'trial_balance_monthly', tb)
    out.put(f'{ACC}/Financial statements.xlsx', xlsx_bytes({'Income statement (annual)': is_y, 'Income statement (monthly)': is_m,
                                                             'Balance sheet (monthly)': bs, 'Cash flow by source': cf}),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    sink.df(ACC, 'vendors', pd.DataFrame(finance.VENDORS, columns=['vendor', 'category', 'default_account', 'payment_terms_days']))
    sink.df(ACC, 'customers', pd.DataFrame({'customer': sorted(set(finance.STORE.values()) | {'Ad mediation partners'})}))
    sink.df(ACC, 'ap_invoices', ap, about='Supplier invoices and their payment.')
    sink.df(ACC, 'purchase_orders', po)
    sink.df(ACC, 'ar_store_statements', ar, about='Monthly store statements (sales, refunds, fees, net due) and payouts.')
    sink.df(ACC, 'fixed_assets', fa)
    sink.df(ACC, 'employees', emp.drop(columns=[]), about='Staff roster (fictional people).')
    sink.df(ACC, 'payroll_register', pay, part=True, about='Monthly pay per employee.')
    sink.df(ACC, 'budget_vs_actual', budget(is_y))
    sink.df(CRM, 'ua_campaigns', ua, about='Paid user-acquisition campaigns: installs and spend (ties to accounts acquired).')
    acc_, con, opp, act = crm.b2b(emp)
    sink.df(CRM, 'partner_accounts', acc_)
    sink.df(CRM, 'partner_contacts', con)
    sink.df(CRM, 'opportunities', opp)
    sink.df(CRM, 'activities', act)
    sink.df(SUP, 'agents', agents)
    projects, versions, sprints, issues, logs = pm.build(emp, END)
    sink.df(PMF, 'projects', projects)
    sink.df(PMF, 'releases', versions)
    sink.df(PMF, 'sprints', sprints)
    sink.df(PMF, 'issues', issues, part=True)
    sink.df(PMF, 'worklogs', logs, part=True)
    mail.finance_mail(box, ap, emp, is_m)
    mail.company_mail(box, emp, sprints, projects)
    index = box.close()
    for p in sorted(work.rglob('*.mbox')):
        out.put(f'{MAIL}/{p.parent.name}/{p.name}', p.read_bytes(), 'application/mbox')
    sink.df(MAIL, 'message_index', index, part=True, about='Every e-mail with the record it belongs to.')
    shutil.rmtree(work, ignore_errors=True)
    sink.df(DWF, 'dim_date', dw.dim_date())
    sink.df(DWF, 'dim_game', dw.dim_game())
    sink.df(DWF, 'dim_platform', dw.dim_platform())
    sink.df(DWF, 'dim_country', dw.dim_country())
    sink.df(DWF, 'dim_employee', emp[['employee_id', 'department', 'title', 'level', 'location', 'hire_date', 'termination_date', 'project']])
    sink.df(DWF, 'dim_campaign', ua[['campaign_id', 'month', 'game', 'platform', 'country_tier', 'channel']])
    sink.df(DWF, 'fact_ua_spend', ua[['campaign_id', 'month', 'installs', 'spend_usd', 'cpi_usd']])
    sink.df(DWF, 'fact_finance_monthly', dw.finance_fact(gl), part=True)
    sink.df(DWF, 'fact_headcount_monthly', dw.headcount_fact(emp, finance.MONTHS))
    out.put(f'{DWF}/schema.sql', dw.DDL.encode('utf-8'), 'text/plain')
    cat = pd.DataFrame(sink.catalog)
    out.put(f'{ROOT}/Data dictionary.xlsx', xlsx_bytes({'About': pd.DataFrame({'Little Games Incorporated': [
        SYNTHETIC_NOTE, f'Records from 2015-01-01 to {END}.', 'Systems: CRM, Accounting, Customer Support, Project Management, '
        'Data Warehouse, Email Exports. All systems share the same ids (player_id, txn_id, ticket_id, employee_id, campaign_id).']}),
        'Tables': cat}), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    log(f'all done in {time.time() - t0:.0f}s')


def budget(is_y):
    """Budget set each December for the next year (prior actual x plan growth), against actual."""
    rng = np.random.default_rng(SEED + 31)
    years = [c for c in is_y.columns if c not in ('section', 'line')]
    rows = []
    for _, r in is_y.iterrows():
        for i in range(1, len(years)):
            prev, act = r[years[i - 1]], r[years[i]]
            bud = prev * rng.normal(1.12, .06)
            rows.append({'year': years[i], 'section': r.section, 'line': r.line, 'budget_usd': round(bud, 2), 'actual_usd': round(act, 2),
                         'variance_usd': round(act - bud, 2), 'variance_pct': round(100 * (act - bud) / abs(bud), 1) if bud else None})
    return pd.DataFrame(rows)


if __name__ == '__main__':
    main()
