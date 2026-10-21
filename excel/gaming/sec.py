"""Listed gaming companies: every reported financial line (annual and quarterly, as far back as filed) and the
operating KPIs they publish in their earnings releases (DAU, MAU, bookings, payers, ARPDAU, hours engaged, ...)."""
import io
import os
import re
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .out import FINANCE, PERSON, Table, log

H = {'User-Agent': os.environ.get('SEC_USER_AGENT') or 'GamingData research contact@example.com'}
# (name, CIK, segment)
COMPANIES = [
    ('Activision Blizzard', 718877, 'Publisher'), ('Electronic Arts', 712515, 'Publisher'),
    ('Take-Two Interactive', 946581, 'Publisher'), ('Roblox', 1315098, 'Platform'), ('Playtika', 1828016, 'Mobile'),
    ('Zynga', 1439404, 'Mobile'), ('Glu Mobile', 1366246, 'Mobile'), ('SciPlay', 1760717, 'Mobile'),
    ('PLAYSTUDIOS', 1823878, 'Mobile'), ('Skillz (Firy)', 1801661, 'Mobile'), ('DoubleDown Interactive', 1799567, 'Mobile'),
    ('GDEV', 1848739, 'Mobile'), ('Snail', 1886894, 'Publisher'), ('NetEase', 1110646, 'Publisher'),
    ('Sea (Garena)', 1703399, 'Publisher'), ('Bilibili', 1723690, 'Platform'), ('Gravity', 1313310, 'Publisher'),
    ('The9', 1296774, 'Publisher'), ('Sohu (Changyou)', 1734107, 'Publisher'), ('HUYA', 1728190, 'Streaming'),
    ('DouYu', 1762417, 'Streaming'), ('Unity Software', 1810806, 'Engine & Tools'), ('AppLovin', 1751008, 'Engine & Tools'),
    ('GameStop', 1326380, 'Retail'), ('Corsair Gaming', 1743759, 'Hardware'), ('GameSquare', 1714562, 'Esports & Media')]
KPI = re.compile(r'(?i)\b(daily active users?|DAUs?|monthly active users?|MAUs?|monthly active payers?|monthly unique payers?|'
                 r'MUPs?|daily paying users?|daily payers?|average daily payers|paying users?|payers|net bookings|bookings|'
                 r'average bookings per|ABPDAU|ARPDAU|average (daily )?revenue per|ARPPU|ARPU|hours engaged|engagement hours|'
                 r'recurrent consumer spending|live services|users?\b.*\b(million|thousand)|downloads|installs|subscribers|'
                 r'concurrent users|(?:player|user|day ?\d+) retention|retention rate)\b')
LAST = [0.0]


def sec(url):
    for i in range(6):
        wait = LAST[0] + 0.15 - time.time()
        if wait > 0:
            time.sleep(wait)
        LAST[0] = time.time()
        try:
            r = requests.get(url, headers=H, timeout=90)
        except requests.RequestException:
            time.sleep(5 * (i + 1))
            continue
        if r.status_code in (429, 503):
            time.sleep(30 * (i + 1))
            continue
        return r
    return None


def _days(a, b):
    from datetime import date
    return (date.fromisoformat(b) - date.fromisoformat(a)).days


def financials(cik):
    """{'annual': {label: {end: value}}, 'quarterly': {...}, 'units': {label: unit}}"""
    r = sec(f'https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json')
    if r is None or not r.ok:
        return None
    facts = r.json().get('facts') or {}
    out = {'annual': {}, 'quarterly': {}, 'units': {}, 'tags': {}}
    for tax in ('us-gaap', 'ifrs-full'):
        for tag, f in (facts.get(tax) or {}).items():
            units = f.get('units') or {}
            unit = 'USD' if 'USD' in units else next(iter(units), None)
            if unit is None:
                continue
            label = f.get('label') or tag
            best = {'annual': {}, 'quarterly': {}}
            for v in units[unit]:
                form, end, start = v.get('form') or '', v.get('end'), v.get('start')
                if form.startswith(('10-K', '20-F', '40-F')) and v.get('fp') == 'FY':
                    kind = 'annual' if start is None or 330 <= _days(start, end) <= 380 else None
                elif form.startswith('10-Q'):
                    kind = 'quarterly' if start is None or 80 <= _days(start, end) <= 100 else None
                else:
                    kind = None
                if kind:
                    prev = best[kind].get(end)
                    if prev is None or v.get('filed', '') >= prev[1]:
                        best[kind][end] = (v['val'], v.get('filed', ''))
            for kind in best:
                if len(best[kind]) >= 2:
                    out[kind].setdefault(label, {}).update({e: x[0] for e, x in best[kind].items()})
                    out['units'][label] = unit
                    out['tags'][label] = f'{tax}:{tag}'
    return out


def _filings(cik):
    r = sec(f'https://data.sec.gov/submissions/CIK{cik:010d}.json')
    if r is None or not r.ok:
        return []
    j = r.json()
    blocks = [j['filings']['recent']] + [
        (sec(f'https://data.sec.gov/submissions/{f["name"]}').json()) for f in j['filings'].get('files') or []]
    rows = []
    for b in blocks:
        for i in range(len(b.get('form') or [])):
            form, items = b['form'][i], (b.get('items') or [''] * len(b['form']))[i] or ''
            desc = ((b.get('primaryDocDescription') or [''] * len(b['form']))[i] or '').lower()
            if (form == '8-K' and '2.02' in items) or (form == '6-K' and re.search(r'result|earn|quarter', desc)):
                rows.append((b['filingDate'][i], b['accessionNumber'][i]))
    return rows


def _num(s):
    s = str(s).strip().replace(',', '').replace('$', '').replace('%', '')
    neg = s.startswith('(') and s.endswith(')')
    s = s.strip('()').strip()
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def kpis(name, cik):
    import pandas as pd
    rows = []
    for filed, acc in _filings(cik):
        base = f'https://www.sec.gov/Archives/edgar/data/{cik}/{acc.replace("-", "")}'
        idx = sec(f'{base}/index.json')
        if idx is None or not idx.ok:
            continue
        docs = [d['name'] for d in idx.json().get('directory', {}).get('item', [])
                if re.search(r'(?i)ex.?99|exhibit.?99|press', d['name']) and d['name'].lower().endswith(('.htm', '.html'))]
        for doc in docs[:2]:
            r = sec(f'{base}/{doc}')
            if r is None or not r.ok:
                continue
            try:
                tables = pd.read_html(io.StringIO(r.text))
            except Exception:  # noqa: BLE001 - a document without tables
                continue
            for t in tables:
                t = t.dropna(how='all').dropna(axis=1, how='all')
                if t.shape[0] < 2 or t.shape[1] < 2:
                    continue
                cells = [[str(x) for x in row] for row in t.values.tolist()]
                header = [' '.join(str(x) for x in c) if isinstance(c, tuple) else str(c) for c in t.columns]
                for line in cells:
                    label = str(line[0]).strip()
                    if re.search(r'(?i)(quarter|months|year|fy|q[1-4]|20\d\d)', ' '.join(line[1:])) and not KPI.search(label):
                        header = line
                        continue
                    if (not KPI.search(label) or label == 'nan' or label.endswith(':') or label.startswith('(')
                            or re.search(r'(?i)payments?|expense|adjusted ebitda|compensation', label)):
                        continue
                    seen = set()
                    for j, cell in enumerate(line[1:], 1):
                        v = _num(cell)
                        if v is None or (j, v) in seen:
                            continue
                        seen.add((j, v))
                        period = str(header[j]).strip() if header and j < len(header) and str(header[j]) != 'nan' else ''
                        rows.append({'company': name, 'filed': filed, 'metric': re.sub(r'\s+', ' ', label)[:120],
                                     'period': re.sub(r'\s+', ' ', period)[:60], 'value': v,
                                     'document': f'{acc}/{doc}'})
    # one value per (metric, period, value) as first reported
    uniq, seen = [], set()
    for r in sorted(rows, key=lambda r: r['filed']):
        k = (r['metric'].lower(), r['period'].lower(), r['value'])
        if k not in seen:
            seen.add(k)
            uniq.append(r)
    return uniq


def _sheet(wb, title, data, units, tags):
    ws = wb.create_sheet(title)
    ends = sorted({e for v in data.values() for e in v})
    head = ['Line item', 'Unit', 'Tag'] + ends
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3864')
    for label in sorted(data, key=str.lower):
        ws.append([label, units.get(label), tags.get(label)] + [data[label].get(e) for e in ends])
    ws.freeze_panes = 'D2'
    ws.column_dimensions['A'].width = 60
    return ends


SUMMARY = {'revenue': ['Revenues', 'Revenue', 'Revenue from Contract with Customer, Excluding Assessed Tax', 'Sales Revenue, Net'],
           'gross_profit': ['Gross Profit', 'Gross profit'], 'rnd': ['Research and Development Expense'],
           'sales_marketing': ['Selling and Marketing Expense', 'Marketing and Advertising Expense', 'Advertising Expense'],
           'gna': ['General and Administrative Expense'], 'operating_income': ['Operating Income (Loss)', 'Profit (loss) from operating activities'],
           'net_income': ['Net Income (Loss) Attributable to Parent', 'Net Income (Loss)', 'Profit (loss)'],
           'deferred_revenue': ['Contract with Customer, Liability', 'Deferred Revenue, Current', 'Contract with Customer, Liability, Current'],
           'cash': ['Cash and Cash Equivalents, at Carrying Value', 'Cash and cash equivalents'],
           'sbc': ['Share-based Payment Arrangement, Noncash Expense', 'Allocated Share-based Compensation Expense']}


def companies(out, only=None):
    panel, summary = [], []
    for name, cik, seg in COMPANIES:
        if only and name not in only:
            continue
        fin = financials(cik)
        k = kpis(name, cik)
        if not fin and not k:
            log(f'  ! {name}: nothing found')
            continue
        wb = Workbook()
        wb.remove(wb.active)
        if fin:
            years = _sheet(wb, 'Annual', fin['annual'], fin['units'], fin['tags'])
            _sheet(wb, 'Quarterly', fin['quarterly'], fin['units'], fin['tags'])
            for end in years:
                row = {'company': name, 'segment': seg, 'fiscal_year_end': end}
                for key, labels in SUMMARY.items():
                    row[key] = next((fin['annual'][lab][end] for lab in labels if lab in fin['annual'] and end in fin['annual'][lab]), None)
                summary.append(row)
        ws = wb.create_sheet('Reported KPIs')
        cols = ['filed', 'metric', 'period', 'value', 'document']
        ws.append(cols)
        for r in k:
            ws.append([r[c] for c in cols])
        d = wb.create_sheet('About')
        d.append([name])
        d.append(['Annual and Quarterly: every reported financial line item by period end, as far back as filed.'])
        d.append(['Reported KPIs: operating metrics from the quarterly earnings-release tables, as first reported.'])
        wb.properties.creator = PERSON
        buf = io.BytesIO()
        wb.save(buf)
        out.put(f'{FINANCE}/{seg}/{name}/{name} - financials and KPIs.xlsx', buf.getvalue(),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        for r in k:
            panel.append({**r, 'segment': seg})
        log(f'  {name}: {len(fin["annual"]) if fin else 0} annual lines, {len(k)} KPI values')
    out.table(f'{FINANCE}/All Companies', 'gaming_company_kpis', Table(
        'Reported operating KPIs, all companies', [('company', 'Company'), ('segment', 'Segment'), ('filed', 'Release filed'),
                                                   ('metric', 'Metric as labelled by the company'), ('period', 'Period column'),
                                                   ('value', 'Value as printed (units as in the release)'),
                                                   ('document', 'Filing document')], panel))
    for r in summary:
        rev = r.get('revenue')
        for key in ('gross_profit', 'rnd', 'sales_marketing', 'gna', 'operating_income', 'net_income', 'sbc'):
            r[f'{key}_pct_revenue'] = round(100 * r[key] / rev, 1) if rev and r.get(key) is not None else None
    cols = [('company', 'Company'), ('segment', 'Segment'), ('fiscal_year_end', 'Fiscal year end')] + [
        (k, f'{k.replace("_", " ").title()} (reported currency)') for k in SUMMARY] + [
        (f'{k}_pct_revenue', f'{k.replace("_", " ").title()} as % of revenue') for k in
        ('gross_profit', 'rnd', 'sales_marketing', 'gna', 'operating_income', 'net_income', 'sbc')]
    out.table(f'{FINANCE}/All Companies', 'gaming_company_annual_summary', Table(
        'Annual summary, all companies', cols, summary, 'Headline lines and cost structure per fiscal year.'))
    out.save_state('company_summary', summary)
