"""Output for the gaming data: tables written as CSV plus a formatted XLSX (with a data dictionary), uploaded to the
Shared Drive "Gaming Operations Data" under the operational category folders. A dry run writes to ./samples instead."""
import csv
import gzip
import io
import json
import os
import re
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PLAYERS = '1 Players & Engagement'
MONEY = '2 Monetisation'
MARKETING = '3 Marketing & Community'
PLATFORMS = '4 Platforms & Distribution'
DEV = '5 Development & Live Ops'
FINANCE = '6 Company Finance & KPIs'
BENCH = '7 Industry Benchmarks'
PRIVATE = 'Gaming Data - private'
XLSX_MAX_ROWS = 1_000_000
PERSON = 'Salvador Sekwa'


def log(*a):
    print(time.strftime('[%H:%M:%S]'), *a, flush=True)


_FIRSTS = None


def person_free(text):
    """Individual people's names (a known first name followed by an uncommon capitalised word) become Salvador Sekwa;
    studio and company names stay as they are."""
    global _FIRSTS
    if not text or not isinstance(text, str):
        return text
    if _FIRSTS is None:
        try:
            from fdocs import anon
            _FIRSTS = anon.first_names()
        except Exception:  # noqa: BLE001 - without the list nothing is replaced
            _FIRSTS = set()
    from wordfreq import zipf_frequency

    def one(m):
        first, last = m.group(1), m.group(2)
        if first.lower() in _FIRSTS and zipf_frequency(last.lower(), 'en') < 3.2 and not re.search(
                r'(?i)games?|studios?|entertainment|interactive|software|productions?|digital|media|labs?|works', last):
            return PERSON
        return m.group(0)
    return re.sub(r"\b([A-Z][a-z]+)\s+(?:[A-Z]\.\s+)?([A-Z][a-z]+(?:-[A-Z][a-z]+)?)\b", one, text)


class Table:
    """rows: list of dicts; columns: [(key, description)]."""

    def __init__(self, title, columns, rows, about=''):
        self.title, self.columns, self.rows, self.about = title, columns, rows, about

    def csv_bytes(self):
        buf = io.StringIO()
        w = csv.writer(buf)
        keys = [k for k, _ in self.columns]
        w.writerow(keys)
        for r in self.rows:
            w.writerow(['' if r.get(k) is None else r.get(k) for k in keys])
        return buf.getvalue().encode('utf-8')

    def xlsx_bytes(self):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Data')
        keys = [k for k, _ in self.columns]
        ws.freeze_panes = 'A2'
        for i, k in enumerate(keys, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(k) + 4))
        from openpyxl.cell import WriteOnlyCell
        head = []
        for k in keys:
            c = WriteOnlyCell(ws, value=k)
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F3864')
            c.alignment = Alignment(wrap_text=True, vertical='top')
            head.append(c)
        ws.append(head)
        for r in self.rows:
            ws.append([r.get(k) for k in keys])
        ws.auto_filter.ref = f'A1:{get_column_letter(len(keys))}{len(self.rows) + 1}'
        d = wb.create_sheet('Data dictionary')
        d.column_dimensions['A'].width = 28
        d.column_dimensions['B'].width = 100
        d.append([self.title])
        if self.about:
            d.append([self.about])
        d.append([f'{len(self.rows):,} rows; updated {time.strftime("%Y-%m-%d")}'])
        d.append([])
        d.append(['Column', 'Meaning'])
        for k, desc in self.columns:
            d.append([k, desc])
        wb.properties.creator = PERSON
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


class Out:
    def __init__(self, dry_run):
        self.dry = dry_run
        self.root = Path('samples')
        self.folders = {}
        if not dry_run:
            from models.store import Store
            self.store = Store(root_name=None, private_name=PRIVATE, app='gaming-data', state_name='unused.json.gz')
            self.files = self.store.drive.files()
            self.folders[''] = self.store.root

    # --- Drive paths --------------------------------------------------------
    def _folder(self, path):
        if path in self.folders:
            return self.folders[path]
        parent, _, name = path.rpartition('/')
        fid = self.store._folder(self._folder(parent), name)
        self.folders[path] = fid
        return fid

    def _find(self, path):
        from models.store import _retry
        folder, _, name = path.rpartition('/')
        parent = self.store.private if folder == PRIVATE else self._folder(folder)
        q = f"name = '{name.replace(chr(39), chr(92) + chr(39))}' and '{parent}' in parents and trashed = false"
        found = _retry(lambda: self.files.list(q=q, fields='files(id)', pageSize=2).execute())['files']
        return parent, name, (found[0]['id'] if found else None)

    def put(self, path, data, mime='application/octet-stream'):
        if self.dry:
            p = self.root / path
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_bytes(data)
            return
        from googleapiclient.http import MediaIoBaseUpload
        from models.store import _retry
        parent, name, fid = self._find(path)
        media = MediaIoBaseUpload(io.BytesIO(data), mimetype=mime, resumable=True, chunksize=32 * 2**20)
        req = (self.files.update(fileId=fid, media_body=media) if fid else
               self.files.create(body={'name': name, 'parents': [parent]}, media_body=media, fields='id'))
        resp = None
        while resp is None:
            _, resp = _retry(lambda: req.next_chunk())

    def get(self, path):
        if self.dry:
            p = self.root / path
            return p.read_bytes() if p.exists() else None
        from models.store import _retry
        _, _, fid = self._find(path)
        if not fid:
            return None
        return _retry(lambda: self.files.get_media(fileId=fid).execute())

    # --- tables and state -----------------------------------------------------
    def table(self, folder, base, table, xlsx=True):
        self.put(f'{folder}/{base}.csv', table.csv_bytes(), 'text/csv')
        if xlsx and len(table.rows) <= XLSX_MAX_ROWS:
            self.put(f'{folder}/{base}.xlsx', table.xlsx_bytes(),
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        log(f'  wrote {folder}/{base} ({len(table.rows):,} rows)')

    def read_csv(self, path):
        data = self.get(path)
        if not data:
            return []
        return list(csv.DictReader(io.StringIO(data.decode('utf-8'))))

    def state(self, name):
        data = self.get(f'{PRIVATE}/{name}.json.gz')
        return json.loads(gzip.decompress(data)) if data else None

    def save_state(self, name, obj):
        self.put(f'{PRIVATE}/{name}.json.gz', gzip.compress(json.dumps(obj, separators=(',', ':')).encode()), 'application/gzip')


def dry_run():
    return os.environ.get('DRY_RUN', 'false') == 'true'
