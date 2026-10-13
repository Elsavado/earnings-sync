"""Unlevered DCF valuation model.

Ten-year projection with revenue growth fading linearly from the company's historical median
to the terminal growth rate, margins and reinvestment at historical levels, WACC from CAPM
with an industry unlevered beta relevered at the company's market D/E, Gordon-growth
terminal value with the implied exit multiple, equity bridge, and a WACC x growth sensitivity
grid whose centre must equal the base case (checked).
"""
from . import xbrl
from .three import prepare
from openpyxl.styles import Font

from .xl import (NAVY, MULT, NOTE, NUM, PCT, PRICE, REDACTED, Inputs, Sheet, checks, col, cover, data_sheet,
                 new_workbook)

YEARS = 10
DATA_ROWS = [('revenue', 'Revenue'), ('cogs', 'Cost of revenue'), ('ebit', 'Operating income'), ('da', 'Depreciation & amortisation'),
             ('capex', 'Capital expenditure'), ('ar', 'Accounts receivable'), ('inventory', 'Inventory'),
             ('ap', 'Accounts payable'), ('cash', 'Cash and equivalents'), ('debt', 'Long-term debt'),
             ('debt_current', 'Current portion of debt'), ('interest', 'Interest expense'), ('equity', "Shareholders' equity")]


def build(facts, sector, market, sic=None):
    got = prepare(facts)
    if not got:
        return None
    ends, data = got
    data['debt_current'] = [v or 0.0 for v in data['debt_current']]
    revenue, ebit = data['revenue'], data['ebit']
    if sum(ebit) <= 0 or ebit[-1] <= 0:
        return None  # a cash-flow DCF on a loss-making history needs bespoke assumptions
    shares = xbrl.dei_latest(facts, 'EntityCommonStockSharesOutstanding', 'shares')
    float_ = xbrl.dei_latest(facts, 'EntityPublicFloat', 'USD')
    nh = len(ends)
    last = int(ends[-1][:4])
    periods = [f'FY{e[:4]}A' for e in ends] + [f'FY{last + k}E' for k in range(1, YEARS + 1)]
    n = len(periods)
    F0, FN = nh, n - 1
    wb = new_workbook()
    head = lambda t: f'{REDACTED} - {t}'  # noqa: E731

    ds = data_sheet(wb, head('Reported figures (Form 10-K)'), list(ends), DATA_ROWS, data)
    D = ds.ref

    bu, de_ind, industry, beta_note = market.beta(sic)
    mi = Inputs(wb, 'Market Inputs', head('Market and valuation inputs'))
    mi.section('Cost of capital')
    mi.add('rf', 'Risk-free rate', market.rf, PCT, market.notes['rf'])
    mi.add('erp', 'Equity risk premium', market.erp, PCT, market.notes['erp'])
    mi.add('bu', 'Unlevered beta (industry)', round(bu, 3), '0.00', beta_note)
    mi.add('tax', 'Marginal tax rate', 0.25, PCT, 'US federal 21% plus a typical state rate')
    mi.add('spread', 'Minimum credit spread over risk-free', 0.015, PCT, 'Floor for the pre-tax cost of debt')
    mi.section('Terminal value and conventions')
    mi.add('g', 'Terminal growth rate', 0.025, PCT, 'Long-run nominal growth; must stay below WACC')
    mi.add('mid', 'Mid-year discounting (1 = on, 0 = off)', 1, '0', 'Cash flows assumed to arrive mid-year')
    mi.section('Equity bridge')
    mi.add('shares', 'Shares outstanding (millions)', round(shares / 1e6, 3) if shares else 0, NUM,
           'From the latest 10-K cover page (dei:EntityCommonStockSharesOutstanding)')
    mi.add('mcap', 'Market value of equity (USD millions)', round(float_ / 1e6, 3) if float_ else 0, NUM,
           'Public float from the 10-K cover page (excludes insider holdings); used for capital-structure weights')
    mi.section('Sensitivity steps')
    mi.add('w_step', 'WACC step', 0.005, PCT)
    mi.add('g_step', 'Terminal growth step', 0.005, PCT)
    M = mi.ref

    asm = Sheet(wb, 'Assumptions', head('Operating drivers (derived from the historical years)'), ['Driver'])
    pj = Sheet(wb, 'Projections', head(f'Unlevered free cash flow projection ({YEARS} years)'), periods, nh)
    for key, label, bold, sec in [
        ('revenue', 'Revenue', True, 'Operating build'), ('growth', '  growth %', False, None), ('ebit', 'EBIT', True, None),
        ('margin', '  EBIT margin', False, None), ('da', 'Depreciation & amortisation', False, None), ('ebitda', 'EBITDA', True, None),
        ('taxes', 'Less: taxes on EBIT', False, 'Free cash flow'), ('nopat', 'NOPAT', True, None), ('da2', 'Plus: D&A', False, None),
        ('capex', 'Less: capital expenditure', False, None), ('nwc', 'Net working capital', False, None),
        ('dnwc', 'Less: increase in net working capital', False, None), ('ufcf', 'Unlevered free cash flow', True, None),
        ('year', 'Forecast year', False, 'Discounting'), ('period', 'Discount period', False, None),
        ('df', 'Discount factor', False, None), ('pv', 'Present value of UFCF', True, None)]:
        pj.row(key, label, bold, section=sec)
    P = pj.ref
    hi = nh - 1
    drivers = [
        ('g0', 'Starting revenue growth (median of historical years)', f'IFERROR(MEDIAN({pj.span("growth", 1, hi)}),0)'),
        ('margin', 'EBIT margin (historical average)', f'IFERROR(SUM({pj.span("ebit", 0, hi)})/SUM({pj.span("revenue", 0, hi)}),0)'),
        ('da_pct', 'D&A % revenue', f'IFERROR(SUM({pj.span("da", 0, hi)})/SUM({pj.span("revenue", 0, hi)}),0)'),
        ('capex_pct', 'Capital expenditure % revenue', f'IFERROR(-SUM({pj.span("capex", 0, hi)})/SUM({pj.span("revenue", 0, hi)}),0)'),
        ('nwc_pct', 'Net working capital % revenue', f'IFERROR(SUM({pj.span("nwc", 0, hi)})/SUM({pj.span("revenue", 0, hi)}),0)'),
    ]
    for key, label, f in drivers:
        asm.row(key, label)
        asm.set(key, 0, f'={f}', PCT)
    A = lambda k: asm.ref(k, 0, absolute=True)  # noqa: E731
    asm.ws.cell(row=asm.next_row + 1, column=1,
                value='Revenue growth fades in a straight line from the starting rate to the terminal growth rate by the final '
                      'forecast year; margins and reinvestment rates stay at their historical levels.').font = NOTE

    for j in range(n):
        p, h = j - 1, j < nh
        pj.set('revenue', j, f'={D("revenue", j)}' if h else f'={P("revenue", p)}*(1+{P("growth", j)})', bold=True)
        if j:
            pj.set('growth', j, f'=IFERROR({P("revenue", j)}/{P("revenue", p)}-1,0)' if h else
                   f'={A("g0")}+({M("g")}-{A("g0")})*{P("year", j)}/{YEARS}', PCT)
        pj.set('ebit', j, f'={D("ebit", j)}' if h else f'={P("revenue", j)}*{A("margin")}', bold=True)
        pj.set('margin', j, f'=IFERROR({P("ebit", j)}/{P("revenue", j)},0)', PCT)
        pj.set('da', j, f'={D("da", j)}' if h else f'={P("revenue", j)}*{A("da_pct")}')
        pj.set('ebitda', j, f'={P("ebit", j)}+{P("da", j)}', bold=True)
        pj.set('taxes', j, f'=-MAX(0,{P("ebit", j)})*{M("tax")}')
        pj.set('nopat', j, f'={P("ebit", j)}+{P("taxes", j)}', bold=True)
        pj.set('da2', j, f'={P("da", j)}')
        pj.set('capex', j, f'=-{D("capex", j)}' if h else f'=-{P("revenue", j)}*{A("capex_pct")}')
        pj.set('nwc', j, f'={D("ar", j)}+{D("inventory", j)}-{D("ap", j)}' if h else f'={P("revenue", j)}*{A("nwc_pct")}')
        if j:
            pj.set('dnwc', j, f'=-({P("nwc", j)}-{P("nwc", p)})')
            pj.set('ufcf', j, f'={P("nopat", j)}+{P("da2", j)}+{P("capex", j)}+{P("dnwc", j)}', bold=True)
        if not h:
            pj.set('year', j, f'=N({P("year", p)})+1', '0')
            pj.set('period', j, f'={P("year", j)}-{M("mid")}/2', '0.0')
            pj.set('pv', j, f'={P("ufcf", j)}*{P("df", j)}', bold=True)

    # WACC
    wc = Inputs(wb, 'WACC', head('Weighted average cost of capital'))
    wc.section('Capital structure (market values)')
    wc.add('debt', 'Total debt (latest year)', f'={D("debt", hi)}+{D("debt_current", hi)}', NUM, formula=True)
    wc.add('equity', 'Equity value used for weights',
           f'=IF({M("mcap")}>0,{M("mcap")},{D("equity", hi)})', NUM, 'Public float if reported, otherwise book equity', formula=True)
    wc.add('de', 'Debt / equity', f'=IFERROR({wc.ref("debt")}/{wc.ref("equity")},0)', PCT, formula=True)
    wc.add('wd', 'Weight of debt', f'=IFERROR({wc.ref("debt")}/({wc.ref("debt")}+{wc.ref("equity")}),0)', PCT, formula=True)
    wc.add('we', 'Weight of equity', f'=1-{wc.ref("wd")}', PCT, formula=True)
    wc.section('Cost of equity (CAPM)')
    wc.add('bl', 'Relevered beta', f'={M("bu")}*(1+(1-{M("tax")})*{wc.ref("de")})', '0.00',
           'Unlevered industry beta relevered at the company D/E (Hamada)', formula=True)
    wc.add('ke', 'Cost of equity', f'={M("rf")}+{wc.ref("bl")}*{M("erp")}', PCT, formula=True)
    wc.section('Cost of debt')
    wc.add('kd_hist', 'Historical interest rate on debt',
           f'=IFERROR({D("interest", hi)}/({D("debt", hi)}+{D("debt_current", hi)}),0)', PCT, formula=True)
    wc.add('kd', 'Pre-tax cost of debt', f'=MAX({wc.ref("kd_hist")},{M("rf")}+{M("spread")})', PCT,
           'Higher of the historical rate and risk-free plus minimum spread', formula=True)
    wc.add('kd_at', 'After-tax cost of debt', f'={wc.ref("kd")}*(1-{M("tax")})', PCT, formula=True)
    wc.section('Result')
    wc.add('wacc', 'WACC', f'={wc.ref("we")}*{wc.ref("ke")}+{wc.ref("wd")}*{wc.ref("kd_at")}', PCT, formula=True)
    W = wc.ref('wacc')
    for j in range(F0, n):
        pj.set('df', j, f'=1/(1+{W})^{P("period", j)}', '0.0000')

    # Valuation
    va = Inputs(wb, 'Valuation', head('Enterprise and equity value'))
    ufcf_n, ebitda_n = P('ufcf', FN), P('ebitda', FN)
    va.section('Enterprise value')
    va.add('sum_pv', 'Sum of PV of forecast cash flows', f'=SUM({pj.span("pv", F0, FN)})', NUM, formula=True)
    va.add('tv', 'Terminal value (Gordon growth, end of final year)', f'={ufcf_n}*(1+{M("g")})/({W}-{M("g")})', NUM, formula=True)
    va.add('pv_tv', 'PV of terminal value', f'={va.ref("tv")}/(1+{W})^{P("year", FN)}', NUM, formula=True)
    va.add('ev', 'Enterprise value', f'={va.ref("sum_pv")}+{va.ref("pv_tv")}', NUM, formula=True)
    va.add('tv_share', 'Terminal value % enterprise value', f'=IFERROR({va.ref("pv_tv")}/{va.ref("ev")},0)', PCT, formula=True)
    va.add('exit_mult', 'Implied exit EV / EBITDA', f'=IFERROR({va.ref("tv")}/{ebitda_n},0)', MULT, formula=True)
    va.section('Equity bridge')
    va.add('debt', 'Less: total debt', f'=-{wc.ref("debt")}', NUM, formula=True)
    va.add('cash', 'Plus: cash and equivalents', f'={D("cash", hi)}', NUM, formula=True)
    va.add('eqv', 'Equity value', f'={va.ref("ev")}+{va.ref("debt")}+{va.ref("cash")}', NUM, formula=True)
    va.add('ps', 'Equity value per share (USD)', f'=IFERROR({va.ref("eqv")}/{M("shares")},0)', PRICE, formula=True)
    va.add('vs_mkt', 'Equity value vs market value of equity', f'=IFERROR({va.ref("eqv")}/{M("mcap")}-1,0)', PCT,
           'Positive = model value above the public float value (float excludes insiders)', formula=True)
    va.section('Implied multiples')
    va.add('ev_rev', 'EV / revenue (latest year)', f'=IFERROR({va.ref("ev")}/{D("revenue", hi)},0)', MULT, formula=True)
    va.add('ev_ebitda', 'EV / EBITDA (latest year)', f'=IFERROR({va.ref("ev")}/{P("ebitda", hi)},0)', MULT, formula=True)

    # Sensitivity: equity value for WACC x terminal growth, each cell computed in closed form.
    se = wb.create_sheet('Sensitivity')
    se.sheet_view.showGridLines = False
    se['A1'] = head('Equity value sensitivity (USD millions)')
    se['A1'].font = Font(color=NAVY, bold=True, size=13)
    se['A2'] = 'Rows: WACC. Columns: terminal growth. Every cell re-discounts the projected cash flows.'
    se['A2'].font = NOTE
    se.column_dimensions['A'].width = 14
    steps = [-2, -1, 0, 1, 2]
    for k, s in enumerate(steps):
        c = se.cell(row=4, column=2 + k, value=f'={M("g")}+({s})*{M("g_step")}')
        c.number_format = PCT
        c = se.cell(row=5 + k, column=1, value=f'={W}+({s})*{M("w_step")}')
        c.number_format = PCT
        se.column_dimensions[col(k)].width = 14
    ufcf_rng, per_rng = pj.span('ufcf', F0, FN), pj.span('period', F0, FN)
    bridge = f'{va.ref("debt")}+{va.ref("cash")}'
    for r in range(5):
        for k in range(5):
            w, g = f'$A{5 + r}', f'{col(k)}$4'
            f = (f'=SUMPRODUCT({ufcf_rng}/(1+{w})^{per_rng})+{ufcf_n}*(1+{g})/({w}-{g})/(1+{w})^{P("year", FN)}+{bridge}')
            se.cell(row=5 + r, column=2 + k, value=f).number_format = NUM
    centre = 'Sensitivity!$D$7'

    checks(wb, [('WACC exceeds terminal growth (0 = yes)', f'IF({W}>{M("g")},0,1)'),
                ('Sensitivity centre equals base-case equity value', f'{centre}-{va.ref("eqv")}'),
                ('PV of cash flows equals SUMPRODUCT re-calculation',
                 f'{va.ref("sum_pv")}-SUMPRODUCT({ufcf_rng}/(1+{W})^{per_rng})')])
    cover(wb, 'Discounted cash flow (DCF) valuation model', f'{sector} (beta peer group: {industry})',
          f'Unlevered DCF built from the company\'s reported 10-K figures: {YEARS}-year projection with growth fading to the '
          'terminal rate, WACC from CAPM using a relevered industry beta, Gordon-growth terminal value with the implied '
          'exit multiple, an equity bridge and a WACC x growth sensitivity table. Market inputs are dated on the Market Inputs sheet.',
          [('Historical Data', 'Reported annual figures (USD millions) - hard-coded source data'),
           ('Market Inputs', 'Risk-free rate, ERP, beta, tax, terminal growth and share data (inputs)'),
           ('Assumptions', 'Operating drivers calculated from history'),
           ('Projections', 'Revenue to unlevered free cash flow and discounting'),
           ('WACC', 'Capital structure, CAPM cost of equity, cost of debt, WACC'),
           ('Valuation', 'Enterprise value, equity bridge, per-share value, implied multiples'),
           ('Sensitivity', 'Equity value across WACC and terminal growth'),
           ('Checks', 'Integrity checks')])
    wb.properties.title = f'{REDACTED} - DCF valuation model'
    wb.properties.creator = REDACTED
    return wb
