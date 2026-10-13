"""Leveraged buyout model.

Acquisition at the latest fiscal year end: sources & uses, a SOFR-based term loan with
mandatory amortisation and an excess-cash sweep, fixed-rate senior notes, five-year operating
projection from the company's historical ratios, credit statistics, sponsor returns (MOIC and
IRR, the IRR computed two ways and checked), and an IRR grid over entry and exit multiples.
Interest runs on opening balances, so the model has no circular references.
"""
from openpyxl.styles import Font

from . import xbrl
from .three import prepare
from .xl import (MULT, NAVY, NOTE, NUM, PCT, REDACTED, Inputs, Sheet, checks, col, cover, data_sheet, new_workbook)

YEARS = 5
DATA_ROWS = [('revenue', 'Revenue'), ('ebit', 'Operating income'), ('da', 'Depreciation & amortisation'),
             ('capex', 'Capital expenditure'), ('ar', 'Accounts receivable'), ('inventory', 'Inventory'),
             ('ap', 'Accounts payable'), ('cash', 'Cash and equivalents'), ('debt', 'Long-term debt'),
             ('debt_current', 'Current portion of debt'), ('tax', 'Income tax expense'), ('pretax', 'Pre-tax income')]


def build(facts, sector, market, sic=None):
    try:
        if 6000 <= int(sic or 0) <= 6799:
            return None  # leveraged buyouts of banks, insurers and REITs are not modelled this way
    except ValueError:
        pass
    got = prepare(facts)
    if not got:
        return None
    ends, data = got
    data['debt_current'] = [v or 0.0 for v in data['debt_current']]
    ebitda = [e + d for e, d in zip(data['ebit'], data['da'])]
    if ebitda[-1] <= 0 or sum(ebitda) <= 0 or ebitda[-1] / data['revenue'][-1] < 0.05:
        return None
    # Entry multiple: today's EV / EBITDA from the cover-page float, kept within a range a
    # sponsor would underwrite. Chosen here so the workbook carries it as a plain input.
    float_ = xbrl.dei_latest(facts, 'EntityPublicFloat', 'USD')
    net_debt = data['debt'][-1] + data['debt_current'][-1] - data['cash'][-1]
    entry = 10.0
    note = 'Assumed entry multiple (no market value reported)'
    if float_:
        cur = (float_ / 1e6 + net_debt) / ebitda[-1]
        entry = round(min(16.0, max(6.0, cur)), 1)
        note = f'Current EV/EBITDA from the 10-K public float is {cur:.1f}x; entry set within a 6x-16x underwriting range'
    senior_x = round(min(4.0, entry * 0.35), 2)
    notes_x = round(min(2.0, entry * 0.15), 2)

    nh = len(ends)
    hi = nh - 1
    last = int(ends[-1][:4])
    periods = [f'FY{e[:4]}A' for e in ends] + [f'FY{last + k}E' for k in range(1, YEARS + 1)]
    n = len(periods)
    F0, FN = nh, n - 1
    wb = new_workbook()
    head = lambda t: f'{REDACTED} - {t}'  # noqa: E731

    ds = data_sheet(wb, head('Reported figures (Form 10-K)'), list(ends), DATA_ROWS, data)
    D = ds.ref

    ti = Inputs(wb, 'Transaction Inputs', head('Transaction, financing and exit inputs'))
    ti.section('Purchase price')
    ti.add('entry', 'Entry EV / EBITDA', entry, MULT, note)
    ti.add('tx_fee', 'Transaction fees % enterprise value', 0.02, PCT, 'Advisory, legal and other deal costs')
    ti.add('min_cash', 'Minimum operating cash % revenue', 0.03, PCT, 'Cash left on balance sheet at close and maintained')
    ti.section('Financing')
    ti.add('sofr', 'SOFR', market.sofr, PCT, market.notes['sofr'])
    ti.add('tl_x', 'Term loan B (x LTM EBITDA)', senior_x, MULT, 'Senior secured, floating')
    ti.add('tl_spread', 'Term loan spread over SOFR', 0.035, PCT)
    ti.add('tl_amort', 'Term loan mandatory amortisation % original principal p.a.', 0.01, PCT)
    ti.add('sweep', 'Excess cash flow sweep %', 0.75, PCT, 'Share of surplus cash applied to prepay the term loan')
    ti.add('sn_x', 'Senior notes (x LTM EBITDA)', notes_x, MULT, 'Unsecured, fixed coupon, bullet at maturity')
    ti.add('sn_rate', 'Senior notes coupon', round(market.rf + 0.035, 4), PCT, 'Risk-free rate plus 350 bp at pricing')
    ti.add('fin_fee', 'Financing fees % of debt raised', 0.025, PCT)
    ti.section('Operations and exit')
    ti.add('tax', 'Cash tax rate', 0.25, PCT)
    ti.add('exit', 'Exit EV / EBITDA', entry, MULT, 'Set equal to entry (no multiple expansion)')
    ti.add('hold', 'Holding period (years)', YEARS, '0', 'Exit at the end of the final forecast year')
    ti.add('step', 'Sensitivity step (x EBITDA)', 1.0, MULT)
    T = ti.ref

    asm = Sheet(wb, 'Assumptions', head('Operating drivers (derived from the historical years)'), ['Driver'])
    op = Sheet(wb, 'Operating Model', head('Operating projection and free cash flow'), periods, nh)
    for key, label, bold, sec in [
        ('revenue', 'Revenue', True, 'Income statement'), ('growth', '  growth %', False, None), ('ebitda', 'EBITDA', True, None),
        ('margin', '  EBITDA margin', False, None), ('da', 'Depreciation & amortisation', False, None), ('ebit', 'EBIT', True, None),
        ('interest', 'Interest expense', False, None), ('pretax', 'Pre-tax income', True, None), ('taxes', 'Cash taxes', False, None),
        ('ni', 'Net income', True, None), ('capex', 'Capital expenditure', False, 'Free cash flow'), ('nwc', 'Net working capital', False, None),
        ('dnwc', 'Increase in net working capital', False, None), ('fcf', 'Free cash flow available for debt service', True, None)]:
        op.row(key, label, bold, section=sec)
    O = op.ref
    for key, label, f in [
        ('growth', 'Revenue growth (median of historical years)', f'IFERROR(MEDIAN({op.span("growth", 1, hi)}),0)'),
        ('margin', 'EBITDA margin (historical average)', f'IFERROR(SUM({op.span("ebitda", 0, hi)})/SUM({op.span("revenue", 0, hi)}),0)'),
        ('da_pct', 'D&A % revenue', f'IFERROR(SUM({op.span("da", 0, hi)})/SUM({op.span("revenue", 0, hi)}),0)'),
        ('capex_pct', 'Capital expenditure % revenue', f'IFERROR(-SUM({op.span("capex", 0, hi)})/SUM({op.span("revenue", 0, hi)}),0)'),
        ('nwc_pct', 'Net working capital % revenue', f'IFERROR(SUM({op.span("nwc", 0, hi)})/SUM({op.span("revenue", 0, hi)}),0)')]:
        asm.row(key, label)
        asm.set(key, 0, f'={f}', PCT)
    A = lambda k: asm.ref(k, 0, absolute=True)  # noqa: E731

    # Sources & uses
    su = Inputs(wb, 'Sources & Uses', head('Sources and uses of funds at close'))
    ltm = O('ebitda', hi)
    su.section('Uses')
    su.add('ev', 'Enterprise value (entry multiple x LTM EBITDA)', f'={T("entry")}*{ltm}', NUM, formula=True)
    su.add('tx_fees', 'Transaction fees', f'={su.ref("ev")}*{T("tx_fee")}', NUM, formula=True)
    su.add('fin_fees', 'Financing fees', f'=({T("tl_x")}+{T("sn_x")})*{ltm}*{T("fin_fee")}', NUM, formula=True)
    su.add('cash_bs', 'Cash to balance sheet', f'={D("revenue", hi)}*{T("min_cash")}', NUM, formula=True)
    su.add('uses', 'Total uses', f'=SUM(B{su.rows["ev"]}:B{su.rows["cash_bs"]})', NUM, formula=True)
    su.section('Sources')
    su.add('tl', 'Term loan B', f'={T("tl_x")}*{ltm}', NUM, formula=True)
    su.add('sn', 'Senior notes', f'={T("sn_x")}*{ltm}', NUM, formula=True)
    su.add('equity', 'Sponsor equity', f'={su.ref("uses")}-{su.ref("tl")}-{su.ref("sn")}', NUM, formula=True)
    su.add('sources', 'Total sources', f'=SUM(B{su.rows["tl"]}:B{su.rows["equity"]})', NUM, formula=True)
    su.add('eq_pct', 'Equity % of total sources', f'=IFERROR({su.ref("equity")}/{su.ref("sources")},0)', PCT, formula=True)
    su.add('lev', 'Total leverage at close (x LTM EBITDA)', f'=IFERROR(({su.ref("tl")}+{su.ref("sn")})/{ltm},0)', MULT, formula=True)

    # Debt schedule
    db = Sheet(wb, 'Debt Schedule', head('Debt schedule and cash'), periods, nh)
    for key, label, bold, sec in [
        ('cash_open', 'Opening cash', False, 'Cash'), ('fcf', 'Free cash flow', False, None), ('mand', 'Less: mandatory amortisation', False, None),
        ('avail', 'Cash available for sweep (above minimum)', False, None), ('sweep', 'Less: optional prepayment (sweep)', False, None),
        ('cash', 'Closing cash', True, None),
        ('tl_open', 'Opening balance', False, 'Term loan B'), ('tl_mand', 'Mandatory amortisation', False, None),
        ('tl_sweep', 'Optional prepayment', False, None), ('tl', 'Closing balance', True, None), ('tl_rate', 'Interest rate', False, None),
        ('tl_int', 'Interest expense', False, None),
        ('sn_open', 'Opening balance', False, 'Senior notes'), ('sn', 'Closing balance', True, None), ('sn_int', 'Interest expense', False, None),
        ('total', 'Total debt', True, 'Credit statistics'), ('net', 'Net debt', False, None), ('lev', 'Total debt / EBITDA', False, None),
        ('net_lev', 'Net debt / EBITDA', False, None), ('cov', 'EBITDA / interest', False, None)]:
        db.row(key, label, bold, section=sec)
    B = db.ref

    for j in range(n):
        p, h = j - 1, j < nh
        op.set('revenue', j, f'={D("revenue", j)}' if h else f'={O("revenue", p)}*(1+{A("growth")})', bold=True)
        if j:
            op.set('growth', j, f'=IFERROR({O("revenue", j)}/{O("revenue", p)}-1,0)', PCT)
        op.set('ebitda', j, f'={D("ebit", j)}+{D("da", j)}' if h else f'={O("revenue", j)}*{A("margin")}', bold=True)
        op.set('margin', j, f'=IFERROR({O("ebitda", j)}/{O("revenue", j)},0)', PCT)
        op.set('da', j, f'={D("da", j)}' if h else f'={O("revenue", j)}*{A("da_pct")}')
        op.set('ebit', j, f'={O("ebitda", j)}-{O("da", j)}', bold=True)
        op.set('capex', j, f'=-{D("capex", j)}' if h else f'=-{O("revenue", j)}*{A("capex_pct")}')
        op.set('nwc', j, f'={D("ar", j)}+{D("inventory", j)}-{D("ap", j)}' if h else f'={O("revenue", j)}*{A("nwc_pct")}')
        if h:
            continue
        op.set('interest', j, f'={B("tl_int", j)}+{B("sn_int", j)}')
        op.set('pretax', j, f'={O("ebit", j)}-{O("interest", j)}', bold=True)
        op.set('taxes', j, f'=-MAX(0,{O("pretax", j)})*{T("tax")}')
        op.set('ni', j, f'={O("pretax", j)}+{O("taxes", j)}', bold=True)
        op.set('dnwc', j, f'=-({O("nwc", j)}-{O("nwc", p)})')
        op.set('fcf', j, f'={O("ni", j)}+{O("da", j)}+{O("capex", j)}+{O("dnwc", j)}', bold=True)

        first = j == F0
        db.set('cash_open', j, f'={su.ref("cash_bs")}' if first else f'={B("cash", p)}')
        db.set('fcf', j, f'={O("fcf", j)}')
        db.set('mand', j, f'=-{B("tl_mand", j)}')
        db.set('avail', j, f'=MAX(0,{B("cash_open", j)}+{B("fcf", j)}+{B("mand", j)}-{O("revenue", j)}*{T("min_cash")})')
        db.set('sweep', j, f'=-{B("tl_sweep", j)}')
        db.set('cash', j, f'={B("cash_open", j)}+{B("fcf", j)}+{B("mand", j)}+{B("sweep", j)}', bold=True)
        db.set('tl_open', j, f'={su.ref("tl")}' if first else f'={B("tl", p)}')
        db.set('tl_mand', j, f'=MIN({B("tl_open", j)},{su.ref("tl")}*{T("tl_amort")})')
        db.set('tl_sweep', j, f'=MIN({B("tl_open", j)}-{B("tl_mand", j)},{B("avail", j)}*{T("sweep")})')
        db.set('tl', j, f'={B("tl_open", j)}-{B("tl_mand", j)}-{B("tl_sweep", j)}', bold=True)
        db.set('tl_rate', j, f'={T("sofr")}+{T("tl_spread")}', PCT)
        db.set('tl_int', j, f'={B("tl_open", j)}*{B("tl_rate", j)}')
        db.set('sn_open', j, f'={su.ref("sn")}' if first else f'={B("sn", p)}')
        db.set('sn', j, f'={B("sn_open", j)}', bold=True)
        db.set('sn_int', j, f'={B("sn_open", j)}*{T("sn_rate")}')
        db.set('total', j, f'={B("tl", j)}+{B("sn", j)}', bold=True)
        db.set('net', j, f'={B("total", j)}-{B("cash", j)}')
        db.set('lev', j, f'=IFERROR({B("total", j)}/{O("ebitda", j)},0)', MULT)
        db.set('net_lev', j, f'=IFERROR({B("net", j)}/{O("ebitda", j)},0)', MULT)
        db.set('cov', j, f'=IFERROR({O("ebitda", j)}/{O("interest", j)},0)', MULT)

    # Returns
    rt = Inputs(wb, 'Returns', head('Sponsor returns'))
    rt.section('Exit')
    rt.add('ebitda', 'Exit-year EBITDA', f'={O("ebitda", FN)}', NUM, formula=True)
    rt.add('ev', 'Exit enterprise value', f'={rt.ref("ebitda")}*{T("exit")}', NUM, formula=True)
    rt.add('net', 'Less: net debt at exit', f'=-{B("net", FN)}', NUM, formula=True)
    rt.add('eq', 'Exit equity value', f'={rt.ref("ev")}+{rt.ref("net")}', NUM, formula=True)
    rt.section('Returns')
    rt.add('moic', 'Multiple of invested capital (MOIC)', f'=IFERROR({rt.ref("eq")}/{su.ref("equity")},0)', MULT, formula=True)
    rt.add('irr', 'IRR', f'=IFERROR({rt.ref("moic")}^(1/{T("hold")})-1,0)', PCT, formula=True)
    # Equity cash flows: investment at close, proceeds at exit.
    r0 = rt.next_row + 1
    rt.ws.cell(row=r0, column=1, value='Sponsor equity cash flows').font = Font(bold=True, color=NAVY)
    for k in range(YEARS + 1):
        rt.ws.cell(row=r0 + 1, column=2 + k, value=f'={op.ws.cell(row=3, column=2 + hi + k).coordinate}'
                   .replace('=', f"='Operating Model'!")).font = Font(bold=True)
        if k == 0:
            v = f'=-{su.ref("equity")}'
        elif k == YEARS:
            v = f'={rt.ref("eq")}'
        else:
            v = f'={su.ref("equity")}-{su.ref("equity")}'
        rt.ws.cell(row=r0 + 2, column=2 + k, value=v).number_format = NUM
    flows = f"'Returns'!B{r0 + 2}:{col(YEARS)}{r0 + 2}"
    rt.next_row = r0 + 4
    rt.add('irr_xl', 'IRR (Excel IRR function on the cash flows)', f'=IFERROR(IRR({flows}),0)', PCT, formula=True)

    # Sensitivity: IRR over entry and exit multiples, closed form (debt is sized on EBITDA, so
    # the debt paydown does not depend on the purchase multiple).
    se = wb.create_sheet('Sensitivity')
    se.sheet_view.showGridLines = False
    se['A1'] = head('IRR sensitivity: entry multiple (rows) x exit multiple (columns)')
    se['A1'].font = Font(bold=True, size=13, color=NAVY)
    se['A2'] = 'Each cell rebuilds the sponsor equity cheque and exit proceeds for that pair of multiples.'
    se['A2'].font = NOTE
    se.column_dimensions['A'].width = 14
    steps = [-2, -1, 0, 1, 2]
    for k, s in enumerate(steps):
        se.cell(row=4, column=2 + k, value=f'={T("exit")}+({s})*{T("step")}').number_format = MULT
        se.cell(row=5 + k, column=1, value=f'={T("entry")}+({s})*{T("step")}').number_format = MULT
        se.column_dimensions[col(k)].width = 12
    other_uses = f'({su.ref("fin_fees")}+{su.ref("cash_bs")}-{su.ref("tl")}-{su.ref("sn")})'
    for r in range(5):
        for k in range(5):
            e, x = f'$A{5 + r}', f'{col(k)}$4'
            eq_in = f'({e}*{ltm}*(1+{T("tx_fee")})+{other_uses})'
            f = f'=IFERROR((({x}*{rt.ref("ebitda")}+{rt.ref("net")})/{eq_in})^(1/{T("hold")})-1,0)'
            se.cell(row=5 + r, column=2 + k, value=f).number_format = PCT

    checks(wb, [('Sources equal uses', f'{su.ref("sources")}-{su.ref("uses")}'),
                ('IRR formula equals Excel IRR', f'({rt.ref("irr")}-{rt.ref("irr_xl")})*100'),
                ('Sensitivity centre equals base IRR', f'(Sensitivity!$D$7-{rt.ref("irr")})*100'),
                ('Closing cash never below zero', f'-MIN(0,MIN({db.span("cash", F0, FN)}))'),
                ('Term loan never negative', f'-MIN(0,MIN({db.span("tl", F0, FN)}))')])
    cover(wb, 'Leveraged buyout (LBO) model', sector,
          'Sponsor acquisition of the company at its latest fiscal year end, financed with a SOFR term loan and fixed-rate '
          'senior notes. The operating projection comes from the company\'s reported 10-K history; the term loan amortises '
          'and is prepaid from excess cash; returns are shown as MOIC and IRR with an entry/exit multiple sensitivity.',
          [('Historical Data', 'Reported annual figures (USD millions) - hard-coded source data'),
           ('Transaction Inputs', 'Purchase price, financing terms, fees and exit (inputs)'),
           ('Assumptions', 'Operating drivers calculated from history'),
           ('Sources & Uses', 'Funding of the purchase at close'),
           ('Operating Model', 'Revenue to free cash flow available for debt service'),
           ('Debt Schedule', 'Term loan, senior notes, cash sweep and credit statistics'),
           ('Returns', 'Exit value, MOIC and IRR'),
           ('Sensitivity', 'IRR across entry and exit multiples'),
           ('Checks', 'Integrity checks')])
    wb.properties.title = f'{REDACTED} - LBO model'
    wb.properties.creator = REDACTED
    return wb
