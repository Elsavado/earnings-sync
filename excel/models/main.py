"""One collection run: refresh sources, process scraped spreadsheets, build SEC company models,
verify every file, number it and upload it to Drive in folders of twenty.

Environment: GOOGLE_CLIENT_ID / GOOGLE_CLIENT_SECRET / GOOGLE_REFRESH_TOKEN (Drive),
SEC_USER_AGENT (SEC fair-access contact), GITHUB_TOKEN (GitHub API), RUN_MINUTES (budget),
DRY_RUN=true (build and verify, no uploads), LOCAL_COPY_DIR (copies of models 1-100).
"""
import hashlib
import json
import multiprocessing
import os
import re
import shutil
import sys
import tempfile
import time
import traceback
from pathlib import Path

import requests
from openpyxl.workbook import Workbook  # noqa: F401  (import check)

from . import bank, bankent, dcf, enterprise, fpa, govent, irent, lbo, links, lo, sources, three, unique
from .anonymize import Anonymizer
from .market import Market
from .process import process
from .store import Store, folder_name
from .xl import INPUT_SHEETS
from .xlsxml import Package

START = time.time()
BUDGET = float(os.environ.get('RUN_MINUTES', '50')) * 60
DRY = os.environ.get('DRY_RUN', '').lower() == 'true'
LOCAL = os.environ.get('LOCAL_COPY_DIR', 'first100')
SEC_UA = os.environ.get('SEC_USER_AGENT') or 'ExcelModels research contact@example.com'
BANK_SIC = set(range(6020, 6037)) | {6111, 6141, 6153, 6159, 6162, 6172, 6199}
# Enterprise-wide models (enterprise.py): one model type per job, each with its own folder and numbering.
ENT_KIND = os.environ.get('ENT_KIND', '')
ENT_ROOT = os.environ.get('ENT_ROOT', '')  # '' when the Shared Drive itself is 'Enterprise Wide Models'
NAME_KEY = (os.environ.get('ENT_NAME_KEY') or os.urandom(32).hex()).encode()


SHORT = {'integrated': 'Integrated Plan', 'valuation': 'DCF Valuation', 'buyout': 'LBO', 'budget': 'Budget & FP&A',
         'credit': 'Credit Analysis', 'strategic': '10Y Strategic Plan', 'bank_integrated': 'Bank Plan',
         'bank_capital': 'Bank Capital & Stress', 'bank_alm': 'Bank ALM', 'bank_valuation': 'Bank Valuation',
         'bank_strategic': 'Bank 10Y Plan', 'gov_population': 'Population Plan', 'gov_fiscal': 'Fiscal DSA',
         'gov_economy': 'Macro Framework', 'gov_project': 'Project Appraisal', 'ir_enhanced': 'IR Operating Model'}


def short_name(text, limit=24):
    """'Services-Prepackaged Software' -> 'Prepackaged Software'; cut at a word boundary."""
    t = re.sub(r'^(?:Services|Retail|Wholesale)-', '', text or '').replace('&amp;', '&')
    t = re.sub(r'\s*\(.*?\)', '', t).replace(',', ' ')
    t = re.sub(r'\s+', ' ', re.sub(r'\b(?:And|Of|The|For)\b', lambda m: m.group(0).lower(), t.title())).strip(' -')
    if len(t) > limit:
        t = t[:limit].rsplit(' ', 1)[0]
    return t.strip(' -&,') or 'General'


LIST_V = 2  # 2: listed companies plus every other XBRL filer with annual figures (user, 2026-10-24: "add more companies")
FRAME_TAGS = ('Revenues', 'RevenueFromContractWithCustomerExcludingAssessedTax', 'SalesRevenueNet', 'Assets', 'NetIncomeLoss',
              'InterestAndDividendIncomeOperating')


def sec_filers(s):
    """CIKs to model: companies with a ticker first, then every other company that reported annual revenue, assets,
    net income or bank interest income in XBRL in any year since 2010 (unlisted, delisted and debt-only filers)."""
    data = s.get('https://www.sec.gov/files/company_tickers.json', timeout=60).json()
    ciks, seen = [], set()
    for v in data.values():
        if v['cik_str'] not in seen:
            seen.add(v['cik_str'])
            ciks.append(v['cik_str'])
    more = set()
    for tag in FRAME_TAGS:
        for year in range(2010, time.gmtime().tm_year):
            try:
                r = s.get(f'https://data.sec.gov/api/xbrl/frames/us-gaap/{tag}/USD/CY{year}.json', timeout=120)
                time.sleep(0.12)
                if r.ok:
                    more.update(int(d['cik']) for d in r.json().get('data', []) if d.get('cik'))
            except (requests.RequestException, ValueError):
                continue
    return ciks + sorted(more - seen)


def finish_unique(wb, key, extra_inputs=()):
    """The model's own look (unique.py), then the reference check; returns the sheets allowed to hold figures."""
    mapping = unique.apply(wb, wb.profile if getattr(wb, 'profile', None) else unique.Profile(key))
    enterprise.references_ok(wb)
    return {mapping.get(n, n) for n in set(INPUT_SHEETS) | set(extra_inputs)}


def model_code(key):
    """Unique code for a file name that does not reveal the company (keyed hash; the key is a repository secret)."""
    import hmac
    return 'EW-' + hmac.new(NAME_KEY, key.encode(), hashlib.sha256).hexdigest()[:8].upper()
TYPE_LABEL = {'three': 'Three-statement model', 'reit': 'Industry-specific model - REIT', 'bank': 'Industry-specific model - Bank',
              'dcf': 'DCF valuation model', 'lbo': 'LBO model', 'fpa': 'Operating & FP&A model'}


FILE_SECONDS = 180  # hard limit for processing one downloaded spreadsheet
ANON = None  # set before the worker pool forks, so workers inherit the loaded name lists


def _ir_one(args):
    """IR-enhanced model from one downloaded workbook (runs in the killable worker)."""
    src, name, work = args
    got, why = irent.make(src, name, work, ANON)
    if not got:
        return {'ok': False, 'reason': why}
    wb, category, title, originals = got
    path = str(Path(work) / 'model.xlsx')
    wb.save(path)
    return {'ok': True, 'path': path, 'category': category, 'title': title, 'originals': sorted(originals), 'asof': wb.asof}


def _process_one(args):
    src, name, work = args
    res = process(src, name, work, ANON)
    return res.__dict__


class Worker:
    """Runs process() in a child process so a pathological file can be killed."""

    def __init__(self):
        self.pool = multiprocessing.get_context('fork').Pool(1)

    def run(self, args, fn=None):
        try:
            return self.pool.apply_async(fn or _process_one, (args,)).get(timeout=FILE_SECONDS)
        except multiprocessing.TimeoutError:
            self.pool.terminate()
            os.system('pkill -9 -f soffice >/dev/null 2>&1')
            self.pool = multiprocessing.get_context('fork').Pool(1)
            return {'ok': False, 'reason': f'processing took longer than {FILE_SECONDS}s'}

    def close(self):
        self.pool.terminate()


def log(msg):
    print(f'[{time.strftime("%H:%M:%S")}] {msg}', flush=True)


def left():
    return BUDGET - (time.time() - START)


def clean_label(s, limit=150):
    s = re.sub(r'[\\/:*?"<>|]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s[:limit].strip().rstrip('.').strip()


class Run:
    def __init__(self):
        if DRY and not os.environ.get('GOOGLE_REFRESH_TOKEN'):
            self.store = None
        elif ENT_KIND:
            group = enterprise.group_of(ENT_KIND)
            self.store = Store(f'{ENT_ROOT}/{group}/{enterprise.label_of(ENT_KIND)}', 'Enterprise Wide Models - private', dry_run=DRY,
                               app=f'ent-{ENT_KIND}', state_name=f'state-{ENT_KIND}.json.gz')
        else:
            self.store = Store(dry_run=DRY)
        self.state = (self.store.load_state() if self.store else {}) or {}
        st = self.state
        for k, v in {'next': 1, 'seen': {}, 'queued': {}, 'queue': [], 'hashes': {}, 'sec': {'i': 0, 'done': {}}, 'counts': {}}.items():
            st.setdefault(k, v)
        if self.store and not DRY:
            st['next'] = max(st['next'], self.store.max_number() + 1)
            if ENT_KIND:
                for k in self.store.uploaded_keys():
                    st['seen'].setdefault(k, 'ok:already on Drive')
        self.tmp = Path(tempfile.mkdtemp(prefix='xm_'))
        self.made = 0
        self.last_save = time.time()

    def save(self, force=False):
        if self.store and (force or time.time() - self.last_save > 300):
            try:
                self.store.save_state(self.state)
            except Exception as err:  # noqa: BLE001  (progress is saved again later; never stop the run for it)
                log(f'state save failed, will retry: {type(err).__name__}: {err}')
                return
            self.last_save = time.time()

    def publish(self, path, label, ext, key, category, source):
        """Numbers, uploads and records one verified model file."""
        digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        if digest in self.state['hashes']:
            self.state['seen'][key] = 'rej:duplicate'
            return None
        n = self.state['next']
        name = self.store.upload(str(path), n, clean_label(label), ext, key, category) if self.store else f'{n} - {label}.{ext}'
        self.state['next'] = n + 1
        self.state['hashes'][digest] = n
        self.state['seen'][key] = f'ok:{n}'
        c = self.state['counts']
        c[category] = c.get(category, 0) + 1
        c[f'source:{source}'] = c.get(f'source:{source}', 0) + 1
        self.made += 1
        if n <= 100 and LOCAL:
            dest = Path(LOCAL) / folder_name(n)
            dest.mkdir(parents=True, exist_ok=True)
            shutil.copy(path, dest / name)
        log(f'  #{n} {name}')
        self.save()
        return n

    # --- scraped spreadsheets ------------------------------------------------------
    def scraped(self, until):
        s = sources.session()
        q = self.state['queue']
        done = 0
        worker = Worker()
        while q and time.time() < until:
            cand = q.pop(0)
            self.state['queued'].pop(cand['key'], None)
            if cand['key'] in self.state['seen']:
                continue
            work = self.tmp / f'c{done}'
            work.mkdir()
            try:
                r = s.get(cand['url'], timeout=90, stream=True)
                if not r.ok:
                    self.state['seen'][cand['key']] = f'rej:http {r.status_code}'
                    continue
                src = work / 'download.bin'
                size = 0
                t0 = time.time()
                with open(src, 'wb') as f:
                    for chunk in r.iter_content(1 << 16):
                        size += len(chunk)
                        if size > 40 * 1024 * 1024 or time.time() - t0 > 90:
                            break
                        f.write(chunk)
                r.close()
                raw_hash = hashlib.sha256(src.read_bytes()).hexdigest()
                if f'raw:{raw_hash}' in self.state['seen']:
                    self.state['seen'][cand['key']] = 'rej:duplicate download'
                    continue
                self.state['seen'][f'raw:{raw_hash}'] = 1
                res = worker.run((str(src), cand['name'], str(work)))
                if not res.get('ok'):
                    self.state['seen'][cand['key']] = f'rej:{res.get("reason")}'[:120]
                    log(f'  - {cand["source"]}: {cand["name"][:60]}: {res.get("reason")}')
                    continue
                title, cat = res['title'], res['category']
                self.publish(res['path'], f'{cat} - {title}' if title != cat else cat, res['ext'], cand['key'], cat, cand['source'])
            except Exception as err:  # noqa: BLE001
                self.state['seen'][cand['key']] = f'rej:error {type(err).__name__}'
                log(f'  ! {cand["name"][:60]}: {type(err).__name__}: {err}')
            finally:
                shutil.rmtree(work, ignore_errors=True)
                done += 1
        worker.close()
        log(f'scraped: processed {done}, {len(q)} still queued')

    # --- SEC company models ------------------------------------------------------------
    def verify(self, wb, path, allowed=()):
        """Saves, recalculates in LibreOffice and checks a generated model. Returns the recalculated file or None.
        allowed: further sheets that may hold hard-coded figures (a company's own data sheets)."""
        hard = [(ws.title, c.coordinate) for ws in wb.worksheets if ws.title not in INPUT_SHEETS and ws.title not in allowed
                for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
        if hard:
            return None, f'hard-coded numbers outside inputs: {hard[:3]}'
        wb.save(path)
        out = lo.convert([path], path.parent / 'recalc', 'xlsx').get(str(path))
        if not out:
            return None, 'recalculation failed'
        errors, formulas, values = lo.recalculated_errors(out)
        if errors:
            return None, f'formula errors {errors}'
        try:
            total = float(values.get('Checks!B3'))
        except (TypeError, ValueError):
            return None, 'checks not computed'
        if abs(total) > 0.01:
            return None, f'checks fail ({total})'
        empty = lo.empty_parts(out)
        if empty:
            return None, f'empty content: {empty[0]}'
        # Safety net: nothing in a delivered model may point outside the workbook.
        pkg = Package(out)
        links.clean(pkg)
        if links.has_external(pkg):
            return None, 'external link could not be removed'
        if pkg.broken_parts():
            return None, f'invalid XML in {pkg.broken_parts()[0]}'
        pkg.save(out)
        from openpyxl import load_workbook
        try:
            load_workbook(out).close()
        except Exception as err:  # noqa: BLE001
            return None, f'file does not open ({type(err).__name__})'
        return out, None

    def sec(self, until):
        s = requests.Session()
        s.headers['User-Agent'] = SEC_UA
        st = self.state['sec']
        if 'companies' not in st or st.get('listed_at', 0) < time.time() - 7 * 86400 and st['i'] >= len(st['companies']):
            data = s.get('https://www.sec.gov/files/company_tickers.json', timeout=60).json()
            ciks, seen = [], set()
            for v in data.values():
                if v['cik_str'] not in seen:
                    seen.add(v['cik_str'])
                    ciks.append(v['cik_str'])
            st['companies'], st['i'], st['listed_at'] = ciks, 0, time.time()
        while st['i'] < len(st['companies']) and time.time() < until:
            cik = st['companies'][st['i']]
            st['i'] += 1
            done = set(st['done'].get(str(cik), []))
            try:
                sub = s.get(f'https://data.sec.gov/submissions/CIK{cik:010d}.json', timeout=60)
                time.sleep(0.12)
                facts = s.get(f'https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json', timeout=120)
                time.sleep(0.12)
                if not (sub.ok and facts.ok):
                    continue
                sub, facts = sub.json(), facts.json()
            except (requests.RequestException, ValueError):
                continue
            sic = sub.get('sic') or ''
            sector = (sub.get('sicDescription') or 'Unclassified').title()
            try:
                sic_i = int(sic)
            except ValueError:
                sic_i = 0
            plan = []
            if sic_i in BANK_SIC:
                plan.append(('bank', lambda: bank.build(facts, sector, self.market, sic)))
            else:
                plan.append(('reit' if sic_i == 6798 else 'three', lambda: three.build(facts, sector, is_reit=sic_i == 6798)))
            plan += [('dcf', lambda: dcf.build(facts, sector, self.market, sic)),
                     ('lbo', lambda: lbo.build(facts, sector, self.market, sic)),
                     ('fpa', lambda: fpa.build(facts, sector))]
            for kind, make in plan:
                if kind in done or time.time() > until:
                    continue
                key = f'sec:{cik}:{kind}'
                try:
                    wb = make()
                except Exception as err:  # noqa: BLE001
                    log(f'  ! CIK {cik} {kind}: {type(err).__name__}: {err}')
                    traceback.print_exc(limit=2)
                    wb = None
                done.add(kind)
                if wb is None:
                    continue
                work = self.tmp / f'sec_{cik}_{kind}'
                work.mkdir()
                try:
                    out, why = self.verify(wb, work / 'model.xlsx')
                    if not out:
                        self.state['seen'][key] = f'rej:{why}'[:160]
                        log(f'  - CIK {cik} {kind}: {why}')
                        continue
                    self.publish(out, f'{TYPE_LABEL[kind]} - {sector}', 'xlsx', key, TYPE_LABEL[kind].split(' - ')[0], 'sec')
                finally:
                    shutil.rmtree(work, ignore_errors=True)
            st['done'][str(cik)] = sorted(done)
        log(f'sec: company {st["i"]}/{len(st["companies"])}')

    def enterprise(self, until):
        """Enterprise-wide models of one type (ENT_KIND) for every company, at the latest and the previous year end."""
        s = requests.Session()
        s.headers['User-Agent'] = SEC_UA
        st = self.state['sec']
        if ('companies' not in st or st.get('list_v') != LIST_V
                or st.get('listed_at', 0) < time.time() - 7 * 86400 and st['i'] >= len(st['companies'])):
            st['companies'], st['i'], st['listed_at'], st['list_v'] = sec_filers(s), 0, time.time(), LIST_V
            log(f'enterprise {ENT_KIND}: {len(st["companies"])} companies listed')
        label = enterprise.label_of(ENT_KIND)
        while st['i'] < len(st['companies']) and time.time() < until:
            cik = st['companies'][st['i']]
            st['i'] += 1
            keys = {t: 'ent:' + model_code(f'{cik}:{ENT_KIND}:{t}') for t in enterprise.TRIMS}
            todo = [t for t in enterprise.TRIMS if keys[t] not in self.state['seen']]
            if not todo:
                continue
            try:
                sub = s.get(f'https://data.sec.gov/submissions/CIK{cik:010d}.json', timeout=60)
                time.sleep(0.12)
                if not sub.ok:
                    continue
                sub = sub.json()
                sic = sub.get('sic') or ''
                if not enterprise.applies(ENT_KIND, sic):
                    for t in todo:
                        self.state['seen'][keys[t]] = 'rej:not this industry'
                    continue
                facts = s.get(f'https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json', timeout=120)
                time.sleep(0.12)
                if not facts.ok:
                    continue
                facts = facts.json()
            except (requests.RequestException, ValueError):
                continue
            sector = (sub.get('sicDescription') or 'Diversified').title()
            for trim in todo:
                key = keys[trim]
                try:
                    make = bankent.build if ENT_KIND in bankent.KINDS else enterprise.build
                    prof = unique.Profile(key)
                    wb = make(facts, sector, self.market, sic, ENT_KIND, trim, profile=prof)
                    if wb is not None:
                        wb.profile = prof
                        allowed = finish_unique(wb, key)
                except Exception as err:  # noqa: BLE001
                    log(f'  ! {key[4:]}: {type(err).__name__}: {err}')
                    self.state['seen'][key] = f'rej:{type(err).__name__}'
                    continue
                if wb is None:
                    self.state['seen'][key] = 'rej:not applicable'
                    continue
                work = self.tmp / f'ent_{cik}_{trim}'
                work.mkdir()
                try:
                    out, why = self.verify(wb, work / 'model.xlsx', allowed=allowed)
                    if not out:
                        self.state['seen'][key] = f'rej:{why}'[:160]
                        log(f'  - {key[4:]}: {why}')
                        continue
                    self.publish(out, f'{SHORT[ENT_KIND]} - {short_name(sector)} FY{str(wb.asof)[-2:]} - {key[7:13]}', 'xlsx',
                                 key, sector, 'sec')
                finally:
                    shutil.rmtree(work, ignore_errors=True)
        log(f'enterprise {ENT_KIND}: company {st["i"]}/{len(st["companies"])}')

    def ir(self, companies_json):
        """IR-enhanced models: crawl company investor-relations sites for workbooks, then turn each into a model."""
        s = sources.session()
        if len(self.state['queue']) < 400:
            try:
                sources.websites(self.state, s, log, companies_json, max_domains=300, ir_hosts=True,
                                 minutes=max(2, min(20, 0.4 * left() / 60)))
            except Exception as err:  # noqa: BLE001
                log(f'websites: failed ({type(err).__name__}: {err})')
            self.save(force=True)
        q = self.state['queue']
        until = time.time() + left() - 120
        worker = Worker()
        done = 0
        while q and time.time() < until:
            cand = q.pop(0)
            self.state['queued'].pop(cand['key'], None)
            key = 'ir:' + model_code(cand['url'])
            if key in self.state['seen']:
                continue
            work = self.tmp / f'ir{done}'
            work.mkdir()
            try:
                r = s.get(cand['url'], timeout=90, stream=True)
                if not r.ok:
                    self.state['seen'][key] = f'rej:http {r.status_code}'
                    continue
                src = work / 'download.bin'
                size, t0 = 0, time.time()
                with open(src, 'wb') as f:
                    for chunk in r.iter_content(1 << 16):
                        size += len(chunk)
                        if size > 40 * 1024 * 1024 or time.time() - t0 > 90:
                            break
                        f.write(chunk)
                r.close()
                raw = 'raw:' + hashlib.sha256(src.read_bytes()).hexdigest()
                if raw in self.state['seen']:
                    self.state['seen'][key] = 'rej:duplicate download'
                    continue
                self.state['seen'][raw] = 1
                res = worker.run((str(src), cand['name'], str(work)), _ir_one)
                if not res.get('ok'):
                    self.state['seen'][key] = f'rej:{res.get("reason")}'[:120]
                    log(f'  - {key[3:]}: {res.get("reason")}')
                    continue
                from openpyxl import load_workbook
                wb = load_workbook(res['path'])
                allowed = finish_unique(wb, key, res['originals'])
                out, why = self.verify(wb, Path(res['path']).parent / 'final.xlsx', allowed=allowed)
                if not out:
                    self.state['seen'][key] = f'rej:{why}'[:160]
                    log(f'  - {key[3:]}: {why}')
                    continue
                code = model_code(cand['url'])
                self.publish(out, f'{SHORT[ENT_KIND]} - {short_name(res["title"], 26)} FY{str(res["asof"])[-2:]} - {code[3:9]}',
                             'xlsx', key, res['category'], 'website')
            except Exception as err:  # noqa: BLE001
                self.state['seen'][key] = f'rej:error {type(err).__name__}'
                log(f'  ! {key[3:]}: {type(err).__name__}: {err}')
            finally:
                shutil.rmtree(work, ignore_errors=True)
                done += 1
        worker.close()
        log(f'ir: processed {done}, {len(q)} still queued')

    def gov(self, until):
        """Government models of one type: every county, state, country or project with complete statistics."""
        s = requests.Session()
        s.headers['User-Agent'] = SEC_UA
        entities = govent.load(ENT_KIND, s, log)
        st = self.state.setdefault('gov', {})
        done = 0
        for eid, e in entities:
            if time.time() > until:
                break
            key = 'gov:' + model_code(f'{eid}:{ENT_KIND}')
            if key in self.state['seen']:
                done += 1
                continue
            try:
                prof = unique.Profile(key)
                wb = govent.build(ENT_KIND, e, self.market, profile=prof)
                if wb is not None:
                    wb.profile = prof
                    allowed = finish_unique(wb, key)
            except Exception as err:  # noqa: BLE001
                log(f'  ! {key[4:]}: {type(err).__name__}: {err}')
                self.state['seen'][key] = f'rej:{type(err).__name__}'
                continue
            if wb is None:
                self.state['seen'][key] = 'rej:not applicable'
                continue
            work = self.tmp / f'gov_{done}'
            work.mkdir(exist_ok=True)
            try:
                out, why = self.verify(wb, work / 'model.xlsx', allowed=allowed)
                if not out:
                    self.state['seen'][key] = f'rej:{why}'[:160]
                    log(f'  - {key[4:]}: {why}')
                    continue
                cat = e.get('sector') if ENT_KIND == 'gov_project' else ('United States' if e.get('kind') == 'us' else 'Countries')
                place = short_name(e.get('country') or e.get('name'), 22)
                what = f'{place} {short_name(e.get("sector"), 18)}' if ENT_KIND == 'gov_project' else f'{place} {wb.asof}'
                self.publish(out, f'{SHORT[ENT_KIND]} - {what} - {key[7:13]}', 'xlsx', key, cat or 'Other', 'gov')
            finally:
                shutil.rmtree(work, ignore_errors=True)
            done += 1
        st['done'], st['total'] = done, len(entities)
        log(f'government {ENT_KIND}: {done}/{len(entities)} done')

    def go(self):
        log(f'start: next model number {self.state["next"]}, {len(self.state["queue"])} candidates queued')
        if ENT_KIND:
            self.market = Market(log)
            if ENT_KIND in govent.KINDS:
                self.gov(time.time() + left() - 120)
            elif ENT_KIND == irent.KIND:
                self.ir(self.setup_anonymizer())
            else:
                self.enterprise(time.time() + left() - 120)
            self.save(force=True)
            c = self.state['counts']
            log(f'done: {self.made} models this run; next number {self.state["next"]}')
            summary = os.environ.get('GITHUB_STEP_SUMMARY')
            if summary:
                with open(summary, 'a', encoding='utf-8') as f:
                    f.write(f'## Enterprise wide models - {enterprise.label_of(ENT_KIND)}\n\n{self.made} models this run, '
                            f'{self.state["next"] - 1} in total, ' +
                            (f'{self.state.get("gov", {}).get("done", 0)} of {self.state.get("gov", {}).get("total", 0)} done'
                             if ENT_KIND in govent.KINDS else
                             f'company {self.state["sec"]["i"]} of {len(self.state["sec"].get("companies", []))}') +
                            '.\n\n| Group | Models |\n|---|---|\n')
                    for k in sorted(c, key=lambda k: -c[k]):
                        if not k.startswith('source:'):
                            f.write(f'| {k} | {c[k]} |\n')
            return
        if self.store:
            used, limit = self.store.quota()
            log(f'Drive: {used / 2**30:.1f} GB used of {limit / 2**30:.1f} GB')
            if limit and limit - used < 2 * 2**30:
                log('Drive has less than 2 GB free; stopping')
                return
        self.market = Market(log)
        log(f'market: rf {self.market.rf:.2%}, ERP {self.market.erp:.2%}, SOFR {self.market.sofr:.2%}, {len(self.market.betas)} industry betas')
        companies_json = self.setup_anonymizer()

        # Uploads start straight away: SEC company models first (fast, almost all pass), then
        # source collection and scraped files in the last part of the run.
        self.old_collector(companies_json)

    def setup_anonymizer(self):
        tickers = requests.get('https://www.sec.gov/files/company_tickers.json', headers={'User-Agent': SEC_UA}, timeout=60).json()
        names = [v['title'] for v in tickers.values()]
        companies_json = None
        cj = Path(__file__).resolve().parents[2] / 'companies.json'
        if cj.exists():
            companies_json = json.loads(cj.read_text(encoding='utf-8'))
            for c in companies_json.get('companies', []):
                names += [c.get('name', '')] + list(c.get('aliases') or [])
        try:
            import spacy
            nlp = spacy.load('en_core_web_sm', exclude=['parser', 'lemmatizer', 'attribute_ruler', 'tagger', 'senter'])
        except Exception:  # noqa: BLE001
            nlp = None
        global ANON
        self.anonymizer = ANON = Anonymizer(names, [v['title'] for v in list(tickers.values())[:1500]], nlp)
        log(f'anonymizer: {len(self.anonymizer.multi)} multi-word and {len(self.anonymizer.single)} one-word company names')
        return companies_json

    def old_collector(self, companies_json):
        self.sec(time.time() + 0.75 * left())
        self.save(force=True)
        s = sources.session()
        refresh_until = time.time() + min(0.25 * left(), 5 * 60)
        for name, fn in [('github', lambda: sources.github(self.state, s, log)),
                         ('damodaran', lambda: sources.damodaran(self.state, s, log)),
                         ('zenodo', lambda: sources.zenodo(self.state, s, log, max_pages=6)),
                         ('websites', lambda: sources.websites(self.state, s, log, companies_json,
                                                               minutes=max(1, (refresh_until - time.time()) / 60)))]:
            if time.time() > refresh_until:
                break
            try:
                fn()
            except Exception as err:  # noqa: BLE001
                log(f'{name}: failed ({type(err).__name__}: {err})')
        self.save(force=True)
        self.scraped(time.time() + left() - 60)
        self.save(force=True)
        c = self.state['counts']
        log(f'done: {self.made} models this run; next number {self.state["next"]}; totals {json.dumps(c, sort_keys=True)}')
        summary = os.environ.get('GITHUB_STEP_SUMMARY')
        if summary:
            with open(summary, 'a', encoding='utf-8') as f:
                f.write(f'## Excel models\n\n{self.made} models this run, {self.state["next"] - 1} in total.\n\n| Category / source | Count |\n|---|---|\n')
                for k in sorted(c):
                    f.write(f'| {k} | {c[k]} |\n')


if __name__ == '__main__':
    code = 0
    try:
        if os.environ.get('RESET', '').lower() == 'true':
            Store().reset(log)
        if os.environ.get('ENT_SETUP'):
            # One job creates every group and type folder before the parallel jobs start (no duplicate folders).
            for kind in os.environ['ENT_SETUP'].split(','):
                Store(f'{ENT_ROOT}/{enterprise.group_of(kind)}/{enterprise.label_of(kind)}', 'Enterprise Wide Models - private',
                      app=f'ent-{kind}', state_name=f'state-{kind}.json.gz')
                log(f'folders ready: {enterprise.group_of(kind)}/{enterprise.label_of(kind)}')
        else:
            Run().go()
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        code = 1
    sys.stdout.flush()
    os._exit(code)  # do not wait for website-crawler threads still finishing a request
