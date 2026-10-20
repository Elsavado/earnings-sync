"""Enterprise-wide models: the integrated three-statement model (three.py) extended with modules.

Every enterprise model has
- a scenario manager (Base / Upside / Downside, one switch) that moves the forecast drivers,
- supporting schedules (working capital days, PP&E and equity roll-forwards, debt and tax),
- a dashboard with KPIs and charts,
plus the modules of its type:
  integrated  - the above only (REITs keep their FFO/AFFO schedule)
  valuation   - DCF with WACC, perpetuity and exit-multiple terminal values, two sensitivity grids, football field
  buyout      - leveraged buyout on the integrated forecast: sources & uses, debt schedule with sweep, returns, IRR grid
  budget      - monthly budget for the first forecast year, department plan and budget-vs-actual variance
  credit      - credit statistics, synthetic rating, covenant tests, liquidity and debt capacity
  strategic   - ten-year horizon with valuation, credit and budget modules together

Reported figures stay on 'Historical Data'; assumptions live on input sheets (xl.INPUT_SHEETS); everything else is
a formula, verified by a LibreOffice recalculation in main.py (Checks!B3 must be zero).
"""
import re

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from . import three, xbrl
from .xl import BOLD, INPUT, INPUT_FILL, MULT, NAVY, NOTE, NUM, PCT, PRICE, Inputs, Sheet, col

KINDS = {
    'integrated': 'Integrated financial plan',
    'valuation': 'Valuation (DCF)',
    'buyout': 'Leveraged buyout',
    'budget': 'Budget & FP&A',
    'credit': 'Credit & liquidity',
    'strategic': 'Strategic 10-year plan',
}
BANK_KINDS = {}  # filled in by bankent.py
GOV_KINDS = {}  # filled in by govent.py
GROUPS = {'corporate': 'Corporate & Industry Models', 'bank': 'Bank Models', 'gov': 'Government Models'}
TRIMS = (0, 1)  # as of the latest year end, and as of the year before
FIN = range(6000, 6800)  # banks, insurers, brokers, REITs: no unlevered DCF, LBO or corporate credit module
BANK_SIC = set(range(6020, 6037)) | {6111, 6141, 6153, 6159, 6162, 6172, 6199}


def group_of(kind):
    return GROUPS['bank' if kind in BANK_KINDS else 'gov' if kind in GOV_KINDS else 'corporate']


def label_of(kind):
    return BANK_KINDS.get(kind) or GOV_KINDS.get(kind) or KINDS[kind]


def applies(kind, sic):
    """Bank models for banks and lenders; every other model type for every other company."""
    return (_sic(sic) in BANK_SIC) == (kind in BANK_KINDS)
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']

# Interest coverage -> rating and spread (synthetic rating approach for large non-financial firms; editable inputs).
RATINGS = [(-100000, 'D', 0.19), (0.2, 'C', 0.155), (0.65, 'CC', 0.1176), (0.8, 'CCC', 0.0856), (1.25, 'B-', 0.0515),
           (1.5, 'B', 0.0403), (1.75, 'B+', 0.0339), (2.0, 'BB', 0.0275), (2.25, 'BB+', 0.0221), (2.5, 'BBB', 0.0174),
           (3.0, 'A-', 0.0147), (4.25, 'A', 0.0121), (5.5, 'A+', 0.0107), (6.5, 'AA', 0.0092), (8.5, 'AAA', 0.0070)]


def _sic(sic):
    try:
        return int(sic or 0)
    except ValueError:
        return 0


def _plain(f):
    return f[1:] if f.startswith('=') else f


def _title(ws, text, note=None):
    ws['A1'] = text
    ws['A1'].font = Font(bold=True, size=13, color=NAVY)
    if note:
        ws['A2'] = note
        ws['A2'].font = NOTE
    ws.sheet_view.showGridLines = False


# --- scenario manager --------------------------------------------------------------------------------------------
SCEN = [('growth', 'Revenue growth adjustment (added to starting growth)', 0.0, 0.02, -0.03, PCT),
        ('gm', 'Gross margin adjustment (cost of revenue % revenue reduced by)', 0.0, 0.01, -0.015, PCT),
        ('opex', 'Operating expense efficiency (opex % revenue reduced by)', 0.0, 0.005, -0.01, PCT),
        ('capex', 'Capital intensity multiplier (x capex % revenue)', 1.0, 0.95, 1.10, '0.00"x"'),
        ('wc', 'Working capital multiplier (x receivables and inventory ratios)', 1.0, 0.95, 1.10, '0.00"x"')]


def _clamp(x, lo, hi):
    return max(lo, min(hi, x))


def _sized(c):
    """Scenario cases sized by the company's own history: volatile companies get wider cases."""
    import statistics
    d = c['data']
    rev = [x for x in d['revenue'] if x]
    g = [b / a - 1 for a, b in zip(rev, rev[1:]) if a > 0][-6:]
    sd = statistics.pstdev(g) if len(g) >= 2 else 0.05
    gm = [(r - (k or 0)) / r for r, k in zip(d['revenue'], d['cogs']) if r][-6:]
    sd_gm = statistics.pstdev(gm) if len(gm) >= 2 else 0.01
    cx = [(x or 0) / r for x, r in zip(d['capex'], d['revenue']) if r][-6:]
    cv_cx = statistics.pstdev(cx) / max(1e-9, statistics.mean(cx)) if len(cx) >= 2 and statistics.mean(cx) > 0 else 0.1
    j = c.get('profile').rng.uniform(0.9, 1.1) if c.get('profile') else 1.0
    up_g, dn_g = round(_clamp(0.5 * sd * j, 0.01, 0.06), 4), -round(_clamp(0.8 * sd * j, 0.015, 0.10), 4)
    up_m, dn_m = round(_clamp(sd_gm * j, 0.003, 0.03), 4), -round(_clamp(1.5 * sd_gm * j, 0.005, 0.045), 4)
    cx_d = round(_clamp(0.5 * cv_cx * j, 0.03, 0.2), 3)
    return [('growth', 'Revenue growth adjustment (added to starting growth)', 0.0, up_g, dn_g, PCT),
            ('gm', 'Gross margin adjustment (cost of revenue % revenue reduced by)', 0.0, up_m, dn_m, PCT),
            ('opex', 'Operating expense efficiency (opex % revenue reduced by)', 0.0, round(up_m / 2, 4), round(dn_m / 2, 4), PCT),
            ('capex', 'Capital intensity multiplier (x capex % revenue)', 1.0, round(1 - cx_d, 3), round(1 + cx_d, 3), '0.00"x"'),
            ('wc', 'Working capital multiplier (x receivables and inventory ratios)', 1.0, round(1 - cx_d / 2, 3),
             round(1 + cx_d / 1.5, 3), '0.00"x"')]


# --- company-specific content ----------------------------------------------------------------------------------------
def _gaap(facts):
    return (facts.get('facts', facts) or {}).get('us-gaap', {})


def company_lines(wb, c, facts, base=None, base_label='revenue'):
    """The company's own further reported line items (beyond the standard statements), forecast with revenue."""
    from . import xbrl as X
    used = {t for tags in list(X.FLOW_TAGS.values()) + list(X.INSTANT_TAGS.values()) for t in tags}
    ends, nh, n = c['ends'], c['nh'], c['n']
    hi = nh - 1
    base = base or (lambda j: c['is_'].ref('revenue', j))
    rev_last = (c['data']['revenue'][hi] if 'revenue' in c['data'] else c['base_last']) or 1
    known = {round(v, 1) for k, vals in c['data'].items() for v in [vals[hi]] if isinstance(v, (int, float))}
    cands = []
    for tag, body in _gaap(facts).items():
        if tag in used or not (body.get('units') or {}).get('USD'):
            continue
        fl = {e: v for (s0, e), v in X.flows(facts, [tag], 330, 400).items()}
        series, kind = (fl, 'flow') if fl else (X.instants(facts, [tag]), 'stock')
        vals = [series.get(e) for e in ends]
        if vals[-1] in (None, 0) or sum(v is not None for v in vals[-3:]) < min(3, nh):
            continue
        vals = [None if v is None else round(v / 1e6, 3) for v in vals]
        if abs(vals[-1]) < 0.002 * abs(rev_last) or round(vals[-1], 1) in known:
            continue
        label = re.sub(r'\s+', ' ', body.get('label') or re.sub(r'(?<!^)(?=[A-Z])', ' ', tag)).strip()[:70]
        if SOURCE_WORDS.search(label):
            continue
        cands.append((abs(vals[-1]), tag, label, vals, kind))
    cands.sort(reverse=True)
    prof = c.get('profile')
    take = prof.pick([8, 10, 12, 14, 16]) if prof else 10
    picked, labels = [], set()
    for _, tag, label, vals, kind in cands:
        if label.lower() in labels:
            continue
        labels.add(label.lower())
        picked.append((tag, label, vals, kind))
        if len(picked) >= take:
            break
    if len(picked) < 3:
        return [], []
    ds = c['ds']
    from .xl import INPUT as _IN
    for i, (tag, label, vals, kind) in enumerate(picked):
        ds.row(f'cx{i}', label)
        for j, v in enumerate(vals):
            if v is not None:
                ds.set(f'cx{i}', j, v, font=_IN)
    sh = Sheet(wb, 'Company Detail', c['heading']('Company-specific line items (forecast with revenue)'), c['periods'], nh)
    for kind_name, sec in (('flow', 'Income and cash-flow items'), ('stock', 'Balance sheet items')):
        first = True
        for i, (tag, label, vals, kind) in enumerate(picked):
            if kind != kind_name:
                continue
            sh.row(f'v{i}', label, section=sec if first else None)
            sh.row(f'r{i}', f'  % of {base_label}')
            first = False
    w0 = max(0, nh - 3)
    for i, (tag, label, vals, kind) in enumerate(picked):
        for j in range(n):
            if j < nh:
                if vals[j] is not None:
                    sh.set(f'v{i}', j, f'={ds.ref(f"cx{i}", j)}')
            else:
                sh.set(f'v{i}', j, f'={base(j)}*IFERROR(AVERAGE({sh.span(f"r{i}", w0, hi)}),0)')
            sh.set(f'r{i}', j, f'=IFERROR({sh.ref(f"v{i}", j)}/{base(j)},0)', PCT)
    c.setdefault('dash_extra', []).append(('cx_n', 'Company-specific line items modelled',
                                           f"=COUNTIF('Company Detail'!$A$4:$A${sh.next_row},\"  % of {base_label}\")", '0'))
    return [('Company Detail', f'{len(picked)} further line items this company reports, each forecast as a share of {base_label}')], []


def rnd(wb, c, facts):
    """R&D analysis with capitalised R&D, when the company reports research and development spending."""
    from . import xbrl as X
    fl = {e: v for (s0, e), v in X.flows(facts, ['ResearchAndDevelopmentExpense',
                                                  'ResearchAndDevelopmentExpenseExcludingAcquiredInProcessCost', 'ResearchDevelopmentAndRelatedExpenses', 'ResearchAndDevelopmentExpenseSoftwareExcludingAcquiredInProcessCost'], 330, 400).items()}
    ends, nh, n = c['ends'], c['nh'], c['n']
    vals = [None if fl.get(e) is None else round(fl[e] / 1e6, 3) for e in ends]
    if sum(v is not None and v > 0 for v in vals[-3:]) < 3:
        return [], []
    hi = nh - 1
    ds, I, si = c['ds'], c['is_'].ref, c['si']
    from .xl import INPUT as _IN
    ds.row('rd', 'Research and development expense')
    for j, v in enumerate(vals):
        if v is not None:
            ds.set('rd', j, v, font=_IN)
    life = c['profile'].pick([3, 4, 5]) if c.get('profile') else 5
    si.next_row += 1
    si.add('rd_life', 'R&D amortisation life (years, capitalised R&D)', life, '0', 'Typical useful life of research spending')
    L = si.ref('rd_life')
    sh = Sheet(wb, 'R&D Analysis', c['heading']('Research and development: spending, capitalisation and adjusted returns'), c['periods'], nh)
    for key, label, bold, sec in [('rd', 'R&D expense', True, 'Spending'), ('pct', '  % of revenue', False, None),
                                  ('amort', 'Amortisation of capitalised R&D', False, 'Capitalised R&D'),
                                  ('asset', 'Research asset (unamortised R&D)', True, None),
                                  ('ebit', 'Reported EBIT', False, 'Adjusted operating income'),
                                  ('adj', 'EBIT with R&D capitalised', True, None), ('m', 'Adjusted EBIT margin', False, None)]:
        sh.row(key, label, bold, section=sec)
    S = sh.ref
    w0 = max(0, nh - 3)
    for j in range(n):
        if j < nh:
            sh.set('rd', j, f'=N({ds.ref("rd", j)})', bold=True)
        else:
            sh.set('rd', j, f'={I("revenue", j)}*AVERAGE({sh.span("pct", w0, hi)})', bold=True)
        sh.set('pct', j, f'=IFERROR({S("rd", j)}/{I("revenue", j)},0)', PCT)
        terms = '+'.join(f'IF({k}<{L},{S("rd", j - k)},0)' for k in range(0, min(j, 7) + 1))
        sh.set('amort', j, f'=({terms})/{L}')
        unam = '+'.join(f'IF({k}<{L},{S("rd", j - k)}*({L}-{k}-1)/{L},0)' for k in range(0, min(j, 7) + 1))
        sh.set('asset', j, f'={unam}', bold=True)
        sh.set('ebit', j, f'={I("ebit", j)}')
        sh.set('adj', j, f'={S("ebit", j)}+{S("rd", j)}-{S("amort", j)}', bold=True)
        sh.set('m', j, f'=IFERROR({S("adj", j)}/{I("revenue", j)},0)', PCT)
    return [('R&D Analysis', 'R&D intensity, capitalised R&D and R&D-adjusted operating income')], []


def sbc(wb, c):
    """Share-based pay against EBITDA, when the company reports it."""
    d, nh, n = c['data'], c['nh'], c['n']
    vals = d.get('sbc') or []
    if sum(v is not None and v > 0 for v in vals[-3:]) < 3:
        return [], []
    hi = nh - 1
    ds, I = c['ds'], c['is_'].ref
    from .xl import INPUT as _IN
    ds.row('sbc_x', 'Share-based compensation')
    for j, v in enumerate(vals):
        if v is not None:
            ds.set('sbc_x', j, v, font=_IN)
    sh = Sheet(wb, 'Share-Based Pay', c['heading']('Share-based compensation and cash-adjusted earnings'), c['periods'], nh)
    for key, label, bold in [('sbc', 'Share-based compensation', True), ('pct', '  % of revenue', False),
                             ('ebitda', 'EBITDA', False), ('adj', 'EBITDA after share-based pay', True),
                             ('share', 'Share-based pay % EBITDA', False)]:
        sh.row(key, label, bold)
    S = sh.ref
    w0 = max(0, nh - 3)
    for j in range(n):
        sh.set('sbc', j, f'=N({ds.ref("sbc_x", j)})' if j < nh else f'={I("revenue", j)}*AVERAGE({sh.span("pct", w0, hi)})', bold=True)
        sh.set('pct', j, f'=IFERROR({S("sbc", j)}/{I("revenue", j)},0)', PCT)
        sh.set('ebitda', j, f'={I("ebitda", j)}')
        sh.set('adj', j, f'={S("ebitda", j)}-{S("sbc", j)}', bold=True)
        sh.set('share', j, f'=IFERROR({S("sbc", j)}/{S("ebitda", j)},0)', PCT)
    return [('Share-Based Pay', 'Share-based compensation intensity and EBITDA after share-based pay')], []


def scenarios(wb, c):
    si, asm = c['si'], c['asm']
    ws = si.ws
    r = si.next_row + 1
    ws.cell(row=r, column=1, value='Scenario manager').font = Font(bold=True, color=NAVY)
    r += 1
    ws.cell(row=r, column=1, value='Active scenario (1 = Base, 2 = Upside, 3 = Downside)').font = BOLD
    sel = ws.cell(row=r, column=2, value=1)
    sel.font, sel.fill, sel.number_format = INPUT, INPUT_FILL, '0'
    dv = DataValidation(type='list', formula1='"1,2,3"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(sel.coordinate)
    sel_ref = f"'Scenario Inputs'!$B${r}"
    r += 1
    ws.cell(row=r, column=1, value='Active scenario name')
    ws.cell(row=r, column=2, value=f'=CHOOSE({sel_ref},"Base","Upside","Downside")').font = BOLD
    name_ref = f"'Scenario Inputs'!$B${r}"
    r += 2
    for k, h in enumerate(['Driver adjustment', 'Base', 'Upside', 'Downside', 'Live (active case)']):
        cell = ws.cell(row=r, column=1 + k, value=h)
        cell.font = BOLD
    for letter in 'DE':
        ws.column_dimensions[letter].width = 16
    live = {}
    for key, label, base, up, down, fmt in _sized(c):
        r += 1
        ws.cell(row=r, column=1, value=label)
        for k, v in enumerate((base, up, down)):
            cell = ws.cell(row=r, column=2 + k, value=v)
            cell.font, cell.fill, cell.number_format = INPUT, INPUT_FILL, fmt
        cell = ws.cell(row=r, column=5, value=f'=CHOOSE({sel_ref},B{r},C{r},D{r})')
        cell.number_format, cell.font = fmt, BOLD
        live[key] = f"'Scenario Inputs'!$E${r}"
    si.next_row = r + 2

    def adjust(key, fn):
        cell = asm.ws.cell(row=asm.rows[key], column=2)
        cell.value = '=' + fn(_plain(cell.value))
    adjust('growth', lambda f: f'{f}+{live["growth"]}')
    adjust('cogs_pct', lambda f: f'MAX(0,{f}-{live["gm"]})')
    adjust('opex_pct', lambda f: f'{f}-{live["opex"]}')
    adjust('capex_pct', lambda f: f'({f})*{live["capex"]}')
    adjust('ar_pct', lambda f: f'({f})*{live["wc"]}')
    adjust('inv_pct', lambda f: f'({f})*{live["wc"]}')
    asm.ws.cell(row=asm.next_row + 2, column=1,
                value='Drivers include the live scenario adjustments from the Scenario Inputs sheet.').font = NOTE
    c['scenario_name'] = name_ref
    return [], [('Scenario switch is 1, 2 or 3', f'IF(OR({sel_ref}=1,{sel_ref}=2,{sel_ref}=3),0,1)')]


# --- supporting schedules ----------------------------------------------------------------------------------------
def schedules(wb, c):
    I, B, C, R = c['is_'].ref, c['bs'].ref, c['cf'].ref, c['rt'].ref
    n, nh, periods = c['n'], c['nh'], c['periods']
    sc = Sheet(wb, 'Supporting Schedules', c['heading']('Supporting schedules'), periods, nh)
    rows = [('dso', 'Days sales outstanding', 'Working capital'), ('dio', 'Days inventory outstanding', None),
            ('dpo', 'Days payables outstanding', None), ('ccc', 'Cash conversion cycle (days)', None),
            ('nwc', 'Net working capital (excl. cash and debt)', None), ('nwc_pct', '  % revenue', None),
            ('dnwc', 'Change in net working capital', None),
            ('ppe_open', 'Opening PP&E', 'PP&E roll-forward'), ('ppe_capex', 'Plus: capital expenditure', None),
            ('ppe_da', 'Less: depreciation & amortisation', None), ('ppe_other', 'Acquisitions, disposals & other', None),
            ('ppe_close', 'Closing PP&E', None), ('ppe_chk', 'Check: roll-forward vs balance sheet', None),
            ('capex_da', 'Capex / D&A (x)', None), ('capex_pct', 'Capex % revenue', None),
            ('debt_open', 'Opening debt (long-term + revolver)', 'Debt and interest'), ('debt_close', 'Closing debt', None),
            ('debt_net', 'Net borrowing / (repayment)', None), ('int', 'Interest expense', None),
            ('int_rate', 'Average interest rate on debt', None), ('net_debt', 'Net debt', None),
            ('nd_ebitda', 'Net debt / EBITDA (x)', None),
            ('eq_open', 'Opening shareholders\' equity', 'Equity roll-forward'), ('eq_ni', 'Plus: net income', None),
            ('eq_div', 'Less: dividends', None), ('eq_bb', 'Less: share repurchases', None),
            ('eq_other', 'Share issuance, OCI & other', None), ('eq_close', 'Closing shareholders\' equity', None),
            ('eq_chk', 'Check: roll-forward vs balance sheet', None),
            ('pretax', 'Pre-tax income', 'Tax'), ('tax', 'Income tax expense', None), ('etr', 'Effective tax rate', None),
            ('nopat', 'NOPAT (EBIT x (1 - effective rate))', None)]
    for key, label, sec in rows:
        sc.row(key, label, bold=key in ('ccc', 'nwc', 'ppe_close', 'debt_close', 'eq_close', 'nopat'), section=sec)
    S = sc.ref
    for j in range(n):
        p = j - 1
        sc.set('dso', j, f'=IFERROR({B("ar", j)}/{I("revenue", j)}*365,0)', '0.0')
        sc.set('dio', j, f'=IFERROR({B("inventory", j)}/{I("cogs", j)}*365,0)', '0.0')
        sc.set('dpo', j, f'=IFERROR({B("ap", j)}/{I("cogs", j)}*365,0)', '0.0')
        sc.set('ccc', j, f'={S("dso", j)}+{S("dio", j)}-{S("dpo", j)}', '0.0', bold=True)
        sc.set('nwc', j, f'={B("ar", j)}+{B("inventory", j)}+{B("oca", j)}-{B("ap", j)}-{B("ocl", j)}', bold=True)
        sc.set('nwc_pct', j, f'=IFERROR({S("nwc", j)}/{I("revenue", j)},0)', PCT)
        sc.set('ppe_close', j, f'={B("ppe", j)}' if j == 0 else
               f'={S("ppe_open", j)}+{S("ppe_capex", j)}-{S("ppe_da", j)}+{S("ppe_other", j)}', bold=True)
        sc.set('capex_da', j, f'=IFERROR(-{C("capex", j)}/{C("da", j)},0)', MULT)
        sc.set('capex_pct', j, f'=IFERROR(-{C("capex", j)}/{I("revenue", j)},0)', PCT)
        sc.set('debt_close', j, f'={B("debt", j)}+{B("revolver", j)}', bold=True)
        sc.set('int', j, f'={I("interest", j)}')
        sc.set('net_debt', j, f'={S("debt_close", j)}-{B("cash", j)}')
        sc.set('nd_ebitda', j, f'=IFERROR({S("net_debt", j)}/{I("ebitda", j)},0)', MULT)
        sc.set('eq_close', j, f'={B("equity", j)}' if j == 0 else
               f'={S("eq_open", j)}+{S("eq_ni", j)}-{S("eq_div", j)}-{S("eq_bb", j)}+{S("eq_other", j)}', bold=True)
        sc.set('pretax', j, f'={I("pretax", j)}')
        sc.set('tax', j, f'={I("tax", j)}')
        sc.set('etr', j, f'=IFERROR({S("tax", j)}/{S("pretax", j)},0)', PCT)
        sc.set('nopat', j, f'={I("ebit", j)}*(1-MAX(0,MIN(1,{S("etr", j)})))', bold=True)
        if not j:
            continue
        sc.set('dnwc', j, f'={S("nwc", j)}-{S("nwc", p)}')
        sc.set('ppe_open', j, f'={S("ppe_close", p)}')
        sc.set('ppe_capex', j, f'=-{C("capex", j)}')
        sc.set('ppe_da', j, f'={C("da", j)}')
        # History: acquisitions, disposals, impairments and leases explain the rest; the forecast has none.
        sc.set('ppe_other', j, f'={B("ppe", j)}-({S("ppe_open", j)}+{S("ppe_capex", j)}-{S("ppe_da", j)})' if j < nh
               else f'={S("ppe_open", j)}-{S("ppe_open", j)}')
        sc.set('ppe_chk', j, f'=ROUND({S("ppe_close", j)}-{B("ppe", j)},3)')
        sc.set('debt_open', j, f'={S("debt_close", p)}')
        sc.set('debt_net', j, f'={S("debt_close", j)}-{S("debt_open", j)}')
        sc.set('int_rate', j, f'=IFERROR({S("int", j)}/AVERAGE({S("debt_open", j)},{S("debt_close", j)}),0)', PCT)
        sc.set('eq_open', j, f'={S("eq_close", p)}')
        sc.set('eq_ni', j, f'={I("net_income", j)}')
        sc.set('eq_div', j, f'=-{C("div", j)}')
        sc.set('eq_bb', j, f'=-{C("buyback", j)}')
        sc.set('eq_other', j, f'={B("equity", j)}-({S("eq_open", j)}+{S("eq_ni", j)}-{S("eq_div", j)}-{S("eq_bb", j)})' if j < nh
               else f'={S("eq_open", j)}-{S("eq_open", j)}')
        sc.set('eq_chk', j, f'=ROUND({S("eq_close", j)}-{B("equity", j)},3)')
    c['sc'] = sc
    F = range(nh, n)
    items = [(f'PP&E roll-forward ties ({periods[j]})', S('ppe_chk', j)) for j in F]
    items += [(f'Equity roll-forward ties ({periods[j]})', S('eq_chk', j)) for j in F]
    return [('Supporting Schedules', 'Working capital days, PP&E and equity roll-forwards, debt, interest and tax')], items


# --- dashboard ---------------------------------------------------------------------------------------------------
def dashboard(wb, c):
    I, B, C, R = c['is_'].ref, c['bs'].ref, c['cf'].ref, c['rt'].ref
    is_, rt = c['is_'], c['rt']
    n, nh, years = c['n'], c['nh'], c['years']
    hi, fn = nh - 1, n - 1
    db = Inputs(wb, 'Dashboard', c['heading']('Executive dashboard'), index=0)
    db.ws['A2'] = 'Key outputs of the live scenario. Every figure is a formula; change inputs on the input sheets.'
    db.section('Scenario')
    if c.get('scenario_name'):
        db.add('scen', 'Active scenario', f'={c["scenario_name"]}', '@', formula=True)
    db.add('checks', 'Model integrity', '=IF(Checks!$B$3=0,"All checks pass","CHECK ERRORS")', '@', formula=True)
    db.section('Growth')
    db.add('rev_a', f'Revenue, last reported year ({c["periods"][hi]})', f'={I("revenue", hi)}', NUM, formula=True)
    db.add('rev_f', f'Revenue, final forecast year ({c["periods"][fn]})', f'={I("revenue", fn)}', NUM, formula=True)
    db.add('cagr_h', 'Historical revenue CAGR',
           f'=IFERROR(({I("revenue", hi)}/{I("revenue", 0)})^(1/{nh - 1})-1,0)', PCT, formula=True)
    db.add('cagr_f', 'Forecast revenue CAGR', f'=IFERROR(({I("revenue", fn)}/{I("revenue", hi)})^(1/{years})-1,0)', PCT, formula=True)
    db.section('Profitability and cash')
    db.add('em_h', 'EBITDA margin, last reported year', f'={R("em", hi)}', PCT, formula=True)
    db.add('em_f', 'Average EBITDA margin, forecast', f'=AVERAGE({rt.span("em", nh, fn)})', PCT, formula=True)
    db.add('ni_f', 'Cumulative net income, forecast', f'=SUM({is_.span("net_income", nh, fn)})', NUM, formula=True)
    db.add('fcf_f', 'Cumulative free cash flow, forecast', f'=SUM({rt.span("fcf", nh, fn)})', NUM, formula=True)
    db.add('conv', 'FCF conversion (FCF / net income), forecast', f'=IFERROR({db.ref("fcf_f")}/{db.ref("ni_f")},0)', PCT, formula=True)
    db.add('roe_f', 'Return on equity, final forecast year', f'={R("roe", fn)}', PCT, formula=True)
    db.section('Balance sheet')
    db.add('cash_f', 'Cash, final forecast year', f'={B("cash", fn)}', NUM, formula=True)
    db.add('rev_peak', 'Peak revolver draw, forecast', f'=MAX({c["bs"].span("revolver", nh, fn)})', NUM, formula=True)
    db.add('lev_f', 'Total debt / EBITDA, final forecast year', f'={R("lev", fn)}', MULT, formula=True)
    db.add('cov_f', 'EBITDA / interest, final forecast year', f'={R("cov", fn)}', MULT, formula=True)
    extra = c.get('dash_extra', [])
    if extra:
        db.section('Module results')
        for key, label, f, fmt in extra:
            db.add(key, label, f, fmt, formula=True)

    ws = db.ws
    first, last = 2, 1 + n
    lc = LineChart()
    lc.title, lc.y_axis.title, lc.height, lc.width = 'Revenue and EBITDA (USD millions)', 'USD m', 7.5, 16
    for key in ('revenue', 'ebitda'):
        r = is_.rows[key]
        lc.add_data(Reference(is_.ws, min_col=1, max_col=last, min_row=r, max_row=r), from_rows=True, titles_from_data=True)
    lc.set_categories(Reference(is_.ws, min_col=first, max_col=last, min_row=3, max_row=3))
    ws.add_chart(lc, 'E4')
    bc = BarChart()
    bc.title, bc.height, bc.width = 'Free cash flow (USD millions)', 7.5, 16
    r = rt.rows['fcf']
    bc.add_data(Reference(rt.ws, min_col=1, max_col=last, min_row=r, max_row=r), from_rows=True, titles_from_data=True)
    bc.set_categories(Reference(rt.ws, min_col=first, max_col=last, min_row=3, max_row=3))
    ws.add_chart(bc, 'E20')
    return [('Dashboard', 'KPIs of the live scenario with revenue, EBITDA and free cash flow charts')], []


# --- market inputs (valuation, buyout) ---------------------------------------------------------------------------
def market_inputs(wb, c, market, sic, facts):
    if 'mi' in c:
        return c['mi']
    bu, _, industry, beta_note = market.beta(sic)
    shares = xbrl.dei_latest(facts, 'EntityCommonStockSharesOutstanding', 'shares')
    float_ = xbrl.dei_latest(facts, 'EntityPublicFloat', 'USD')
    d, hi = c['data'], c['nh'] - 1
    ebitda = d['ebit'][hi] + d['da'][hi]
    net_debt = d['debt'][hi] + (d['debt_current'][hi] or 0) + d['st_debt'][hi] - d['cash'][hi]
    mult = 10.0
    note = 'Assumed EV / EBITDA (no market value reported)'
    if float_ and ebitda > 0:
        cur = (float_ / 1e6 + net_debt) / ebitda
        mult = round(min(16.0, max(6.0, cur)), 1)
        note = f'Current trading EV/EBITDA is {cur:.1f}x; kept within a 6x-16x range'
    mi = Inputs(wb, 'Market Inputs', c['heading']('Market and valuation inputs'))
    mi.section('Cost of capital')
    mi.add('rf', 'Risk-free rate', market.rf, PCT, 'Current 10-year government bond yield')
    mi.add('erp', 'Equity risk premium', market.erp, PCT, 'Current implied equity risk premium')
    mi.add('bu', 'Unlevered beta (industry)', round(bu, 3), '0.00', f'Industry average unlevered beta ({industry})')
    mi.add('tax', 'Marginal tax rate', 0.25, PCT, 'US federal 21% plus a typical state rate')
    mi.add('spread', 'Minimum credit spread over risk-free', 0.015, PCT, 'Floor for the pre-tax cost of debt')
    mi.add('min_eq', 'Minimum equity weight base (% of total assets)', 0.10, PCT,
           'Used only when no market value is available and book equity is small or negative')
    mi.add('sofr', 'SOFR', market.sofr, PCT, 'Current secured overnight financing rate')
    mi.section('Terminal value')
    mi.add('g', 'Terminal growth rate', 0.025, PCT, 'Long-run nominal growth; must stay below WACC')
    mi.add('exit', 'Exit EV / EBITDA multiple', mult, MULT, note)
    mi.add('mid', 'Mid-year discounting (1 = on, 0 = off)', 1, '0', 'Cash flows assumed to arrive mid-year')
    mi.section('Equity data')
    mi.add('shares', 'Shares outstanding (millions)', round(shares / 1e6, 3) if shares else 0, NUM,
           'Latest shares outstanding')
    mi.add('mcap', 'Market value of equity (USD millions)', round(float_ / 1e6, 3) if float_ else 0, NUM,
           'Market value of equity held by the public; used for capital-structure weights')
    mi.section('Sensitivity steps')
    mi.add('w_step', 'WACC step', 0.005, PCT)
    mi.add('g_step', 'Terminal growth step', 0.005, PCT)
    mi.add('m_step', 'Exit multiple step', 1.0, MULT)
    c['mi'], c['industry'] = mi, industry
    return mi


# --- valuation -----------------------------------------------------------------------------------------------------
def valuation(wb, c, market, sic, facts):
    M = market_inputs(wb, c, market, sic, facts).ref
    I, B, C = c['is_'].ref, c['bs'].ref, c['cf'].ref
    n, nh, periods = c['n'], c['nh'], c['periods']
    hi, F0, FN = nh - 1, nh, n - 1
    head = c['heading']
    dc = Sheet(wb, 'DCF', head('Unlevered free cash flow and discounting (live scenario)'), periods, nh)
    for key, label, bold, sec in [('ebit', 'EBIT', True, 'Free cash flow'), ('taxes', 'Less: taxes on EBIT', False, None),
                                  ('nopat', 'NOPAT', True, None), ('da', 'Plus: depreciation & amortisation', False, None),
                                  ('capex', 'Less: capital expenditure', False, None),
                                  ('wc', 'Less: investment in working capital', False, None),
                                  ('ufcf', 'Unlevered free cash flow', True, None), ('ebitda', 'EBITDA', False, None),
                                  ('year', 'Forecast year', False, 'Discounting'), ('period', 'Discount period', False, None),
                                  ('df', 'Discount factor', False, None), ('pv', 'Present value of UFCF', True, None)]:
        dc.row(key, label, bold, section=sec)
    P = dc.ref
    wc = Inputs(wb, 'WACC', head('Weighted average cost of capital'))
    wc.section('Capital structure')
    wc.add('debt', 'Total debt (last reported year)', f'={B("debt", hi)}+{B("revolver", hi)}', NUM, formula=True)
    wc.add('equity', 'Equity value used for weights', f'=IF({M("mcap")}>0,{M("mcap")},MAX({B("equity", hi)},{B("ta", hi)}*{M("min_eq")}))',
           NUM, 'Market value of equity if available, otherwise book equity (with a minimum)', formula=True)
    wc.add('de', 'Debt / equity', f'=IFERROR({wc.ref("debt")}/{wc.ref("equity")},0)', PCT, formula=True)
    wc.add('wd', 'Weight of debt', f'=IFERROR({wc.ref("debt")}/({wc.ref("debt")}+{wc.ref("equity")}),0)', PCT, formula=True)
    wc.add('we', 'Weight of equity', f'=1-{wc.ref("wd")}', PCT, formula=True)
    wc.section('Cost of equity (CAPM)')
    wc.add('bl', 'Relevered beta (Hamada)', f'={M("bu")}*(1+(1-{M("tax")})*{wc.ref("de")})', '0.00', formula=True)
    wc.add('ke', 'Cost of equity', f'={M("rf")}+{wc.ref("bl")}*{M("erp")}', PCT, formula=True)
    wc.section('Cost of debt')
    wc.add('kd_hist', 'Historical interest rate', f'=IFERROR({I("interest", hi)}/{wc.ref("debt")},0)', PCT, formula=True)
    wc.add('kd', 'Pre-tax cost of debt', f'=MAX({wc.ref("kd_hist")},{M("rf")}+{M("spread")})', PCT, formula=True)
    wc.add('kd_at', 'After-tax cost of debt', f'={wc.ref("kd")}*(1-{M("tax")})', PCT, formula=True)
    wc.section('Result')
    wc.add('wacc', 'WACC', f'={wc.ref("we")}*{wc.ref("ke")}+{wc.ref("wd")}*{wc.ref("kd_at")}', PCT, formula=True)
    W = wc.ref('wacc')
    for j in range(F0, n):
        p = j - 1
        dc.set('ebit', j, f'={I("ebit", j)}', bold=True)
        dc.set('taxes', j, f'=-MAX(0,{P("ebit", j)})*{M("tax")}')
        dc.set('nopat', j, f'={P("ebit", j)}+{P("taxes", j)}', bold=True)
        dc.set('da', j, f'={C("da", j)}')
        dc.set('capex', j, f'={C("capex", j)}')
        dc.set('wc', j, f'={C("wc", j)}')
        dc.set('ufcf', j, f'={P("nopat", j)}+{P("da", j)}+{P("capex", j)}+{P("wc", j)}', bold=True)
        dc.set('ebitda', j, f'={I("ebitda", j)}')
        dc.set('year', j, f'={I("yr", j)}', '0')
        dc.set('period', j, f'={P("year", j)}-{M("mid")}/2', '0.0')
        dc.set('df', j, f'=1/(1+{W})^{P("period", j)}', '0.0000')
        dc.set('pv', j, f'={P("ufcf", j)}*{P("df", j)}', bold=True)
    ufcf_n, ebitda_n, yr_n = P('ufcf', FN), P('ebitda', FN), P('year', FN)
    ufcf_rng, per_rng = dc.span('ufcf', F0, FN), dc.span('period', F0, FN)
    va = Inputs(wb, 'Valuation', head('Enterprise and equity value'))
    va.section('Perpetuity growth method')
    va.add('sum_pv', 'Sum of PV of forecast cash flows', f'=SUM({dc.span("pv", F0, FN)})', NUM, formula=True)
    va.add('tv_g', 'Terminal value (Gordon growth)', f'={ufcf_n}*(1+{M("g")})/({W}-{M("g")})', NUM, formula=True)
    va.add('pv_g', 'PV of terminal value', f'={va.ref("tv_g")}/(1+{W})^{yr_n}', NUM, formula=True)
    va.add('ev_g', 'Enterprise value', f'={va.ref("sum_pv")}+{va.ref("pv_g")}', NUM, formula=True)
    va.add('imp_mult', 'Implied exit EV / EBITDA', f'=IFERROR({va.ref("tv_g")}/{ebitda_n},0)', MULT, formula=True)
    va.section('Exit multiple method')
    va.add('tv_m', 'Terminal value (exit multiple x final EBITDA)', f'={M("exit")}*{ebitda_n}', NUM, formula=True)
    va.add('pv_m', 'PV of terminal value', f'={va.ref("tv_m")}/(1+{W})^{yr_n}', NUM, formula=True)
    va.add('ev_m', 'Enterprise value', f'={va.ref("sum_pv")}+{va.ref("pv_m")}', NUM, formula=True)
    va.add('imp_g', 'Implied perpetual growth',
           f'=IFERROR(({va.ref("tv_m")}*{W}-{ufcf_n})/({va.ref("tv_m")}+{ufcf_n}),0)', PCT, formula=True)
    va.section('Equity bridge (last reported balance sheet)')
    va.add('debt', 'Less: total debt', f'=-{wc.ref("debt")}', NUM, formula=True)
    va.add('nci', 'Less: non-controlling interest & other equity', f'=-{B("nci", hi)}', NUM, formula=True)
    va.add('cash', 'Plus: cash and equivalents', f'={B("cash", hi)}', NUM, formula=True)
    va.add('bridge', 'Net adjustment', f'={va.ref("debt")}+{va.ref("nci")}+{va.ref("cash")}', NUM, formula=True)
    va.add('eq_g', 'Equity value - perpetuity method', f'={va.ref("ev_g")}+{va.ref("bridge")}', NUM, formula=True)
    va.add('eq_m', 'Equity value - exit multiple method', f'={va.ref("ev_m")}+{va.ref("bridge")}', NUM, formula=True)
    va.add('eq', 'Equity value - midpoint of methods', f'=AVERAGE({va.ref("eq_g")},{va.ref("eq_m")})', NUM, formula=True)
    va.add('ps', 'Value per share (USD)', f'=IFERROR({va.ref("eq")}/{M("shares")},0)', PRICE, formula=True)
    va.add('updown', 'Value vs market value of equity', f'=IFERROR({va.ref("eq")}/{M("mcap")}-1,0)', PCT, formula=True)
    va.section('Implied multiples (last reported year)')
    va.add('ev_rev', 'EV / revenue', f'=IFERROR(AVERAGE({va.ref("ev_g")},{va.ref("ev_m")})/{I("revenue", hi)},0)', MULT, formula=True)
    va.add('ev_ebitda', 'EV / EBITDA', f'=IFERROR(AVERAGE({va.ref("ev_g")},{va.ref("ev_m")})/{I("ebitda", hi)},0)', MULT, formula=True)
    va.add('pe', 'Equity value / net income', f'=IFERROR({va.ref("eq")}/{I("net_income", hi)},0)', MULT, formula=True)

    se = wb.create_sheet('Sensitivity')
    _title(se, head('Equity value sensitivity (USD millions)'),
           'Each cell re-discounts the projected cash flows; centre cells equal the base valuation (checked).')
    se.column_dimensions['A'].width = 16
    steps = [-2, -1, 0, 1, 2]
    bridge = va.ref('bridge')
    se['A4'] = 'WACC (rows) x terminal growth (columns)'
    se['A4'].font = BOLD
    se['A12'] = 'WACC (rows) x exit EV / EBITDA (columns)'
    se['A12'].font = BOLD
    for k, s in enumerate(steps):
        se.cell(row=5, column=2 + k, value=f'={M("g")}+({s})*{M("g_step")}').number_format = PCT
        se.cell(row=13, column=2 + k, value=f'={M("exit")}+({s})*{M("m_step")}').number_format = MULT
        for r0 in (6, 14):
            se.cell(row=r0 + k, column=1, value=f'={W}+({s})*{M("w_step")}').number_format = PCT
        se.column_dimensions[col(k)].width = 14
    for r in range(5):
        for k in range(5):
            w, g, m = f'$A{6 + r}', f'{col(k)}$5', f'{col(k)}$13'
            pv = f'SUMPRODUCT({ufcf_rng}/(1+{w})^{per_rng})'
            se.cell(row=6 + r, column=2 + k,
                    value=f'=IFERROR({pv}+{ufcf_n}*(1+{g})/({w}-{g})/(1+{w})^{yr_n}+{bridge},0)').number_format = NUM
            w2 = f'$A{14 + r}'
            pv2 = f'SUMPRODUCT({ufcf_rng}/(1+{w2})^{per_rng})'
            se.cell(row=14 + r, column=2 + k,
                    value=f'={pv2}+{m}*{ebitda_n}/(1+{w2})^{yr_n}+{bridge}').number_format = NUM
    ff = Inputs(wb, 'Football Field', head('Valuation summary - equity value ranges'))
    ff.ws['D3'], ff.ws['E3'] = 'Low', 'High'
    ff.section('Equity value (USD millions)')
    for key, label, rng in [('g', 'DCF - perpetuity growth (sensitivity range)', 'Sensitivity!$B$6:$F$10'),
                            ('m', 'DCF - exit multiple (sensitivity range)', 'Sensitivity!$B$14:$F$18')]:
        ff.add(key, label, f'=MEDIAN({rng})', NUM, 'Median of the grid; low and high in columns D and E', formula=True)
        r = ff.rows[key]
        ff.ws.cell(row=r, column=4, value=f'=MIN({rng})').number_format = NUM
        ff.ws.cell(row=r, column=5, value=f'=MAX({rng})').number_format = NUM
    ff.add('mkt', 'Market value of equity', f'={M("mcap")}', NUM, formula=True)
    c.setdefault('dash_extra', []).extend([
        ('v_wacc', 'WACC', f'={W}', PCT), ('v_eq', 'DCF equity value (midpoint)', f'={va.ref("eq")}', NUM),
        ('v_ps', 'DCF value per share (USD)', f'={va.ref("ps")}', PRICE)])
    sections = [('Market Inputs', 'Risk-free rate, ERP, beta, SOFR, terminal growth, exit multiple, share data (inputs)'),
                ('DCF', 'Unlevered free cash flow from the integrated forecast and discounting'),
                ('WACC', 'Capital structure, CAPM cost of equity, cost of debt and WACC'),
                ('Valuation', 'Perpetuity and exit-multiple methods, equity bridge, per-share value, implied multiples'),
                ('Sensitivity', 'Equity value across WACC x growth and WACC x exit multiple'),
                ('Football Field', 'Valuation ranges by method')]
    items = [('WACC exceeds terminal growth (0 = yes)', f'IF({W}>{M("g")},0,1)'),
             ('Perpetuity grid centre equals base equity value', f'Sensitivity!$D$8-{va.ref("eq_g")}'),
             ('Exit-multiple grid centre equals base equity value', f'Sensitivity!$D$16-{va.ref("eq_m")}'),
             ('PV of cash flows equals SUMPRODUCT re-calculation', f'{va.ref("sum_pv")}-SUMPRODUCT({ufcf_rng}/(1+{W})^{per_rng})')]
    return sections, items


# --- leveraged buyout ----------------------------------------------------------------------------------------------
def buyout(wb, c, market, sic, facts):
    M = market_inputs(wb, c, market, sic, facts).ref
    I, B, C = c['is_'].ref, c['bs'].ref, c['cf'].ref
    n, nh, periods, years = c['n'], c['nh'], c['periods'], c['years']
    hi, F0, FN = nh - 1, nh, n - 1
    head = c['heading']
    entry = None
    for row in c['mi'].ws.iter_rows(min_row=c['mi'].rows['exit'], max_row=c['mi'].rows['exit'], min_col=2, max_col=2):
        entry = row[0].value
    senior_x, notes_x = round(min(4.0, entry * 0.35), 2), round(min(2.0, entry * 0.15), 2)
    ti = Inputs(wb, 'Transaction Inputs', head('Transaction, financing and exit inputs'))
    ti.section('Purchase price')
    ti.add('entry', 'Entry EV / EBITDA', entry, MULT, 'Same basis as the exit multiple on Market Inputs')
    ti.add('tx_fee', 'Transaction fees % enterprise value', 0.02, PCT)
    ti.add('min_cash', 'Minimum operating cash % revenue', 0.03, PCT)
    ti.section('Financing')
    ti.add('tl_x', 'Term loan B (x LTM EBITDA)', senior_x, MULT, 'Senior secured, floating')
    ti.add('tl_spread', 'Term loan spread over SOFR', 0.035, PCT)
    ti.add('tl_amort', 'Mandatory amortisation % original principal p.a.', 0.01, PCT)
    ti.add('sweep', 'Excess cash flow sweep %', 0.75, PCT)
    ti.add('sn_x', 'Senior notes (x LTM EBITDA)', notes_x, MULT, 'Unsecured, fixed coupon, bullet')
    ti.add('sn_rate', 'Senior notes coupon', round(market.rf + 0.035, 4), PCT, 'Risk-free rate plus 350 bp')
    ti.add('fin_fee', 'Financing fees % of debt raised', 0.025, PCT)
    ti.add('rev_spread', 'Revolver spread over SOFR', 0.03, PCT)
    ti.add('mgmt', 'Management equity pool % of exit equity', 0.08, PCT, 'Dilutes the sponsor at exit')
    ti.section('Exit')
    ti.add('exit', 'Exit EV / EBITDA', entry, MULT, 'Equal to entry: no multiple expansion')
    ti.add('step', 'Sensitivity step (x EBITDA)', 1.0, MULT)
    T = ti.ref
    ltm = I('ebitda', hi)
    su = Inputs(wb, 'Sources & Uses', head('Sources and uses of funds at close'))
    su.section('Uses')
    su.add('ev', 'Enterprise value (entry multiple x LTM EBITDA)', f'={T("entry")}*{ltm}', NUM, formula=True)
    su.add('tx_fees', 'Transaction fees', f'={su.ref("ev")}*{T("tx_fee")}', NUM, formula=True)
    su.add('fin_fees', 'Financing fees', f'=({T("tl_x")}+{T("sn_x")})*{ltm}*{T("fin_fee")}', NUM, formula=True)
    su.add('cash_bs', 'Cash to balance sheet', f'={I("revenue", hi)}*{T("min_cash")}', NUM, formula=True)
    su.add('uses', 'Total uses', f'=SUM(B{su.rows["ev"]}:B{su.rows["cash_bs"]})', NUM, formula=True)
    su.section('Sources')
    su.add('tl', 'Term loan B', f'={T("tl_x")}*{ltm}', NUM, formula=True)
    su.add('sn', 'Senior notes', f'={T("sn_x")}*{ltm}', NUM, formula=True)
    su.add('equity', 'Sponsor equity', f'={su.ref("uses")}-{su.ref("tl")}-{su.ref("sn")}', NUM, formula=True)
    su.add('sources', 'Total sources', f'=SUM(B{su.rows["tl"]}:B{su.rows["equity"]})', NUM, formula=True)
    su.add('eq_pct', 'Equity % of sources', f'=IFERROR({su.ref("equity")}/{su.ref("sources")},0)', PCT, formula=True)
    su.add('lev', 'Leverage at close (x LTM EBITDA)', f'=IFERROR(({su.ref("tl")}+{su.ref("sn")})/{ltm},0)', MULT, formula=True)

    lb = Sheet(wb, 'LBO Model', head('Buyout cash flows, debt schedule and credit statistics'), periods, nh)
    for key, label, bold, sec in [
        ('ebitda', 'EBITDA (integrated forecast, live scenario)', True, 'Cash flow'), ('da', 'Depreciation & amortisation', False, None),
        ('ebit', 'EBIT', False, None), ('interest', 'Interest expense (buyout debt)', False, None), ('pretax', 'Pre-tax income', False, None),
        ('taxes', 'Cash taxes', False, None), ('ni', 'Net income', True, None), ('capex', 'Capital expenditure', False, None),
        ('wc', 'Working capital', False, None), ('fcf', 'Free cash flow for debt service', True, None),
        ('cash_open', 'Opening cash', False, 'Cash waterfall'), ('mand', 'Less: mandatory amortisation', False, None),
        ('pre', 'Cash after mandatory debt service', False, None), ('draw', 'Revolver draw / (repayment)', False, None),
        ('avail', 'Cash available for sweep', False, None), ('sweep', 'Less: optional prepayment', False, None),
        ('cash', 'Closing cash', True, None),
        ('rev_open', 'Revolver - opening', False, 'Debt'), ('rev', 'Revolver - closing', False, None), ('rev_int', 'Revolver interest', False, None),
        ('tl_open', 'Term loan - opening', False, None), ('tl_mand', 'Term loan - mandatory amortisation', False, None),
        ('tl_sweep', 'Term loan - optional prepayment', False, None), ('tl', 'Term loan - closing', False, None),
        ('tl_int', 'Term loan interest', False, None), ('sn', 'Senior notes - closing', False, None), ('sn_int', 'Senior notes interest', False, None),
        ('total', 'Total debt', True, 'Credit statistics'), ('net', 'Net debt', False, None), ('lev', 'Total debt / EBITDA', False, None),
        ('cov', 'EBITDA / interest', False, None), ('paid', 'Cumulative debt paydown % of initial debt', False, None)]:
        lb.row(key, label, bold, section=sec)
    L = lb.ref
    lb.set('total', hi, f'={su.ref("tl")}+{su.ref("sn")}', bold=True)
    lb.set('cash', hi, f'={su.ref("cash_bs")}', bold=True)
    lb.set('tl', hi, f'={su.ref("tl")}')
    lb.set('sn', hi, f'={su.ref("sn")}')
    lb.set('rev', hi, f'={su.ref("tl")}-{su.ref("tl")}')
    lb.set('ebitda', hi, f'={ltm}', bold=True)
    for j in range(F0, n):
        p = j - 1
        lb.set('ebitda', j, f'={I("ebitda", j)}', bold=True)
        lb.set('da', j, f'={C("da", j)}')
        lb.set('ebit', j, f'={L("ebitda", j)}-{L("da", j)}')
        lb.set('interest', j, f'={L("rev_int", j)}+{L("tl_int", j)}+{L("sn_int", j)}')
        lb.set('pretax', j, f'={L("ebit", j)}-{L("interest", j)}')
        lb.set('taxes', j, f'=-MAX(0,{L("pretax", j)})*{M("tax")}')
        lb.set('ni', j, f'={L("pretax", j)}+{L("taxes", j)}', bold=True)
        lb.set('capex', j, f'={C("capex", j)}')
        lb.set('wc', j, f'={C("wc", j)}')
        lb.set('fcf', j, f'={L("ni", j)}+{L("da", j)}+{L("capex", j)}+{L("wc", j)}', bold=True)
        lb.set('cash_open', j, f'={L("cash", p)}')
        lb.set('mand', j, f'=-{L("tl_mand", j)}')
        lb.set('pre', j, f'={L("cash_open", j)}+{L("fcf", j)}+{L("mand", j)}')
        lb.set('draw', j, f'=MAX(-{L("rev_open", j)},{I("revenue", j)}*{T("min_cash")}-{L("pre", j)})')
        lb.set('avail', j, f'=MAX(0,{L("pre", j)}+{L("draw", j)}-{I("revenue", j)}*{T("min_cash")})')
        lb.set('sweep', j, f'=-{L("tl_sweep", j)}')
        lb.set('cash', j, f'={L("pre", j)}+{L("draw", j)}+{L("sweep", j)}', bold=True)
        lb.set('rev_open', j, f'={L("rev", p)}')
        lb.set('rev', j, f'={L("rev_open", j)}+{L("draw", j)}')
        lb.set('rev_int', j, f'={L("rev_open", j)}*({M("sofr")}+{T("rev_spread")})')
        lb.set('tl_open', j, f'={L("tl", p)}')
        lb.set('tl_mand', j, f'=MIN({L("tl_open", j)},{su.ref("tl")}*{T("tl_amort")})')
        lb.set('tl_sweep', j, f'=MIN({L("tl_open", j)}-{L("tl_mand", j)},{L("avail", j)}*{T("sweep")})')
        lb.set('tl', j, f'={L("tl_open", j)}-{L("tl_mand", j)}-{L("tl_sweep", j)}')
        lb.set('tl_int', j, f'={L("tl_open", j)}*({M("sofr")}+{T("tl_spread")})')
        lb.set('sn', j, f'={L("sn", p)}')
        lb.set('sn_int', j, f'={L("sn", p)}*{T("sn_rate")}')
        lb.set('total', j, f'={L("rev", j)}+{L("tl", j)}+{L("sn", j)}', bold=True)
        lb.set('net', j, f'={L("total", j)}-{L("cash", j)}')
        lb.set('lev', j, f'=IFERROR({L("total", j)}/{L("ebitda", j)},0)', MULT)
        lb.set('cov', j, f'=IFERROR({L("ebitda", j)}/{L("interest", j)},0)', MULT)
        lb.set('paid', j, f'=IFERROR(1-{L("total", j)}/{L("total", hi)},0)', PCT)
    rs = Inputs(wb, 'Returns', head('Sponsor returns'))
    rs.section('Exit')
    rs.add('ebitda', 'Exit-year EBITDA', f'={L("ebitda", FN)}', NUM, formula=True)
    rs.add('ev', 'Exit enterprise value', f'={rs.ref("ebitda")}*{T("exit")}', NUM, formula=True)
    rs.add('net', 'Less: net debt at exit', f'=-{L("net", FN)}', NUM, formula=True)
    rs.add('eq100', 'Exit equity value (100%)', f'={rs.ref("ev")}+{rs.ref("net")}', NUM, formula=True)
    rs.add('mgmt', 'Less: management equity pool', f'=-MAX(0,{rs.ref("eq100")})*{T("mgmt")}', NUM, formula=True)
    rs.add('eq', 'Sponsor exit equity', f'={rs.ref("eq100")}+{rs.ref("mgmt")}', NUM, formula=True)
    rs.section('Returns')
    rs.add('moic', 'MOIC', f'=IFERROR({rs.ref("eq")}/{su.ref("equity")},0)', MULT, formula=True)
    rs.add('irr', 'IRR', f'=IFERROR(MAX(0,{rs.ref("moic")})^(1/{years})-1,0)', PCT, formula=True)
    r0 = rs.next_row + 1
    rs.ws.cell(row=r0, column=1, value='Sponsor equity cash flows').font = Font(bold=True, color=NAVY)
    for k in range(years + 1):
        rs.ws.cell(row=r0 + 1, column=2 + k, value=f"='LBO Model'!{col(hi + k)}3").font = BOLD
        v = f'=-{su.ref("equity")}' if k == 0 else (f'={rs.ref("eq")}' if k == years else f'={su.ref("equity")}-{su.ref("equity")}')
        rs.ws.cell(row=r0 + 2, column=2 + k, value=v).number_format = NUM
    flows = f"'Returns'!B{r0 + 2}:{col(years)}{r0 + 2}"
    rs.next_row = r0 + 4
    rs.add('irr_xl', 'IRR (Excel IRR on the cash flows)', f'=IFERROR(IRR({flows}),{rs.ref("irr")})', PCT, formula=True)
    rs.add('npv', 'Sponsor NPV at the cost of equity', f'=NPV({M("rf")}+{M("bu")}*{M("erp")},{rs.ws.cell(row=r0 + 2, column=3).coordinate}:'
                                                         f'{col(years)}{r0 + 2})+{rs.ws.cell(row=r0 + 2, column=2).coordinate}', NUM, formula=True)
    se = wb.create_sheet('IRR Sensitivity')
    _title(se, head('IRR: entry multiple (rows) x exit multiple (columns)'),
           'Each cell rebuilds the sponsor cheque and exit proceeds; the centre equals the base IRR (checked).')
    se.column_dimensions['A'].width = 14
    other = f'({su.ref("fin_fees")}+{su.ref("cash_bs")}-{su.ref("tl")}-{su.ref("sn")})'
    for k, s in enumerate([-2, -1, 0, 1, 2]):
        se.cell(row=4, column=2 + k, value=f'={T("exit")}+({s})*{T("step")}').number_format = MULT
        se.cell(row=5 + k, column=1, value=f'={T("entry")}+({s})*{T("step")}').number_format = MULT
        se.column_dimensions[col(k)].width = 12
    for r in range(5):
        for k in range(5):
            e, x = f'$A{5 + r}', f'{col(k)}$4'
            eq_in = f'({e}*{ltm}*(1+{T("tx_fee")})+{other})'
            eq_out = f'({x}*{rs.ref("ebitda")}+{rs.ref("net")})*(1-{T("mgmt")})'
            se.cell(row=5 + r, column=2 + k,
                    value=f'=IFERROR(MAX(0,{eq_out}/{eq_in})^(1/{years})-1,0)').number_format = PCT
    c.setdefault('dash_extra', []).extend([('b_irr', 'Buyout IRR', f'={rs.ref("irr")}', PCT),
                                           ('b_moic', 'Buyout MOIC', f'={rs.ref("moic")}', MULT)])
    sections = [('Market Inputs', 'Rates, beta, tax and multiples (inputs)'),
                ('Transaction Inputs', 'Purchase price, financing terms, fees, management pool and exit (inputs)'),
                ('Sources & Uses', 'Funding of the purchase at close'),
                ('LBO Model', 'Buyout cash flows from the integrated forecast, cash sweep, debt and credit statistics'),
                ('Returns', 'Exit value, MOIC, IRR (two ways) and sponsor NPV'),
                ('IRR Sensitivity', 'IRR across entry and exit multiples')]
    items = [('Sources equal uses', f'{su.ref("sources")}-{su.ref("uses")}'),
             ('Sensitivity centre equals base IRR', f"('IRR Sensitivity'!$D$7-{rs.ref('irr')})*100"),
             ('Closing cash never below zero', f'-MIN(0,MIN({lb.span("cash", F0, FN)}))'),
             ('Term loan never negative', f'-MIN(0,MIN({lb.span("tl", F0, FN)}))')]
    return sections, items


# --- budget & FP&A -----------------------------------------------------------------------------------------------
def budget(wb, c, weights):
    I, C = c['is_'].ref, c['cf'].ref
    nh, periods = c['nh'], c['periods']
    hi, f1 = nh - 1, nh
    head = c['heading']
    bi = Inputs(wb, 'Budget Inputs', head(f'Budget inputs for {periods[f1]}'))
    bi.section('Monthly revenue seasonality (share of the year)')
    for m, w in zip(MONTHS, weights):
        bi.add(f'w_{m}', m, w, PCT, "From the company's quarterly revenue pattern, spread evenly within each quarter"
               if m == 'Jan' and weights[0] != round(1 / 12, 6) else None)
    bi.add('var', 'Variable share of operating expenses', 0.5, PCT, 'Follows monthly revenue; the rest is spread evenly')
    bi.section('Department allocation of operating expenses')
    depts = [('sm', 'Sales & marketing', 0.35), ('rd', 'Research & development', 0.20), ('ga', 'General & administrative', 0.25),
             ('ops', 'Operations & supply chain', 0.20)]
    for key, label, share in depts:
        bi.add(f'd_{key}', label, share, PCT, 'Assumed split; adjust to the company\'s cost structure')
    bi.add('fte_cost', 'Average fully loaded cost per employee (USD millions)', 0.12, NUM)
    W = {m: bi.ref(f'w_{m}') for m in MONTHS}
    mb = Sheet(wb, 'Monthly Budget', head(f'Monthly budget {periods[f1]} (live scenario)'), MONTHS + ['FY total'])
    mb.ws.column_dimensions[col(12)].width = 14
    mrows = [('revenue', 'Revenue', True, 'Income statement'), ('cogs', 'Cost of revenue', False, None), ('gross', 'Gross profit', True, None)]
    mrows += [(f'opex_{k}', f'  {label}', False, 'Operating expenses' if i == 0 else None) for i, (k, label, _) in enumerate(depts)]
    mrows += [('opex', 'Total operating expenses', True, None), ('ebit', 'Operating income (EBIT)', True, None),
              ('da', 'Depreciation & amortisation', False, None), ('ebitda', 'EBITDA', True, None),
              ('gm', 'Gross margin', False, 'Margins and progress'), ('om', 'EBIT margin', False, None),
              ('ytd', 'Year-to-date revenue', False, None), ('ytd_pct', 'Year-to-date % of annual budget', False, None)]
    for key, label, bold, sec in mrows:
        mb.row(key, label, bold, section=sec)
    Mb = mb.ref
    for j, m in enumerate(MONTHS):
        mb.set('revenue', j, f'={I("revenue", f1)}*{W[m]}', bold=True)
        mb.set('cogs', j, f'={Mb("revenue", j)}*IFERROR({I("cogs", f1)}/{I("revenue", f1)},0)')
        mb.set('gross', j, f'={Mb("revenue", j)}-{Mb("cogs", j)}', bold=True)
        for key, _, _ in depts:
            mb.set(f'opex_{key}', j, f'={I("opex", f1)}*{bi.ref("d_" + key)}*({bi.ref("var")}*{W[m]}+(1-{bi.ref("var")})/12)')
        mb.set('opex', j, f'=SUM({col(j)}{mb.rows["opex_" + depts[0][0]]}:{col(j)}{mb.rows["opex_" + depts[-1][0]]})', bold=True)
        mb.set('ebit', j, f'={Mb("gross", j)}-{Mb("opex", j)}', bold=True)
        mb.set('da', j, f'={C("da", f1)}/12')
        mb.set('ebitda', j, f'={Mb("ebit", j)}+{Mb("da", j)}', bold=True)
        mb.set('gm', j, f'=IFERROR({Mb("gross", j)}/{Mb("revenue", j)},0)', PCT)
        mb.set('om', j, f'=IFERROR({Mb("ebit", j)}/{Mb("revenue", j)},0)', PCT)
        mb.set('ytd', j, f'=SUM({mb.span("revenue", 0, j)})')
        mb.set('ytd_pct', j, f'=IFERROR({Mb("ytd", j)}/{Mb("revenue", 12)},0)', PCT)
    for key in ['revenue', 'cogs', 'gross'] + [f'opex_{k}' for k, _, _ in depts] + ['opex', 'ebit', 'da', 'ebitda']:
        mb.set(key, 12, f'=SUM({mb.span(key, 0, 11)})', bold=True)
    mb.set('gm', 12, f'=IFERROR({Mb("gross", 12)}/{Mb("revenue", 12)},0)', PCT)
    mb.set('om', 12, f'=IFERROR({Mb("ebit", 12)}/{Mb("revenue", 12)},0)', PCT)

    dp = Sheet(wb, 'Department Plan', head('Operating expense plan by department'), periods[f1:], 0)
    for key, label, _ in depts:
        dp.row(key, label)
    dp.row('total', 'Total operating expenses', bold=True)
    dp.row('pct', 'Operating expenses % revenue', gap_before=True)
    for key, label, _ in depts:
        dp.row(f'fte_{key}', f'Implied employees - {label}', section='Implied headcount' if key == depts[0][0] else None)
    dp.row('fte', 'Implied employees - total', bold=True)
    for j in range(len(periods) - f1):
        jj = f1 + j
        for key, _, _ in depts:
            dp.set(key, j, f'={I("opex", jj)}*{bi.ref("d_" + key)}')
            dp.set(f'fte_{key}', j, f'=IFERROR({dp.ref(key, j)}/{bi.ref("fte_cost")},0)', '#,##0')
        dp.set('total', j, f'=SUM({col(j)}{dp.rows[depts[0][0]]}:{col(j)}{dp.rows[depts[-1][0]]})', bold=True)
        dp.set('pct', j, f'=IFERROR({dp.ref("total", j)}/{I("revenue", jj)},0)', PCT)
        dp.set('fte', j, f'=SUM({col(j)}{dp.rows["fte_" + depts[0][0]]}:{col(j)}{dp.rows["fte_" + depts[-1][0]]})', '#,##0', bold=True)

    va = Sheet(wb, 'Variance Analysis', head(f'Budget {periods[f1]} vs actual {periods[hi]}'),
               ['Actual', 'Budget', 'Variance', 'Variance %', 'Assessment'])
    va.ws.column_dimensions[col(4)].width = 16
    lines = [('revenue', 'Revenue', I, 'revenue', 1), ('cogs', 'Cost of revenue', I, 'cogs', -1), ('gross', 'Gross profit', I, 'gross', 1),
             ('opex', 'Operating expenses', I, 'opex', -1), ('ebit', 'Operating income', I, 'ebit', 1),
             ('ebitda', 'EBITDA', I, 'ebitda', 1), ('ni', 'Net income', I, 'net_income', 1),
             ('cfo', 'Cash from operations', C, 'cfo', 1), ('capex', 'Capital expenditure', C, 'capex', 1)]
    for key, label, _, _, _ in lines:
        va.row(key, label)
    for key, _, ref, src, sign in lines:
        va.set(key, 0, f'={ref(src, hi)}')
        va.set(key, 1, f'={ref(src, f1)}')
        va.set(key, 2, f'={va.ref(key, 1)}-{va.ref(key, 0)}')
        va.set(key, 3, f'=IFERROR({va.ref(key, 2)}/ABS({va.ref(key, 0)}),0)', PCT)
        good = f'{va.ref(key, 2)}>=0' if sign > 0 else f'{va.ref(key, 2)}<=0'
        va.set(key, 4, f'=IF({good},"Favourable","Unfavourable")', '@')
    c.setdefault('dash_extra', []).extend([
        ('bud_rev', f'Budget revenue {periods[f1]}', f'={Mb("revenue", 12)}', NUM),
        ('bud_fte', 'Implied employees (budget year)', f"={dp.ref('fte', 0)}", '#,##0')])
    sections = [('Budget Inputs', 'Monthly seasonality, cost behaviour, department split and cost per employee (inputs)'),
                ('Monthly Budget', 'Twelve-month budget for the first forecast year with year-to-date tracking'),
                ('Department Plan', 'Operating expenses and implied headcount by department for every forecast year'),
                ('Variance Analysis', 'Budget versus last reported year with favourable / unfavourable flags')]
    items = [('Seasonality weights sum to 100%', f'(SUM({bi.ref("w_Jan")}:{bi.ref("w_Dec").split("!")[1]})-1)*100'),
             ('Department shares sum to 100%',
              f'(SUM({bi.ref("d_" + depts[0][0])}:{bi.ref("d_" + depts[-1][0]).split("!")[1]})-1)*100'),
             ('Monthly revenue adds up to the annual forecast', f'{Mb("revenue", 12)}-{I("revenue", f1)}'),
             ('Monthly operating expenses add up to the annual forecast', f'{Mb("opex", 12)}-{I("opex", f1)}'),
             ('Monthly EBIT adds up to the annual forecast', f'{Mb("ebit", 12)}-{I("ebit", f1)}')]
    return sections, items


# --- credit & liquidity --------------------------------------------------------------------------------------------
def credit(wb, c):
    I, B, R = c['is_'].ref, c['bs'].ref, c['rt'].ref
    n, nh, periods = c['n'], c['nh'], c['periods']
    head = c['heading']
    ci = Inputs(wb, 'Credit Inputs', head('Covenants, liquidity and rating inputs'))
    ci.section('Covenants')
    ci.add('max_lev', 'Maximum net debt / EBITDA', 3.5, MULT, 'Typical investment-grade style maintenance covenant')
    ci.add('min_cov', 'Minimum EBITDA / interest', 3.0, MULT)
    ci.add('target', 'Target total debt / EBITDA for debt capacity', 3.0, MULT)
    ci.section('Liquidity')
    ci.add('rcf', 'Revolving credit commitment % revenue', 0.10, PCT, 'Committed facility size')
    ci.add('min_liq', 'Minimum liquidity % revenue', 0.05, PCT)
    ws = ci.ws
    r = ci.next_row + 1
    ws.cell(row=r, column=1, value='Synthetic rating table (EBIT / interest at least ...)').font = Font(bold=True, color=NAVY)
    for k, h in enumerate(['Coverage from', 'Rating', 'Spread']):
        ws.cell(row=r + 1, column=1 + k, value=h).font = BOLD
    for i, (lo, rating, spread) in enumerate(RATINGS):
        for k, (v, fmt) in enumerate(((lo, '0.00'), (rating, '@'), (spread, PCT))):
            cell = ws.cell(row=r + 2 + i, column=1 + k, value=v)
            cell.font, cell.fill, cell.number_format = INPUT, INPUT_FILL, fmt
    t0, t1 = r + 2, r + 1 + len(RATINGS)
    cov_rng, rat_rng, spr_rng = (f"'Credit Inputs'!${x}${t0}:${x}${t1}" for x in 'ABC')
    C_ = ci.ref
    ca = Sheet(wb, 'Credit Analysis', head('Credit statistics, covenants and liquidity (live scenario)'), periods, nh)
    for key, label, bold, sec in [
        ('ebitda', 'EBITDA', True, 'Earnings and debt'), ('ebit', 'EBIT', False, None), ('interest', 'Interest expense', False, None),
        ('debt', 'Total debt (incl. revolver)', False, None), ('cash', 'Cash', False, None), ('net', 'Net debt', True, None),
        ('lev', 'Total debt / EBITDA', False, 'Leverage and coverage'), ('nlev', 'Net debt / EBITDA', False, None),
        ('cov', 'EBITDA / interest', False, None), ('ecov', 'EBIT / interest', False, None),
        ('fcf_debt', 'Free cash flow / total debt', False, None), ('dcap', 'Debt / (debt + equity)', False, None),
        ('rating', 'Synthetic rating (EBIT coverage)', True, 'Rating'), ('spread', 'Implied credit spread', False, None),
        ('lev_head', 'Leverage covenant headroom (x)', False, 'Covenant tests'), ('lev_test', 'Leverage covenant', False, None),
        ('cov_head', 'Coverage covenant headroom (x)', False, None), ('cov_test', 'Coverage covenant', False, None),
        ('rcf', 'Revolver commitment', False, 'Liquidity'), ('undrawn', 'Undrawn revolver', False, None),
        ('liq', 'Total liquidity (cash + undrawn revolver)', True, None), ('liq_pct', 'Liquidity % revenue', False, None),
        ('liq_test', 'Minimum liquidity test', False, None),
        ('capacity', 'Debt capacity at target leverage', False, 'Debt capacity'), ('headroom', 'Additional debt capacity / (excess)', True, None)]:
        ca.row(key, label, bold, section=sec)
    A_ = ca.ref
    for j in range(n):
        ca.set('ebitda', j, f'={I("ebitda", j)}', bold=True)
        ca.set('ebit', j, f'={I("ebit", j)}')
        ca.set('interest', j, f'={I("interest", j)}')
        ca.set('debt', j, f'={B("debt", j)}+{B("revolver", j)}')
        ca.set('cash', j, f'={B("cash", j)}')
        ca.set('net', j, f'={A_("debt", j)}-{A_("cash", j)}', bold=True)
        ca.set('lev', j, f'=IFERROR({A_("debt", j)}/{A_("ebitda", j)},0)', MULT)
        ca.set('nlev', j, f'=IFERROR({A_("net", j)}/{A_("ebitda", j)},0)', MULT)
        ca.set('cov', j, f'=IFERROR({A_("ebitda", j)}/{A_("interest", j)},99)', MULT)
        ca.set('ecov', j, f'=IF({A_("interest", j)}<=0,99,{A_("ebit", j)}/{A_("interest", j)})', MULT)
        ca.set('fcf_debt', j, f'=IFERROR({R("fcf", j)}/{A_("debt", j)},0)', PCT)
        ca.set('dcap', j, f'=IFERROR({A_("debt", j)}/({A_("debt", j)}+{B("te", j)}),0)', PCT)
        ca.set('rating', j, f'=INDEX({rat_rng},MATCH({A_("ecov", j)},{cov_rng},1))', '@', bold=True)
        ca.set('spread', j, f'=INDEX({spr_rng},MATCH({A_("ecov", j)},{cov_rng},1))', PCT)
        ca.set('lev_head', j, f'={C_("max_lev")}-{A_("nlev", j)}', MULT)
        ca.set('lev_test', j, f'=IF(AND({A_("ebitda", j)}>0,{A_("nlev", j)}<={C_("max_lev")}),"Pass","Breach")', '@')
        ca.set('cov_head', j, f'={A_("cov", j)}-{C_("min_cov")}', MULT)
        ca.set('cov_test', j, f'=IF({A_("cov", j)}>={C_("min_cov")},"Pass","Breach")', '@')
        ca.set('rcf', j, f'={I("revenue", j)}*{C_("rcf")}')
        ca.set('undrawn', j, f'=MAX(0,{A_("rcf", j)}-{B("revolver", j)})')
        ca.set('liq', j, f'={A_("cash", j)}+{A_("undrawn", j)}', bold=True)
        ca.set('liq_pct', j, f'=IFERROR({A_("liq", j)}/{I("revenue", j)},0)', PCT)
        ca.set('liq_test', j, f'=IF({A_("liq_pct", j)}>={C_("min_liq")},"Pass","Shortfall")', '@')
        ca.set('capacity', j, f'=MAX(0,{A_("ebitda", j)})*{C_("target")}')
        ca.set('headroom', j, f'={A_("capacity", j)}-{A_("debt", j)}', bold=True)
    fn = n - 1
    c.setdefault('dash_extra', []).extend([
        ('c_rating', 'Synthetic rating, final forecast year', f'={A_("rating", fn)}', '@'),
        ('c_breach', 'Covenant breaches in the forecast',
         f'=COUNTIF({ca.span("lev_test", nh, fn)},"Breach")+COUNTIF({ca.span("cov_test", nh, fn)},"Breach")', '0'),
        ('c_liq', 'Liquidity, final forecast year', f'={A_("liq", fn)}', NUM)])
    sections = [('Credit Inputs', 'Covenant levels, liquidity requirements and the synthetic rating table (inputs)'),
                ('Credit Analysis', 'Leverage, coverage, synthetic rating, covenant tests, liquidity and debt capacity')]
    items = [(f'Rating lookup resolves ({periods[j]})', f'IF(ISNA(MATCH({A_("ecov", j)},{cov_rng},1)),1,0)') for j in range(nh, n)]
    return sections, items


def seasonality(facts):
    """Twelve monthly revenue weights from the latest reported quarters (each quarter spread evenly)."""
    try:
        got = xbrl.quarterly(facts, min_years=1, max_years=3)
    except Exception:  # noqa: BLE001
        got = None
    if got:
        _, q = got
        shares = [0.0] * 4
        years = 0
        for qs, fy in zip(q['revenue']['q'], q['revenue']['fy']):
            if fy and fy > 0 and all(v is not None and v >= 0 for v in qs) and sum(qs) < fy:
                parts = qs + [fy - sum(qs)]
                shares = [a + b / fy for a, b in zip(shares, parts)]
                years += 1
        if years and min(shares) > 0:
            w = [round(s / years / 3, 6) for s in shares for _ in range(3)]
            w[-1] = round(1 - sum(w[:-1]), 6)
            return w
    w = [round(1 / 12, 6)] * 12
    w[-1] = round(1 - sum(w[:-1]), 6)
    return w


HORIZONS = {'integrated': [5, 6, 7], 'valuation': [7, 8, 10], 'buyout': [5, 6, 7], 'budget': [5, 6], 'credit': [5, 6, 7],
            'strategic': [10, 12, 15]}


def build(facts, sector, market, sic, kind, trim=0, profile=None):
    s = _sic(sic)
    fin = s in FIN
    if kind in ('valuation', 'buyout', 'credit') and fin:
        return None
    years = profile.pick(HORIZONS[kind]) if profile else (10 if kind == 'strategic' else 5)
    mods = {'integrated': [], 'valuation': ['valuation'], 'buyout': ['buyout'], 'budget': ['budget'],
            'credit': ['credit'], 'strategic': ([] if fin else ['valuation', 'credit']) + ['budget']}[kind]

    def extend(wb, c):
        d, hi = c['data'], c['nh'] - 1
        wb.asof = c['ends'][-1][:4]
        if 'buyout' in mods:
            ebitda = d['ebit'][hi] + d['da'][hi]
            if ebitda <= 0 or ebitda / d['revenue'][hi] < 0.05 or sum(e + a for e, a in zip(d['ebit'], d['da'])) <= 0:
                raise _Skip('EBITDA too thin for a buyout')
        if 'valuation' in mods and (sum(d['ebit']) <= 0 or d['ebit'][hi] <= 0):
            raise _Skip('loss-making history: no cash-flow DCF')
        c['profile'] = profile
        sections, items = [], []
        steps = [lambda: scenarios(wb, c), lambda: schedules(wb, c), lambda: company_lines(wb, c, facts)]
        if not fin:
            steps += [lambda: rnd(wb, c, facts), lambda: sbc(wb, c)]
        for fn in steps:
            a, b = fn()
            sections += a
            items += b
        for m in mods:
            if m == 'valuation':
                a, b = valuation(wb, c, market, sic, facts)
            elif m == 'buyout':
                a, b = buyout(wb, c, market, sic, facts)
            elif m == 'budget':
                a, b = budget(wb, c, seasonality(facts))
            else:
                a, b = credit(wb, c)
            sections += [x for x in a if x[0] not in {s_[0] for s_ in sections}]
            items += b
        a, b = dashboard(wb, c)
        return a + sections, items + b

    label = KINDS[kind]
    desc = (f'Enterprise-wide {label.lower()} built on an integrated three-statement model of the company\'s historical '
            f'financial statements, with a {years}-year forecast. A scenario manager (Base / Upside / Downside) drives revenue growth, '
            'margins, capital intensity and working capital across every statement and module. Supporting schedules, '
            'a KPI dashboard and integrity checks are included; all figures outside the input sheets are formulas.')
    try:
        wb = three.build(facts, sector, is_reit=s == 6798, years=years, trim=trim, extend=extend, deep=True,
                         title=f'Enterprise wide model - {label}', description=desc)
    except _Skip:
        return None
    return finish(wb)


# Models never say where their figures come from (user rule, 2026-10-20); people are "Salvador Sekwa".
PERSON = 'Salvador Sekwa'
NEUTRAL = [(re.compile(r'Reported figures \(Form 10-K\)'), 'Historical figures'),
           (re.compile(r'Reported annual figures from Form 10-K XBRL \(USD millions\) - hard-coded source data'),
            'Historical annual figures (USD millions) - hard-coded inputs'),
           (re.compile(r'USD millions \(as reported, SEC XBRL\)'), 'USD millions'),
           (re.compile(r"company's reported 10-K figures"), "company's historical financial statements")]
SOURCE_WORDS = re.compile(r'10-K|10-Q|\bSEC\b|XBRL|EDGAR|\bFRED\b|Damodaran|public float|cover page|\bdei:|public data|'
                          r'publicly available|sec\.gov|filing|World Bank|\bIBRD\b|Census Bureau|census\.gov|worldbank', re.I)


class SourceMention(Exception):
    pass


class BrokenReference(Exception):
    pass


SHEET_REF = re.compile(r"'((?:[^']|'')+)'!|(?<![A-Za-z0-9_.'])([A-Za-z_][A-Za-z0-9_.]*)!")


def references_ok(wb):
    """Every sheet a formula points at exists, and no formula carries #REF! (user rule: no broken references)."""
    names = set(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith('=')):
                    continue
                if '#REF!' in v.upper():
                    raise BrokenReference(f'{ws.title}!{cell.coordinate}: #REF!')
                body = re.sub(r'"[^"]*"', '""', v)  # ignore text inside string literals
                for quoted, bare in SHEET_REF.findall(body):
                    target = quoted.replace("''", "'") if quoted else bare
                    if target not in names:
                        raise BrokenReference(f'{ws.title}!{cell.coordinate}: sheet {target!r} does not exist')
    for chart_ws in wb.worksheets:
        for chart in getattr(chart_ws, '_charts', []):
            for s in chart.series:
                for ref in (s.val, s.cat, s.tx):
                    f = getattr(getattr(ref, 'numRef', None) or getattr(ref, 'strRef', None), 'f', None)
                    m = f and re.match(r"'?((?:[^'!]|'')+)'?!", f)
                    if m and m.group(1).replace("''", "'") not in names:
                        raise BrokenReference(f'chart on {chart_ws.title}: sheet {m.group(1)!r} does not exist')


def finish(wb):
    """Neutral wording everywhere; refuses a workbook that still names a data source."""
    if wb is None:
        return None
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and not v.startswith('='):  # formulas such as =K10-K12 are not source wording
                    for rx, new in NEUTRAL:
                        v = rx.sub(new, v)
                    if v != cell.value:
                        cell.value = v
                    if SOURCE_WORDS.search(v):
                        raise SourceMention(f'{ws.title}!{cell.coordinate}: {v[:80]}')
    references_ok(wb)
    wb.properties.creator = PERSON
    wb.properties.lastModifiedBy = PERSON
    for field in ('title', 'subject', 'description', 'keywords'):
        if SOURCE_WORDS.search(getattr(wb.properties, field) or ''):
            setattr(wb.properties, field, None)
    return wb


class _Skip(Exception):
    pass
